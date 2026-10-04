"""
V3 Institutional Generator for Training Programs & Question Banks
- Exactly unified, distinct lessons from reference manuals (NO duplicated lessons).
- Each lesson appears ONCE per table.
- Theory is strictly 2.0 hours per lesson ("اما النظري فساعتين فقط").
- Practical is distributed up to 4.0 hours per lesson (0, 2, or 4) for lessons requiring practical training.
- Exact total quotas met:
  - Prep: 56h (28 Th + 28 Pr), 7 days, 14 lessons, 700 questions
  - Med T1: 96h (48 Th + 48 Pr), 12 days, 24 lessons, 1200 questions
  - Med T2: 112h (56 Th + 56 Pr), 14 days, 28 lessons, 1400 questions
  - Final T1: 96h (48 Th + 48 Pr), 12 days, 24 lessons, 1200 questions
  - Final T2: 80h (40 Th + 40 Pr), 10 days, 20 lessons, 1000 questions
"""

import os
import sys
from copy import copy
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

ARABIC_DAYS = [
    "الأول", "الثاني", "الثالث", "الرابع", "الخامس", "السادس",
    "السابع", "الثامن", "التاسع", "العاشر", "الحادي عشر", "الثانى عشر",
    "الثالث عشر", "الرابع عشر", "الخامس عشر", "السادس عشر"
]

def create_thin_border():
    thin = Side(border_style='thin', color='A0A0A0')
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def apply_rtl(ws):
    try:
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = True
    except Exception:
        pass

# =========================================================================
# 1. PREP SCHEDULE (7 days, 14 distinct lessons, 28h Th + 28h Pr = 56h)
# =========================================================================
PREP_SCHEDULE = [
    # Day 1
    {"day": 1, "period": "ف1 - ف2", "title": "التركيب العام والخواص الفنية والتكتيكية للمعدة", "th": 2, "pr": 2, "p_from": 1, "p_to": 3, "q": 50},
    {"day": 1, "period": "ف3 - ف4", "title": "فكرة عامة والغرض من المعدة وأوضاع الإطلاق", "th": 2, "pr": 2, "p_from": 1, "p_to": 2, "q": 50},
    # Day 2
    {"day": 2, "period": "ف1",       "title": "احتياطات الأمان والتحذيرات الواجب مراعاتها عند العمل على المعدة", "th": 2, "pr": 0, "p_from": 3, "p_to": 4, "q": 25},
    {"day": 2, "period": "ف2 - ف4", "title": "إجراءات التعامل مع عدم خروج الصاروخ بعد الضغطة الثانية", "th": 2, "pr": 4, "p_from": 4, "p_to": 5, "q": 75},
    # Day 3
    {"day": 3, "period": "ف1",       "title": "الغرض من الصاروخ ومكوناته الرئيسية والخواص الفنية", "th": 2, "pr": 0, "p_from": 5, "p_to": 6, "q": 25},
    {"day": 3, "period": "ف2 - ف4", "title": "رأس التوجيه الذاتي والجهاز العدسي الجيروسكوبي", "th": 2, "pr": 4, "p_from": 6, "p_to": 8, "q": 75},
    # Day 4 (Midterm exam at day 4)
    {"day": 4, "period": "ف1",       "title": "قرص التعديل والمقاومة الفوتوغرافية وجهاز الموازنة", "th": 2, "pr": 0, "p_from": 8, "p_to": 10, "q": 25},
    {"day": 4, "period": "ف2 - ف4", "title": "ملفات رأس التوجيه الذاتي ونظام التثبيت الكهربي", "th": 2, "pr": 4, "p_from": 10, "p_to": 12, "q": 75},
    # Day 5
    {"day": 5, "period": "ف1",       "title": "نظام التوجيه والتتبع الآلي والدليل الآلي وقنطرة هيوستن", "th": 2, "pr": 0, "p_from": 13, "p_to": 15, "q": 25},
    {"day": 5, "period": "ف2 - ف4", "title": "الجزء الخاص بالدفات ومستودع الغاز ومصدر التغذية المحمول", "th": 2, "pr": 4, "p_from": 15, "p_to": 18, "q": 75},
    # Day 6
    {"day": 6, "period": "ف1",       "title": "القسم القتالي وجهاز التفجير ورأس التدمير والمحركات", "th": 2, "pr": 0, "p_from": 18, "p_to": 22, "q": 25},
    {"day": 6, "period": "ف2 - ف4", "title": "القاذف ومصدر التغذية الأرضي ومواصفاته وفكرة عمله", "th": 2, "pr": 4, "p_from": 22, "p_to": 25, "q": 75},
    # Day 7 (Final exam at day 7)
    {"day": 7, "period": "ف1 - ف2", "title": "مجموعة الإطلاق وخواصها الفنية ومكوناتها وطريقة الاستخدام", "th": 2, "pr": 2, "p_from": 25, "p_to": 27, "q": 50},
    {"day": 7, "period": "ف3 - ف4", "title": "التعاون بين عناصر المعدة ومراحل الإطلاق حتى التدمير", "th": 2, "pr": 2, "p_from": 27, "p_to": 31, "q": 50},
]

# =========================================================================
# 2. MEDIUM TERM 1 (12 days, 24 distinct lessons, 48h Th + 48h Pr = 96h)
# =========================================================================
MED_T1_SCHEDULE = [
    # Day 1
    {"day": 1, "period": "ف1 - ف2", "title": "التركيب العام والخواص الفنية والتكتيكية لمعدة الضبع الأسود", "th": 2, "pr": 2, "p_from": 1, "p_to": 3, "q": 50},
    {"day": 1, "period": "ف3 - ف4", "title": "فكرة عامة عن المعدة وأغراضها وأوضاع الإطلاق", "th": 2, "pr": 2, "p_from": 1, "p_to": 2, "q": 50},
    # Day 2
    {"day": 2, "period": "ف1",       "title": "الخواص الفنية والتكتيكية التفصيلية للمعدة العادية والمعدلة", "th": 2, "pr": 0, "p_from": 2, "p_to": 4, "q": 25},
    {"day": 2, "period": "ف2 - ف4", "title": "احتياطات الأمان والتحذيرات الواجب مراعاتها عند العمل على المعدة", "th": 2, "pr": 4, "p_from": 3, "p_to": 4, "q": 75},
    # Day 3
    {"day": 3, "period": "ف1",       "title": "الغرض من الصاروخ وأوضاع إطلاقه ومكوناته الرئيسية", "th": 2, "pr": 0, "p_from": 5, "p_to": 6, "q": 25},
    {"day": 3, "period": "ف2 - ف4", "title": "إجراءات التعامل مع الأعطال وحالات عدم خروج الصاروخ من القاذف", "th": 2, "pr": 4, "p_from": 4, "p_to": 5, "q": 75},
    # Day 4
    {"day": 4, "period": "ف1",       "title": "رأس التوجيه الذاتي الغرض والخواص الفنية لرأس التوجيه", "th": 2, "pr": 0, "p_from": 6, "p_to": 7, "q": 25},
    {"day": 4, "period": "ف2 - ف4", "title": "الجهاز العدسي الجيروسكوبي ومكوناته ومهامه الفنية", "th": 2, "pr": 4, "p_from": 7, "p_to": 8, "q": 75},
    # Day 5
    {"day": 5, "period": "ف1",       "title": "قرص التعديل وشكل وطريقة عمله وإنتاج إشارة الخطأ", "th": 2, "pr": 0, "p_from": 8, "p_to": 9, "q": 25},
    {"day": 5, "period": "ف2 - ف4", "title": "المقاومة الفوتوغرافية وجهاز موازنة دوران العضو الدوار", "th": 2, "pr": 4, "p_from": 9, "p_to": 10, "q": 75},
    # Day 6 (Midterm exam at day 6)
    {"day": 6, "period": "ف1 - ف2", "title": "ملفات رأس التوجيه الذاتي وتصنيفاتها والمجال المغناطيسي", "th": 2, "pr": 2, "p_from": 10, "p_to": 11, "q": 50},
    {"day": 6, "period": "ف3 - ف4", "title": "نظام أخذ العضو الدوار لسرعته العادية وطريقة العمل", "th": 2, "pr": 2, "p_from": 10, "p_to": 11, "q": 50},
    # Day 7
    {"day": 7, "period": "ف1",       "title": "نظام التثبيت الكهربي ومكوناته وطريقة عمله", "th": 2, "pr": 0, "p_from": 11, "p_to": 12, "q": 25},
    {"day": 7, "period": "ف2 - ف4", "title": "نظام التوجيه والتتبع الآلي وفكرة عمله ونظرية الاقتراب التناسبي", "th": 2, "pr": 4, "p_from": 13, "p_to": 14, "q": 75},
    # Day 8
    {"day": 8, "period": "ف1",       "title": "الدليل الآلي ودائرة العمل الكهربية ومكوناته والمعدل والمستخلص", "th": 2, "pr": 0, "p_from": 14, "p_to": 15, "q": 25},
    {"day": 8, "period": "ف2 - ف4", "title": "جهاز تحديد السرعة الزاوية وفكرة عمله وقنطرة هيوستن", "th": 2, "pr": 4, "p_from": 15, "p_to": 15, "q": 75},
    # Day 9
    {"day": 9, "period": "ف1 - ف2", "title": "الجزء الخاص بالدفات ومكوناته ومستودع الغاز المضغوط", "th": 2, "pr": 2, "p_from": 15, "p_to": 17, "q": 50},
    {"day": 9, "period": "ف3 - ف4", "title": "مصدر التغذية المحمول على سطح الصاروخ ومكوناته والتوربين", "th": 2, "pr": 2, "p_from": 17, "p_to": 18, "q": 50},
    # Day 10
    {"day": 10, "period": "ف1",      "title": "القسم القتالي وجهاز التفجير ووسائل الأمان وطريقة عمله", "th": 2, "pr": 0, "p_from": 18, "p_to": 20, "q": 25},
    {"day": 10, "period": "ف2 - ف4", "title": "رأس التدمير ومكوناتها والعبوة المتفجرة والخاص بالمحركات", "th": 2, "pr": 4, "p_from": 20, "p_to": 22, "q": 75},
    # Day 11
    {"day": 11, "period": "ف1 - ف2", "title": "القاذف الغرض منه والمواصفات العامة وحماية الرامي", "th": 2, "pr": 2, "p_from": 22, "p_to": 23, "q": 50},
    {"day": 11, "period": "ف3 - ف4", "title": "مصدر التغذية الأرضي المواصفات والمكونات والتفاعل الكيميائي", "th": 2, "pr": 2, "p_from": 23, "p_to": 25, "q": 50},
    # Day 12 (Final exam at day 12)
    {"day": 12, "period": "ف1 - ف2", "title": "مجموعة الإطلاق الغرض والخواص الفنية ومكوناتها وذراع الأمان", "th": 2, "pr": 2, "p_from": 25, "p_to": 27, "q": 50},
    {"day": 12, "period": "ف3 - ف4", "title": "التعاون بين عناصر المعدة ومراحل طيران الصاروخ حتى الهدف", "th": 2, "pr": 2, "p_from": 27, "p_to": 31, "q": 50},
]

# =========================================================================
# 3. MEDIUM TERM 2 (14 days, 28 distinct lessons, 56h Th + 56h Pr = 112h)
# =========================================================================
MED_T2_SCHEDULE = [
    # Day 1
    {"day": 1, "period": "ف1 - ف2", "title": "المواصفات الفنية المتقدمة لمعدة الضبع الأسود وتطبيقاتها التكتيكية", "th": 2, "pr": 2, "p_from": 1, "p_to": 3, "q": 50},
    {"day": 1, "period": "ف3 - ف4", "title": "أوضاع الإطلاق التكتيكية وإجراءات التجهيز للرماية في الميدان", "th": 2, "pr": 2, "p_from": 1, "p_to": 3, "q": 50},
    # Day 2
    {"day": 2, "period": "ف1",       "title": "قواعد واحتياطات الأمان المشددة أثناء التعامل مع الصواريخ الحية", "th": 2, "pr": 0, "p_from": 3, "p_to": 4, "q": 25},
    {"day": 2, "period": "ف2 - ف4", "title": "تحليل الأعطال الميدانية وحالات الاستعصاء عند الضغطة الثانية", "th": 2, "pr": 4, "p_from": 4, "p_to": 5, "q": 75},
    # Day 3
    {"day": 3, "period": "ف1",       "title": "المخطط الهيكلي لمنظومة الصاروخ وتوزيع الكتل ومراكز الثقل", "th": 2, "pr": 0, "p_from": 5, "p_to": 6, "q": 25},
    {"day": 3, "period": "ف2 - ف4", "title": "الفحص الفني الدقيق لمكونات رأس التوجيه الذاتي البصرية والميكانيكية", "th": 2, "pr": 4, "p_from": 6, "p_to": 7, "q": 75},
    # Day 4
    {"day": 4, "period": "ف1",       "title": "الديناميكا الحركية للجهاز العدسي الجيروسكوبي ومحاور الدوران", "th": 2, "pr": 0, "p_from": 7, "p_to": 8, "q": 25},
    {"day": 4, "period": "ف2 - ف4", "title": "معايرة وضبط جهاز موازنة دوران العضو الدوار عملياً", "th": 2, "pr": 4, "p_from": 8, "p_to": 10, "q": 75},
    # Day 5
    {"day": 5, "period": "ف1",       "title": "تحليل الدوائر الكهربية لملفات رأس التوجيه وإشارات التوجيه", "th": 2, "pr": 0, "p_from": 10, "p_to": 11, "q": 25},
    {"day": 5, "period": "ف2 - ف4", "title": "اختبارات نظام أخذ العضو الدوار لسرعته المقررة في زمن البدء", "th": 2, "pr": 4, "p_from": 10, "p_to": 11, "q": 75},
    # Day 6
    {"day": 6, "period": "ف1",       "title": "دائرة التثبيت الكهربي واستقرار محور الجيروسكوب أثناء التتبع", "th": 2, "pr": 0, "p_from": 11, "p_to": 12, "q": 25},
    {"day": 6, "period": "ف2 - ف4", "title": "تطبيقات الاقتراب التناسبي ومسارات ملاحقة الأهداف الجوية السريعة", "th": 2, "pr": 4, "p_from": 13, "p_to": 14, "q": 75},
    # Day 7 (Midterm exam at day 7)
    {"day": 7, "period": "ف1 - ف2", "title": "المخطط التفصيلي للدليل الآلي ووظائف التعديل والاستخلاص للإشارة", "th": 2, "pr": 2, "p_from": 14, "p_to": 15, "q": 50},
    {"day": 7, "period": "ف3 - ف4", "title": "فحص قنطرة هيوستن وقياس السرعة الزاوية لخط البصر بدقة", "th": 2, "pr": 2, "p_from": 15, "p_to": 15, "q": 50},
    # Day 8
    {"day": 8, "period": "ف1",       "title": "الدائرة الهيدروليكية والغازية لمشغلات الدفات والزعانف", "th": 2, "pr": 0, "p_from": 15, "p_to": 16, "q": 25},
    {"day": 8, "period": "ف2 - ف4", "title": "إجراءات شحن وتفريغ واختبار مستودع الغاز المضغوط للدفات", "th": 2, "pr": 4, "p_from": 16, "p_to": 17, "q": 75},
    # Day 9
    {"day": 9, "period": "ف1",       "title": "الخصائص الكهربية والكيميائية لمصدر التغذية المحمول وعمله", "th": 2, "pr": 0, "p_from": 17, "p_to": 18, "q": 25},
    {"day": 9, "period": "ف2 - ف4", "title": "اختبار سلامة التوربين الغازي والمولد الكهربي الصغير على الصاروخ", "th": 2, "pr": 4, "p_from": 17, "p_to": 18, "q": 75},
    # Day 10
    {"day": 10, "period": "ف1",      "title": "ميكانيزم تأمين وتعمير جهاز التفجير الميكانيكي والكهربي", "th": 2, "pr": 0, "p_from": 18, "p_to": 20, "q": 25},
    {"day": 10, "period": "ف2 - ف4", "title": "دراسة الرأس الحربية والموجة الانفجارية وتأثير الشظايا الموجهة", "th": 2, "pr": 4, "p_from": 20, "p_to": 22, "q": 75},
    # Day 11
    {"day": 11, "period": "ف1",      "title": "الخصائص البالستية للمحرك الدافع ومحرك المسير وفوهات العادم", "th": 2, "pr": 0, "p_from": 20, "p_to": 22, "q": 25},
    {"day": 11, "period": "ف2 - ف4", "title": "إجراءات الفحص الدوري وتفتيش أنبوب القاذف ونظافة التلامسات", "th": 2, "pr": 4, "p_from": 22, "p_to": 23, "q": 75},
    # Day 12
    {"day": 12, "period": "ف1 - ف2", "title": "تركيب وتوصيل واختبار مصدر التغذية الأرضي ومؤشرات الجاهزية", "th": 2, "pr": 2, "p_from": 23, "p_to": 25, "q": 50},
    {"day": 12, "period": "ف3 - ف4", "title": "الفحص الكهربي لمجموعة الإطلاق ومفاتيح التنشيط والزناد", "th": 2, "pr": 2, "p_from": 25, "p_to": 27, "q": 50},
    # Day 13
    {"day": 13, "period": "ف1 - ف2", "title": "التدريب العملي على تسلسل الإشارات الصوتية والضوئية قبل الإطلاق", "th": 2, "pr": 2, "p_from": 27, "p_to": 29, "q": 50},
    {"day": 13, "period": "ف3 - ف4", "title": "توقيتات وحدة التأخير وانفصال الصاروخ وإشعال المحرك الرئيسي", "th": 2, "pr": 2, "p_from": 29, "p_to": 30, "q": 50},
    # Day 14 (Final exam at day 14)
    {"day": 14, "period": "ف1 - ف2", "title": "تكتيكات الاشتباك في بيئات التشويش الحراري والمناورات الجوية المعقدة", "th": 2, "pr": 2, "p_from": 30, "p_to": 31, "q": 50},
    {"day": 14, "period": "ف3 - ف4", "title": "إجراءات الصيانة الوقائية السنوية والتخزين طويل الأجل للمنظومة", "th": 2, "pr": 2, "p_from": 31, "p_to": 33, "q": 50},
]

# =========================================================================
# 4. FINAL TERM 1 (12 days, 24 distinct lessons, 48h Th + 48h Pr = 96h)
# From مرجع الايجلا
# =========================================================================
FINAL_T1_SCHEDULE = [
    # Day 1
    {"day": 1, "period": "ف1 - ف2", "title": "الخواص الفنية والتكتيكية للنظام الصاروخي ايجلا ومكونات المنظومة", "th": 2, "pr": 2, "p_from": 1, "p_to": 2, "q": 50},
    {"day": 1, "period": "ف3 - ف4", "title": "القدرات التكتيكية لمعدة الايجلا ونظرية التوجيه والتحكم الحديثة", "th": 2, "pr": 2, "p_from": 2, "p_to": 3, "q": 50},
    # Day 2
    {"day": 2, "period": "ف1",       "title": "دور دائرة الإزاحة في زيادة فاعلية الاشتباك مع الطائرات النفاثة", "th": 2, "pr": 0, "p_from": 3, "p_to": 4, "q": 25},
    {"day": 2, "period": "ف2 - ف4", "title": "مناطق الإطلاق والتدمير والعوامل المؤثرة على دقة الإصابة", "th": 2, "pr": 4, "p_from": 4, "p_to": 5, "q": 75},
    # Day 3
    {"day": 3, "period": "ف1",       "title": "الوصف التفصيلي لمكونات الصاروخ إيجلا ورأس التوجيه الذاتي", "th": 2, "pr": 0, "p_from": 5, "p_to": 6, "q": 25},
    {"day": 3, "period": "ف2 - ف4", "title": "رأس التوجيه الذاتي والمسمار الإيروديناميكي وحساسات الأشعة تحت الحمراء", "th": 2, "pr": 4, "p_from": 6, "p_to": 7, "q": 75},
    # Day 4
    {"day": 4, "period": "ف1",       "title": "وحدة منسق التتبع والجيروسكوب ونظرية عمل المستشعر البصري", "th": 2, "pr": 0, "p_from": 7, "p_to": 8, "q": 25},
    {"day": 4, "period": "ف2 - ف4", "title": "وظيفة المقاومات الفوتوغرافية والفلتر المجزأ لتصفية الإشارات الكاذبة", "th": 2, "pr": 4, "p_from": 8, "p_to": 9, "q": 75},
    # Day 5
    {"day": 5, "period": "ف1",       "title": "مكونات وحدة التحكم الإلكترونية ومسارات معالجة الإشارة", "th": 2, "pr": 0, "p_from": 9, "p_to": 10, "q": 25},
    {"day": 5, "period": "ف2 - ف4", "title": "المشغل الميكانيكي وأسطح التحكم بالدفات وآلية التوجيه الهوائي", "th": 2, "pr": 4, "p_from": 9, "p_to": 10, "q": 75},
    # Day 6 (Midterm exam at day 6)
    {"day": 6, "period": "ف1 - ف2", "title": "مصدر التغذية المحمول بالصاروخ والمولد التوربيني الصغير", "th": 2, "pr": 2, "p_from": 11, "p_to": 12, "q": 50},
    {"day": 6, "period": "ف3 - ف4", "title": "جهاز الإحساس بالتغير الزاوي الغرض والمكونات والدائرة الكهربية", "th": 2, "pr": 2, "p_from": 12, "p_to": 13, "q": 50},
    # Day 7
    {"day": 7, "period": "ف1",       "title": "مكبر جهاز الإحساس بالتغير الزاوي ووحدة الإمداد بالغاز المضغوط", "th": 2, "pr": 0, "p_from": 13, "p_to": 14, "q": 25},
    {"day": 7, "period": "ف2 - ف4", "title": "موتور توجيه الصاروخ في المرحلة الابتدائية وإلغاء زاوية الخطأ", "th": 2, "pr": 4, "p_from": 14, "p_to": 15, "q": 75},
    # Day 8
    {"day": 8, "period": "ف1",       "title": "وحدة التسليح ووحدة إفقاد الاستقرار ومراحل الأمان الميكانيكي", "th": 2, "pr": 0, "p_from": 15, "p_to": 16, "q": 25},
    {"day": 8, "period": "ف2 - ف4", "title": "رأس التدمير الغرض والمكونات وتركيب الشحنة المشكلة والمتشظية", "th": 2, "pr": 4, "p_from": 16, "p_to": 17, "q": 75},
    # Day 9
    {"day": 9, "period": "ف1 - ف2", "title": "وسائل الأمان والتفجير ووحدة الإحساس بالهدف المغناطيسية والحرارية", "th": 2, "pr": 2, "p_from": 17, "p_to": 19, "q": 50},
    {"day": 9, "period": "ف3 - ف4", "title": "منظومة المحركات بالصاروخ والمحرك الدافع ومواصفات الدفع النفاث", "th": 2, "pr": 2, "p_from": 20, "p_to": 21, "q": 50},
    # Day 10
    {"day": 10, "period": "ف1",      "title": "المحرك ثنائي الوظيفة ومؤخر عمل الشحنة وغرفة الاحتراق", "th": 2, "pr": 0, "p_from": 21, "p_to": 23, "q": 25},
    {"day": 10, "period": "ف2 - ف4", "title": "وحدة الأجنحة الخلفية ونظام انفراج الريش فور مغادرة القاذف", "th": 2, "pr": 4, "p_from": 23, "p_to": 24, "q": 75},
    # Day 11
    {"day": 11, "period": "ف1 - ف2", "title": "تسلسل عمل أجزاء المعدة قبل مغادرة الصاروخ للقاذف وتغذية النظم", "th": 2, "pr": 2, "p_from": 24, "p_to": 25, "q": 50},
    {"day": 11, "period": "ف3 - ف4", "title": "عمل مكونات الصاروخ أثناء الطيران والتوجيه الذاتي نحو الهدف", "th": 2, "pr": 2, "p_from": 25, "p_to": 27, "q": 50},
    # Day 12 (Final exam at day 12)
    {"day": 12, "period": "ف1 - ف2", "title": "مجموعة الإطلاق المكونات ونظرية العمل وخطوات الإطلاق والزناد", "th": 2, "pr": 2, "p_from": 28, "p_to": 30, "q": 50},
    {"day": 12, "period": "ف3 - ف4", "title": "مصدر التغذية الأرضي والقاذف وحسابات الاشتباك الميداني", "th": 2, "pr": 2, "p_from": 30, "p_to": 33, "q": 50},
]

# =========================================================================
# 5. FINAL TERM 2 (10 days, 20 distinct lessons, 40h Th + 40h Pr = 80h)
# From مرجع الايجلا
# =========================================================================
FINAL_T2_SCHEDULE = [
    # Day 1
    {"day": 1, "period": "ف1 - ف2", "title": "المواصفات التكتيكية المتقدمة لمنظومة إيجلا ومقارنتها بالمنظومات المماثلة", "th": 2, "pr": 2, "p_from": 1, "p_to": 3, "q": 50},
    {"day": 1, "period": "ف3 - ف4", "title": "حسابات مناطق القتل والمناورة والتفادي للأهداف الجوية المعادية", "th": 2, "pr": 2, "p_from": 2, "p_to": 5, "q": 50},
    # Day 2
    {"day": 2, "period": "ف1",       "title": "الفحص الفني والمعايرة الميدانية لرأس التوجيه الذاتي البصري", "th": 2, "pr": 0, "p_from": 5, "p_to": 8, "q": 25},
    {"day": 2, "period": "ف2 - ف4", "title": "تحليل أداء المشغل الميكانيكي والديناميكا الهوائية لدفات التوجيه", "th": 2, "pr": 4, "p_from": 9, "p_to": 11, "q": 75},
    # Day 3
    {"day": 3, "period": "ف1",       "title": "اختبارات كفاءة مصدر التغذية المحمول واستقرار جهد الدوائر الإلكترونية", "th": 2, "pr": 0, "p_from": 11, "p_to": 13, "q": 25},
    {"day": 3, "period": "ف2 - ف4", "title": "التحقق العملي من استجابة جهاز ومكبر التغير الزاوي لإشارات التوجيه", "th": 2, "pr": 4, "p_from": 12, "p_to": 14, "q": 75},
    # Day 4
    {"day": 4, "period": "ف1",       "title": "ديناميكا عمل موتور التوجيه الابتدائي وإجراءات تصحيح مسار الإطلاق", "th": 2, "pr": 0, "p_from": 14, "p_to": 15, "q": 25},
    {"day": 4, "period": "ف2 - ف4", "title": "وسائل التأمين متعددة المراحل لوحدة التسليح ومولد الانفجار", "th": 2, "pr": 4, "p_from": 15, "p_to": 17, "q": 75},
    # Day 5 (Midterm exam at day 5)
    {"day": 5, "period": "ف1 - ف2", "title": "اختبار حساسية وحدة الإحساس بالهدف واستجابة الفيوز التقاربي", "th": 2, "pr": 2, "p_from": 17, "p_to": 19, "q": 50},
    {"day": 5, "period": "ف3 - ف4", "title": "الفحص الميداني للشحنات الدافعة بالمحركين الدافع وثنائي الوظيفة", "th": 2, "pr": 2, "p_from": 20, "p_to": 23, "q": 50},
    # Day 6
    {"day": 6, "period": "ف1",       "title": "معايرة ميكانيزم فتح الأجنحة الخلفية والتأكد من اتزان الصاروخ", "th": 2, "pr": 0, "p_from": 23, "p_to": 24, "q": 25},
    {"day": 6, "period": "ف2 - ف4", "title": "المحاكاة العملية لتسلسل أحداث إطلاق الصاروخ من وضع الكتف", "th": 2, "pr": 4, "p_from": 24, "p_to": 27, "q": 75},
    # Day 7
    {"day": 7, "period": "ف1",       "title": "الفحص الفني لمجموعة الإطلاق ومقابس التوصيل والموصلات الكهربية", "th": 2, "pr": 0, "p_from": 28, "p_to": 30, "q": 25},
    {"day": 7, "period": "ف2 - ف4", "title": "إجراءات التعامل مع ظروف عدم الإطلاق وفك مصدر التغذية الأرضي بأمان", "th": 2, "pr": 4, "p_from": 28, "p_to": 31, "q": 75},
    # Day 8
    {"day": 8, "period": "ف1 - ف2", "title": "تكتيكات الاشتباك مع طائرات الهليكوبتر في وضع الثبات والتحليق المنخفض", "th": 2, "pr": 2, "p_from": 3, "p_to": 6, "q": 50},
    {"day": 8, "period": "ف3 - ف4", "title": "الرماية في مواجهة طائرات مقاتلة تطلق مشاعل حرارية خداعية مكثفة", "th": 2, "pr": 2, "p_from": 6, "p_to": 9, "q": 50},
    # Day 9
    {"day": 9, "period": "ف1 - ف2", "title": "التنسيق والاتصال بين رماة الإيجلا ونقاط المراقبة والإنذار المبكر", "th": 2, "pr": 2, "p_from": 28, "p_to": 31, "q": 50},
    {"day": 9, "period": "ف3 - ف4", "title": "إجراءات الصيانة الأسبوعية والشهرية وفحص عوازل الرطوبة بالقاذف", "th": 2, "pr": 2, "p_from": 31, "p_to": 33, "q": 50},
    # Day 10 (Final exam at day 10)
    {"day": 10, "period": "ف1 - ف2", "title": "التدريب الميداني المتكامل على سيناريوهات الدفاع الجوي عن الأهداف الحيوية", "th": 2, "pr": 2, "p_from": 24, "p_to": 31, "q": 50},
    {"day": 10, "period": "ف3 - ف4", "title": "التفتيش الختامي للجاهزية الفنية والقتالية وإجراءات التخزين السنوي", "th": 2, "pr": 2, "p_from": 31, "p_to": 33, "q": 50},
]


def render_term_table(ws, start_row, term_label, schedule, border_cell, header_font, subheader_font, data_font,
                      header_fill, subheader_fill, midterm_after_day=None):
    """
    Renders training table with V3 clean format:
    - Col A: الترم
    - Col B: اليوم (Merged across the lessons belonging to that day!)
    - Col C: المحاضرة (e.g. ف1 - ف2, ف1, ف2 - ف4)
    - Col D: اسم الموضوع (Distinct, non-repeated lesson title)
    - Col E: نظري (2.0)
    - Col F: عملي (0 / - or 2 or 4)
    - Col G: من (Page from)
    - Col H: الي (Page to)
    - Col I: عدد الاسئلة
    """
    r3 = start_row
    r4 = start_row + 1

    # Row 3 Main Header
    headers_r3 = [
        (1, "الترم"),
        (2, "اليوم"),
        (3, "المحاضرة"),
        (4, "اسم الموضوع"),
        (5, "عدد الساعات"),
        (7, "الصفحة في المرجع الموحد"),
        (9, "عدد الاسئلة")
    ]
    for c_idx, title in headers_r3:
        cell = ws.cell(r3, c_idx, title)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = header_fill
        cell.border = border_cell

    # Merges in Row 3-4
    ws.merge_cells(start_row=r3, start_column=1, end_row=r4, end_column=1) # A3:A4
    ws.merge_cells(start_row=r3, start_column=2, end_row=r4, end_column=2) # B3:B4
    ws.merge_cells(start_row=r3, start_column=3, end_row=r4, end_column=3) # C3:C4
    ws.merge_cells(start_row=r3, start_column=4, end_row=r4, end_column=4) # D3:D4
    ws.merge_cells(start_row=r3, start_column=5, end_row=r3, end_column=6) # E3:F3 (عدد الساعات)
    ws.merge_cells(start_row=r3, start_column=7, end_row=r3, end_column=8) # G3:H3 (الصفحة)
    ws.merge_cells(start_row=r3, start_column=9, end_row=r4, end_column=9) # I3:I4 (عدد الاسئلة)

    # Sub-headers in Row 4
    subheaders = [(5, "نظري"), (6, "عملي"), (7, "من"), (8, "الي")]
    for col_idx, sub_title in subheaders:
        c = ws.cell(r4, col_idx, sub_title)
        c.font = subheader_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = subheader_fill
        c.border = border_cell

    # Ensure borders on all header cells
    for r in [r3, r4]:
        for c in range(1, 10):
            ws.cell(r, c).border = border_cell
            if ws.cell(r, c).fill.fill_type is None:
                ws.cell(r, c).fill = header_fill

    current_r = r4 + 1
    term_start_r = current_r

    # Group schedule items by day
    days_dict = {}
    for item in schedule:
        d = item['day']
        days_dict.setdefault(d, []).append(item)

    total_th = 0
    total_pr = 0
    total_q = 0

    exam_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    exam_font = Font(name="Calibri", size=11, bold=True, color="7F6000")

    for d_num in sorted(days_dict.keys()):
        items_in_day = days_dict[d_num]
        day_start_r = current_r
        day_text = ARABIC_DAYS[d_num - 1] if d_num <= len(ARABIC_DAYS) else f"اليوم {d_num}"

        for it in items_in_day:
            # Col A: Term
            ws.cell(current_r, 1, term_label).alignment = Alignment(horizontal="center", vertical="center")
            # Col B: Day
            ws.cell(current_r, 2, day_text).alignment = Alignment(horizontal="center", vertical="center")
            # Col C: Period / Lecture
            ws.cell(current_r, 3, it['period']).alignment = Alignment(horizontal="center", vertical="center")
            # Col D: Topic
            ws.cell(current_r, 4, it['title']).alignment = Alignment(horizontal="right", vertical="center")
            # Col E: Theory hours
            ws.cell(current_r, 5, it['th']).alignment = Alignment(horizontal="center", vertical="center")
            # Col F: Practical hours
            pr_val = it['pr'] if it['pr'] > 0 else "-"
            ws.cell(current_r, 6, pr_val).alignment = Alignment(horizontal="center", vertical="center")
            # Col G & H: Pages
            ws.cell(current_r, 7, it['p_from']).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(current_r, 8, it['p_to']).alignment = Alignment(horizontal="center", vertical="center")
            # Col I: Questions count
            ws.cell(current_r, 9, it['q']).alignment = Alignment(horizontal="center", vertical="center")

            for c in range(1, 10):
                cell = ws.cell(current_r, c)
                cell.font = data_font
                cell.border = border_cell

            total_th += it['th']
            total_pr += it['pr']
            total_q += it['q']
            current_r += 1

        day_end_r = current_r - 1
        if day_end_r > day_start_r:
            ws.merge_cells(start_row=day_start_r, start_column=2, end_row=day_end_r, end_column=2)

        # Check for midterm exam row
        if midterm_after_day is not None and d_num == midterm_after_day:
            ws.cell(current_r, 1, term_label).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(current_r, 2, "امتحان منتصف الترم").alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells(start_row=current_r, start_column=2, end_row=current_r, end_column=9)
            for c in range(1, 10):
                cell = ws.cell(current_r, c)
                cell.border = border_cell
                cell.fill = exam_fill
                cell.font = exam_font
            current_r += 1

    # Merge Col A for the whole term
    ws.merge_cells(start_row=term_start_r, start_column=1, end_row=current_r - 1, end_column=1)

    # Final Exam / Summary row
    final_r = current_r
    ws.cell(final_r, 1, "")
    ws.cell(final_r, 2, "امتحان ختامى الترم").alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=final_r, start_column=2, end_row=final_r, end_column=4)
    ws.cell(final_r, 5, total_th).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(final_r, 6, total_pr).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(final_r, 7, "-").alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=final_r, start_column=7, end_row=final_r, end_column=8)
    ws.cell(final_r, 9, total_q).alignment = Alignment(horizontal="center", vertical="center")

    summary_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    summary_font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    for c in range(1, 10):
        cell = ws.cell(final_r, c)
        cell.border = border_cell
        cell.fill = summary_fill
        cell.font = summary_font

    return final_r + 1


def build_question_bank_sheet(ws, q_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill):
    apply_rtl(ws)
    headers = ["م", "السؤال", "النوع", "الدرجة", "خيار 1", "خيار 2", "خيار 3", "خيار 4", "الإجابة الصحيحة", "الموضوع", "المحاضرة", "الصفحة"]
    ws.append(headers)
    for col_idx in range(1, 13):
        c = ws.cell(1, col_idx)
        c.font = qb_header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = qb_header_fill
        c.border = border_cell
    ws.row_dimensions[1].height = 28

    for r_idx, row_data in enumerate(q_rows, 2):
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 20
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(r_idx, col_idx)
            c.font = qb_data_font
            c.border = border_cell
            if col_idx in [1, 3, 4, 11, 12]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [2, 5, 6, 7, 8, 9, 10]:
                c.alignment = Alignment(horizontal="right", vertical="center")

    widths = {1: 8, 2: 45, 3: 12, 4: 8, 5: 22, 6: 22, 7: 22, 8: 22, 9: 22, 10: 38, 11: 14, 12: 10}
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def main():
    print("Executing V3 Institutional Generator with clean, non-duplicated lessons...")
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active) # Remove default sheet

    border_cell = create_thin_border()
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    subheader_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    qb_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    qb_data_font = Font(name="Calibri", size=10)

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    subheader_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    qb_header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")

    # =========================================================================
    # 1. SHEET: برنامج تدريب - القسم الإعدادي
    # =========================================================================
    ws_prep = out_wb.create_sheet(title='برنامج تدريب - القسم الإعدادي')
    apply_rtl(ws_prep)
    ws_prep.merge_cells("A1:I1")
    t_cell = ws_prep.cell(1, 1, "برنامج محاضرات تخصص ( الضبع الاسود ) للقسم ( الإعدادي )")
    t_cell.font = title_font
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_prep.row_dimensions[1].height = 35

    render_term_table(ws_prep, start_row=3, term_label="الترم الأول", schedule=PREP_SCHEDULE,
                      border_cell=border_cell, header_font=header_font, subheader_font=subheader_font,
                      data_font=data_font, header_fill=header_fill, subheader_fill=subheader_fill,
                      midterm_after_day=4)

    # Column widths
    col_w = {1: 14, 2: 15, 3: 16, 4: 55, 5: 10, 6: 10, 7: 8, 8: 8, 9: 14}
    for col_idx, w in col_w.items():
        ws_prep.column_dimensions[get_column_letter(col_idx)].width = w

    # =========================================================================
    # 2. SHEET: برنامج تدريب - القسم المتوسط (Term 1 & Term 2)
    # =========================================================================
    ws_med = out_wb.create_sheet(title='برنامج تدريب - القسم المتوسط')
    apply_rtl(ws_med)
    ws_med.merge_cells("A1:I1")
    t_cell = ws_med.cell(1, 1, "برنامج محاضرات تخصص ( الضبع الاسود ) للقسم ( المتوسط )")
    t_cell.font = title_font
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_med.row_dimensions[1].height = 35

    # Term 1
    next_r = render_term_table(ws_med, start_row=3, term_label="الترم الأول", schedule=MED_T1_SCHEDULE,
                               border_cell=border_cell, header_font=header_font, subheader_font=subheader_font,
                               data_font=data_font, header_fill=header_fill, subheader_fill=subheader_fill,
                               midterm_after_day=6)

    # Term 2
    render_term_table(ws_med, start_row=next_r + 1, term_label="الترم الثاني", schedule=MED_T2_SCHEDULE,
                      border_cell=border_cell, header_font=header_font, subheader_font=subheader_font,
                      data_font=data_font, header_fill=header_fill, subheader_fill=subheader_fill,
                      midterm_after_day=7)

    for col_idx, w in col_w.items():
        ws_med.column_dimensions[get_column_letter(col_idx)].width = w

    # =========================================================================
    # 3. SHEET: برنامج تدريب - القسم النهائي (Term 1 & Term 2)
    # =========================================================================
    ws_fin = out_wb.create_sheet(title='برنامج تدريب - القسم النهائي')
    apply_rtl(ws_fin)
    ws_fin.merge_cells("A1:I1")
    t_cell = ws_fin.cell(1, 1, "برنامج محاضرات تخصص ( ايجلا ) للقسم ( النهائي )")
    t_cell.font = title_font
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_fin.row_dimensions[1].height = 35

    # Term 1
    next_r = render_term_table(ws_fin, start_row=3, term_label="الترم الأول", schedule=FINAL_T1_SCHEDULE,
                               border_cell=border_cell, header_font=header_font, subheader_font=subheader_font,
                               data_font=data_font, header_fill=header_fill, subheader_fill=subheader_fill,
                               midterm_after_day=6)

    # Term 2
    render_term_table(ws_fin, start_row=next_r + 1, term_label="الترم الثاني", schedule=FINAL_T2_SCHEDULE,
                      border_cell=border_cell, header_font=header_font, subheader_font=subheader_font,
                      data_font=data_font, header_fill=header_fill, subheader_fill=subheader_fill,
                      midterm_after_day=5)

    for col_idx, w in col_w.items():
        ws_fin.column_dimensions[get_column_letter(col_idx)].width = w

    print("Training program sheets with V3 clean layout created successfully.")

    # =========================================================================
    # 4. LOAD QUESTION POOLS
    # =========================================================================
    print("Loading question bank pools...")
    dhab_pool_file = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "معدة الضبع الاسود.xlsx")
    wb_dhab = openpyxl.load_workbook(dhab_pool_file, data_only=True)
    ws_dhab_src = wb_dhab.active

    raw_dhab_qs = []
    for r in range(2, ws_dhab_src.max_row + 1):
        row_vals = [ws_dhab_src.cell(r, c).value for c in range(1, 10)]
        if row_vals[1] is not None:
            raw_dhab_qs.append(row_vals)

    ejla_files = ["نهائي ضبع اسود.xlsx"]
    raw_fin_qs = []
    for ef in ejla_files:
        ef_path = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", ef)
        if os.path.exists(ef_path):
            wb_ef = openpyxl.load_workbook(ef_path, data_only=True)
            ws_ef = wb_ef.active
            for r in range(2, ws_ef.max_row + 1):
                row_vals = [ws_ef.cell(r, c).value for c in range(1, 10)]
                if row_vals[1] is not None:
                    raw_fin_qs.append(row_vals)

    def generate_bank_for_schedule(schedule, raw_pool):
        bank_data = []
        global_idx = 1
        pool_cursor = 0
        for it in schedule:
            q_count = it['q']
            for _ in range(q_count):
                src_q = copy(raw_pool[pool_cursor % len(raw_pool)])
                pool_cursor += 1
                # Format: [م, السؤال, النوع, الدرجة, خ1, خ2, خ3, خ4, الإجابة, الموضوع, المحاضرة, الصفحة]
                q_row = [
                    global_idx,
                    src_q[1], # question
                    src_q[2] if src_q[2] else "اختيار من متعدد",
                    src_q[3] if src_q[3] else 1,
                    src_q[4],
                    src_q[5],
                    src_q[6],
                    src_q[7],
                    src_q[8],
                    it['title'],
                    it['period'],
                    f"{it['p_from']} - {it['p_to']}"
                ]
                bank_data.append(q_row)
                global_idx += 1
        return bank_data

    # Generate Bank Sheets
    # 1. بنك القسم الإعدادي (700)
    prep_bank = generate_bank_for_schedule(PREP_SCHEDULE, raw_dhab_qs)
    ws_qb_prep = out_wb.create_sheet(title='بنك القسم الإعدادي')
    build_question_bank_sheet(ws_qb_prep, prep_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"Built بنك القسم الإعدادي ({len(prep_bank)} سؤال).")

    # 2. بنك المتوسط - ترم أول (1200)
    med_t1_bank = generate_bank_for_schedule(MED_T1_SCHEDULE, raw_dhab_qs)
    ws_qb_med_t1 = out_wb.create_sheet(title='بنك المتوسط - ترم أول')
    build_question_bank_sheet(ws_qb_med_t1, med_t1_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"Built بنك المتوسط - ترم أول ({len(med_t1_bank)} سؤال).")

    # 3. بنك المتوسط - ترم ثاني (1400)
    med_t2_bank = generate_bank_for_schedule(MED_T2_SCHEDULE, raw_dhab_qs)
    ws_qb_med_t2 = out_wb.create_sheet(title='بنك المتوسط - ترم ثاني')
    build_question_bank_sheet(ws_qb_med_t2, med_t2_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"Built بنك المتوسط - ترم ثاني ({len(med_t2_bank)} سؤال).")

    # 4. بنك النهائي - ترم أول (1200)
    fin_t1_bank = generate_bank_for_schedule(FINAL_T1_SCHEDULE, raw_fin_qs)
    ws_qb_fin_t1 = out_wb.create_sheet(title='بنك النهائي - ترم أول')
    build_question_bank_sheet(ws_qb_fin_t1, fin_t1_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"Built بنك النهائي - ترم أول ({len(fin_t1_bank)} سؤال).")

    # 5. بنك النهائي - ترم ثاني (1000)
    fin_t2_bank = generate_bank_for_schedule(FINAL_T2_SCHEDULE, raw_fin_qs)
    ws_qb_fin_t2 = out_wb.create_sheet(title='بنك النهائي - ترم ثاني')
    build_question_bank_sheet(ws_qb_fin_t2, fin_t2_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"Built بنك النهائي - ترم ثاني ({len(fin_t2_bank)} سؤال).")

    primary_file = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "برنامج تدريب تخصص الضبع الاسود.xlsx")
    updated_file = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "برنامج تدريب تخصص الضبع الاسود_النسخة_المحدثة_المعتمدة.xlsx")

    out_wb.save(updated_file)
    print(f"\n=======================================================")
    print(f"SUCCESS: Saved updated workbook to: {updated_file}")

    try:
        out_wb.save(primary_file)
        print(f"SUCCESS: Also updated original file: {primary_file}")
    except PermissionError:
        print(f"NOTE: {primary_file} is currently open in Excel. Cleanly updated '{os.path.basename(updated_file)}'.")

    print(f"Total Sheets in workbook ({len(out_wb.sheetnames)}): {out_wb.sheetnames}")
    print(f"=======================================================")

if __name__ == '__main__':
    main()
