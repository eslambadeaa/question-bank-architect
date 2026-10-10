import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'C:\Users\MaximuM-Tech\Downloads\برنامج تدريب تخصص الضبع الاسود + بنوك_محدث_LMS.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

def analyze_sheet(sheet_name):
    ws = wb[sheet_name]
    print(f"\n=================== {sheet_name} ===================")
    topics = []
    for r in range(5, ws.max_row + 1):
        lec = ws.cell(r, 3).value
        name = ws.cell(r, 4).value
        ps = ws.cell(r, 7).value
        pe = ws.cell(r, 8).value
        if lec in ['ف1', 'ف3']:
            topics.append({
                'idx': len(topics) + 1,
                'name': name.strip() if name else '',
                'ps': ps,
                'pe': pe
            })
    
    # Let's inspect each topic
    for t in topics:
        print(f"[{t['idx']:02d}] من {t['ps']} إلى {t['pe']} : {t['name']}")

analyze_sheet('برنامج تدريب - القسم المتوسط')
analyze_sheet('برنامج تدريب - القسم النهائي')
