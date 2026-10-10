import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

for tag, fname in [('متوسط', 'بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'), ('نهائي', 'بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx')]:
    wb = openpyxl.load_workbook(r'C:\Users\MaximuM-Tech\Downloads\\' + fname, data_only=True)
    ws = wb['Questions']
    print(f'=== Sample duplicates in {tag} ===')
    seen = {}
    dups = []
    for r in range(2, ws.max_row + 1):
        txt = str(ws.cell(r, 2).value).strip() if ws.cell(r, 2).value else ''
        les = ws.cell(r, 10).value
        if txt in seen:
            dups.append((r, les, txt, seen[txt]))
        else:
            seen[txt] = (r, les)
    print(f'Found {len(dups)} duplicate instances. Sample first 5:')
    for d in dups[:5]:
        print(f'  Row {d[0]} in \"{d[1]}\" is identical to Row {d[3][0]} in \"{d[3][1]}\"')
        print(f'    Text: \"{d[2][:70]}...\"')
