import win32com.client
import openpyxl
import sys
import os
import re

sys.stdout.reconfigure(encoding='utf-8')

# Load syllabus
src_path = os.path.join(r'C:\Users\MaximuM-Tech\Downloads\بنوك', 'معدة الضبع الاسود.xlsx')
src_wb = openpyxl.load_workbook(src_path, data_only=True)
ws_syl = src_wb['الجدول الزمني وفهرس الدروس']

lessons = []
for r in range(3, ws_syl.max_row + 1):
    idx = ws_syl.cell(r, 1).value
    if idx is not None and isinstance(idx, (int, float)):
        lessons.append({
            'id': int(idx),
            'title': str(ws_syl.cell(r, 2).value).strip(),
            'ref': str(ws_syl.cell(r, 6).value).strip()
        })

print(f"Total lessons to match: {len(lessons)}")

# Open Word and build paragraph page map for both files
word = win32com.client.Dispatch('Word.Application')
word.Visible = False

doc_maps = {}
for fname in ['مرجع الضبع الاسود.docx', 'مرجع الايجلا.docx']:
    fpath = os.path.join(r'C:\Users\MaximuM-Tech\Downloads', fname)
    doc = word.Documents.Open(fpath, ReadOnly=True)
    p_map = []
    for i in range(1, doc.Paragraphs.Count + 1):
        p = doc.Paragraphs(i)
        t = p.Range.Text.strip()
        if t:
            pg = p.Range.Information(3) # wdActiveEndPageNumber
            p_map.append((pg, t))
    doc.Close(False)
    doc_maps[fname] = p_map
    print(f"Indexed {len(p_map)} paragraphs for {fname}.")

try:
    word.Quit()
except Exception:
    pass

def find_page_for_lesson(lesson_title, p_map):
    # 1. Exact substring
    for pg, t in p_map:
        if lesson_title in t or t in lesson_title:
            return pg
            
    # 2. Token overlap
    words = [w for w in re.split(r'\s+', lesson_title) if len(w) > 2 and w not in ['المعدة', 'العام', 'فكرة', 'عامة', 'الخواص', 'الفنية', 'طريقة', 'عمل']]
    best_score = 0
    best_page = 1
    for pg, t in p_map:
        score = sum(1 for w in words if w in t)
        if score > best_score:
            best_score = score
            best_page = pg
    return best_page

# Match Preparatory (1..28)
p_map_dhab = doc_maps['مرجع الضبع الاسود.docx']
print("\n--- PREPARATORY LESSONS (مرجع الضبع الاسود) ---")
prep_matches = []
for l in lessons[:28]:
    pg = find_page_for_lesson(l['title'], p_map_dhab)
    prep_matches.append((l['id'], l['title'], pg))
    print(f"Lec {l['id']:2d}: Page {pg:2d} | {l['title']}")

# Match Medium (29..84)
print("\n--- MEDIUM LESSONS (مرجع الضبع الاسود) ---")
med_matches = []
for l in lessons[28:84]:
    pg = find_page_for_lesson(l['title'], p_map_dhab)
    med_matches.append((l['id'] - 28, l['title'], pg))
    print(f"Lec {l['id']-28:2d}: Page {pg:2d} | {l['title']}")

# Match Final (85..124)
p_map_ejla = doc_maps['مرجع الايجلا.docx']
print("\n--- FINAL LESSONS (مرجع الايجلا) ---")
fin_matches = []
for l in lessons[84:124]:
    pg = find_page_for_lesson(l['title'], p_map_ejla)
    fin_matches.append((l['id'] - 84, l['title'], pg))
    print(f"Lec {l['id']-84:2d}: Page {pg:2d} | {l['title']}")
