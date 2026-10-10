import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Verify questions sources
src_main = r'C:\Users\MaximuM-Tech\Downloads\بنوك\معدة الضبع الاسود.xlsx'
wb_main = openpyxl.load_workbook(src_main, data_only=True)

prep_qs = []
ws_prep = wb_main['بنك القسم الإعدادي']
for r in range(2, ws_prep.max_row + 1):
    vals = [ws_prep.cell(r, c).value for c in range(1, 12)]
    if any(vals):
        prep_qs.append(vals)

med_qs = []
ws_med = wb_main['بنك القسم المتوسط']
for r in range(2, ws_med.max_row + 1):
    vals = [ws_med.cell(r, c).value for c in range(1, 12)]
    if any(vals):
        med_qs.append(vals)

fin_qs = []
ws_fin = wb_main['بنك القسم النهائي']
for r in range(2, ws_fin.max_row + 1):
    vals = [ws_fin.cell(r, c).value for c in range(1, 12)]
    if any(vals):
        fin_qs.append(vals)

print(f"Loaded from main: Prep={len(prep_qs)}, Med={len(med_qs)}, Fin={len(fin_qs)}")

# Load from Eagla files
eagla_pool = []
p_bank = r'C:\Users\MaximuM-Tech\Downloads\بنك'
for fname in ['Eagla_Bank_1_Technical_200.xlsx', 'Eagla_Bank_2_Tactics_200.xlsx', 'Eagla_Bank_3_Firing_200.xlsx', 'Eagla_Bank_4_Fire_200.xlsx']:
    fpath = os.path.join(p_bank, fname)
    if os.path.exists(fpath):
        wb_e = openpyxl.load_workbook(fpath, data_only=True)
        ws_e = wb_e['Examples']
        for r in range(2, ws_e.max_row + 1):
            q_txt = ws_e.cell(r, 2).value
            if q_txt:
                vals = [
                    len(eagla_pool) + 1,
                    q_txt,
                    ws_e.cell(r, 3).value or "",
                    ws_e.cell(r, 4).value or "",
                    ws_e.cell(r, 5).value or "",
                    ws_e.cell(r, 6).value or "",
                    ws_e.cell(r, 7).value or "",
                    ws_e.cell(r, 8).value or "",
                    ws_e.cell(r, 9).value or "متوسط",
                    ws_e.cell(r, 10).value or "النظام الصاروخي إيجلا",
                    ws_e.cell(r, 11).value or "اختيار من متعدد"
                ]
                eagla_pool.append(vals)

print(f"Loaded from Eagla files: {len(eagla_pool)} questions.")
