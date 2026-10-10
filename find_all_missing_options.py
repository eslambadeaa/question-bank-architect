import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

def audit_file(f_path):
    print(f"\n==========================================")
    print(f"Auditing: {f_path.split('\\')[-1]}")
    print(f"==========================================")
    wb = openpyxl.load_workbook(f_path)
    ws_q = wb['Questions']
    
    missing_list = []
    for r in range(2, ws_q.max_row + 1):
        q_type = str(ws_q.cell(r, 1).value).strip() if ws_q.cell(r, 1).value else ''
        if q_type.startswith('1'):
            optA = ws_q.cell(r, 6).value
            optB = ws_q.cell(r, 7).value
            optC = ws_q.cell(r, 8).value
            optD = ws_q.cell(r, 9).value
            if not optA or not optB or not optC or not optD:
                missing_list.append((r, ws_q.cell(r, 2).value, ws_q.cell(r, 5).value, optA, optB, optC, optD, ws_q.cell(r, 10).value))
                
    print(f"Found {len(missing_list)} MCQ rows with missing options.")
    for item in missing_list:
        print(f"Row {item[0]} [{item[7]}]: {item[1]}")
        print(f"   Ans: {item[2]} | A: {item[3]} | B: {item[4]} | C: {item[5]} | D: {item[6]}")

audit_file(r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx')
audit_file(r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx')
