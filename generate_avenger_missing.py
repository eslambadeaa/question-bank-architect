import json
import sys
import openpyxl
import gemini_service
import time

sys.stdout.reconfigure(encoding='utf-8')

client = gemini_service.get_genai_client()

# Read full manual text
with open('e:/bank/manual_avenger_full_text.txt', 'r', encoding='utf-8') as f:
    manual_text = f.read()

# Load current bank
src_path = r'C:\Users\MaximuM-Tech\Downloads\avenger_bank_from_drive.xlsx'
wb = openpyxl.load_workbook(src_path, data_only=True)
ws = wb['Questions']

lessons = []
for r in range(2, 692):
    les = ws.cell(r, 10).value
    if les not in lessons:
        lessons.append(les)

cache_file = 'e:/bank/avenger_missing_questions_cache.json'
try:
    with open(cache_file, 'r', encoding='utf-8') as f:
        missing_cache = json.load(f)
except Exception:
    missing_cache = {}

print(f'Starting generation of missing questions for {len(lessons)} lessons...')

for idx, les in enumerate(lessons, 1):
    q_rows = [r for r in range(2, 692) if ws.cell(r, 10).value == les]
    cur_mcq = [r for r in q_rows if 'متعدد' in str(ws.cell(r, 1).value)]
    cur_tf = [r for r in q_rows if 'صح' in str(ws.cell(r, 1).value)]
    req_mcq = 13 if idx % 2 == 1 else 12
    req_tf = 12 if idx % 2 == 1 else 13
    need_mcq = max(0, req_mcq - len(cur_mcq))
    need_tf = max(0, req_tf - len(cur_tf))
    
    if need_mcq == 0 and need_tf == 0:
        continue
        
    cache_key = str(idx)
    if cache_key in missing_cache and len(missing_cache[cache_key].get('mcq', [])) >= need_mcq and len(missing_cache[cache_key].get('tf', [])) >= need_tf:
        print(f'Lesson {idx} ({les}) already in cache.')
        continue

    print(f'Generating for L{idx:02d} ({les}): need {need_mcq} MCQs and {need_tf} TFs...')
    
    existing_texts = [ws.cell(r, 2).value for r in q_rows]
    
    prompt = f"""أنت خبير عسكري وأكاديمي متخصص في الدفاع الجوي المصري ومعدة الإطلاق أفنجر (Avenger).
المطلوب منك توليد أسئلة دقيقة واحترافية غير مكررة نهائياً عن الدرس التالي:
عنوان الدرس: "{les}"

الأسئلة المطلوبة بالتحديد:
- عدد أسئلة الاختيار من متعدد المطلوبة: {need_mcq}
- عدد أسئلة الصواب والخطأ المطلوبة: {need_tf}

الأسئلة الموجودة حالياً بالفعل في هذا الدرس (يُمنع منعاً باتاً تكرارها أو صياغة شيء مشابه لها):
{chr(10).join(f'- {t}' for t in existing_texts[:15])}

المرجع المعرفي للمعدة (استند إليه في استخراج معلومات حقيقية وأرقام ومصطلحات دقيقة):
{manual_text[:15000]}

قواعد الإخراج الصارمة (JSON فقط بدون أي مقدمات أو علامات إضافية):
أخرج كائن JSON بالهيكل التالي فقط:
{{
  "mcq": [
    {{
      "text": "نص السؤال الدقيق",
      "ans": "A أو B أو C أو D",
      "opt_a": "الخيار أ",
      "opt_b": "الخيار ب",
      "opt_c": "الخيار ج",
      "opt_d": "الخيار د"
    }}
  ],
  "tf": [
    {{
      "text": "نص العبارة بدقة متناهية",
      "ans": "True أو False"
    }}
  ]
}}
"""
    response_text = ""
    for attempt in range(3):
        try:
            resp = client.models.generate_content(
                model='gemini-2.5-flash',
                contents=prompt,
                config=dict(response_mime_type="application/json")
            )
            response_text = resp.text
            data = json.loads(response_text)
            missing_cache[cache_key] = data
            with open(cache_file, 'w', encoding='utf-8') as f:
                json.dump(missing_cache, f, ensure_ascii=False, indent=2)
            print(f'  Successfully generated and cached L{idx}.')
            break
        except Exception as e:
            print(f'  Attempt {attempt+1} error: {e}')
            time.sleep(2)

print('All missing questions generated!')
