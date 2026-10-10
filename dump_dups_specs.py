import openpyxl, json, sys

sys.stdout.reconfigure(encoding='utf-8')

def extract_dups_spec(filepath, out_json):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb['Questions']
    seen = {}
    dups = []
    all_existing_texts = set()
    for r in range(2, ws.max_row + 1):
        txt = str(ws.cell(r, 2).value).strip() if ws.cell(r, 2).value else ''
        all_existing_texts.add(txt)
        qtype = str(ws.cell(r, 1).value).strip() if ws.cell(r, 1).value else ''
        diff = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value else ''
        les = str(ws.cell(r, 10).value).strip() if ws.cell(r, 10).value else ''
        if txt not in seen:
            seen[txt] = r
        else:
            dups.append({
                'row': r,
                'qtype': qtype,
                'diff': diff,
                'lesson': les,
                'orig_row': seen[txt]
            })
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(dups, f, ensure_ascii=False, indent=2)
    print(f"{filepath}: {len(dups)} duplicates saved to {out_json}")

extract_dups_spec(r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx', r'e:\bank\med_dups_spec.json')
extract_dups_spec(r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx', r'e:\bank\fin_dups_spec.json')
