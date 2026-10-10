import openpyxl
import collections
import sys

sys.stdout.reconfigure(encoding='utf-8')

for fname, total_q in [('بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx', 1300), ('بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx', 1100)]:
    fpath = r'C:\Users\MaximuM-Tech\Downloads\\' + fname
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb['Questions']
    seen = {}
    rows_to_replace = collections.defaultdict(list)
    for r in range(2, ws.max_row + 1):
        txt = str(ws.cell(r, 2).value).strip() if ws.cell(r, 2).value else ''
        les = ws.cell(r, 10).value
        q_type = ws.cell(r, 1).value
        diff = ws.cell(r, 4).value
        if txt not in seen:
            seen[txt] = (r, les)
        else:
            rows_to_replace[les].append((r, q_type, diff, txt, seen[txt]))
            
    print('='*60)
    print(f'{fname}: {sum(len(v) for v in rows_to_replace.values())} rows to replace across {len(rows_to_replace)} lessons')
    for les, r_list in list(rows_to_replace.items()):
        mcq_c = sum(1 for x in r_list if 'متعدد' in str(x[1]))
        tf_c = sum(1 for x in r_list if 'صح' in str(x[1]))
        print(f'  - \"{les}\": {len(r_list)} سؤال (MCQ: {mcq_c}, TF: {tf_c})')
