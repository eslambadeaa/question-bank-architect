import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

for fname in [r'C:\Users\MaximuM-Tech\Downloads\تكتيك و استخدام.docx', r'C:\Users\MaximuM-Tech\Downloads\قواعد خدمة الميدان صار محمول 29-11.docx']:
    doc = docx.Document(fname)
    print(f"=== {fname} ===")
    print("Paragraphs:", len(doc.paragraphs))
    for p in doc.paragraphs[:15]:
        if p.text.strip():
            print("  ", p.text.strip())
