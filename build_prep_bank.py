import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment, Border, Side
import shutil
import sys

sys.stdout.reconfigure(encoding='utf-8')

template_path = r'C:\Users\MaximuM-Tech\Downloads\نموذج تسجيل الاسئلة (1).xlsx'
source_lms_path = r'C:\Users\MaximuM-Tech\Downloads\برنامج تدريب تخصص الضبع الاسود + بنوك_محدث_LMS.xlsx'
output_path = r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_الاعدادي_الضبع_الاسود_LMS.xlsx'

wb_source = openpyxl.load_workbook(source_lms_path, data_only=True)

# 4-lesson repeating pattern
CYCLE = [
    # L1 (13 MCQ, 12 TF)
    (
        ['1 - سهل']*3 + ['2 - متوسط']*3 + ['3 - صعب']*2 + ['4 - صعب جدا']*3 + ['5 - تفوق']*2,  # 13 MCQ
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*1 + ['5 - تفوق']*1   # 12 TF
    ),
    # L2 (12 MCQ, 13 TF)
    (
        ['1 - سهل']*3 + ['2 - متوسط']*3 + ['3 - صعب']*2 + ['4 - صعب جدا']*3 + ['5 - تفوق']*1,  # 12 MCQ
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*1 + ['5 - تفوق']*2   # 13 TF
    ),
    # L3 (13 MCQ, 12 TF)
    (
        ['1 - سهل']*4 + ['2 - متوسط']*5 + ['3 - صعب']*2 + ['4 - صعب جدا']*1 + ['5 - تفوق']*1,  # 13 MCQ
        ['1 - سهل']*4 + ['2 - متوسط']*3 + ['3 - صعب']*1 + ['4 - صعب جدا']*3 + ['5 - تفوق']*1   # 12 TF
    ),
    # L4 (12 MCQ, 13 TF)
    (
        ['1 - سهل']*5 + ['2 - متوسط']*4 + ['3 - صعب']*1 + ['4 - صعب جدا']*1 + ['5 - تفوق']*1,  # 12 MCQ
        ['1 - سهل']*3 + ['2 - متوسط']*4 + ['3 - صعب']*3 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1   # 13 TF
    )
]

# Specifically calibrate L13 (idx 12) and L14 (idx 13) to hit exact 30.00%, 30.00%, 10.00%
L13_PATTERN = (
    ['1 - سهل']*3 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*2 + ['5 - تفوق']*2,  # 13 MCQ
    ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*1 + ['5 - تفوق']*1   # 12 TF
)
L14_PATTERN = (
    ['1 - سهل']*4 + ['2 - متوسط']*3 + ['3 - صعب']*2 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1,  # 12 MCQ
    ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*1 + ['4 - صعب جدا']*3 + ['5 - تفوق']*1   # 13 TF
)

def convert_mcq_to_tf(mcq_item, make_true=True):
    ans_key = str(mcq_item['ans']).strip()
    correct_opt = str(mcq_item.get('opt' + ans_key, '')).strip()
    
    incorrect_key = 'A' if ans_key != 'A' else 'B'
    incorrect_opt = str(mcq_item.get('opt' + incorrect_key, '')).strip()
    
    q_stem = mcq_item['text'].strip()
    if q_stem.endswith('؟') or q_stem.endswith('?'):
        q_stem = q_stem[:-1].strip()
        
    starters = ['ما هو ', 'ما هي ', 'كم يبلغ ', 'أين يقع ', 'متى يتم ', 'هل يمكن ', 'ما وظيفة ', 'ما الغرض من ', 'كم عدد ']
    cleaned_stem = q_stem
    for s in starters:
        if s in cleaned_stem:
            cleaned_stem = cleaned_stem.replace(s, '')
            break
            
    if make_true:
        text = f"{cleaned_stem} هو: {correct_opt}." if not cleaned_stem.endswith(':') else f"{cleaned_stem} {correct_opt}."
        ans = "True"
    else:
        text = f"{cleaned_stem} هو: {incorrect_opt}." if not cleaned_stem.endswith(':') else f"{cleaned_stem} {incorrect_opt}."
        ans = "False"
        
    return {
        'text': text,
        'exp': mcq_item['exp'],
        'diff': mcq_item['diff'],
        'ans': ans,
        'optA': 'True',
        'optB': 'False',
        'optC': None,
        'optD': None,
        'lesson': mcq_item['lesson']
    }

def process_prep():
    print("=======================================================")
    print("Processing Section: القسم الإعدادي")
    print("Source sheet: بنك القسم الإعدادي")
    print(f"Output path: {output_path}")
    print("=======================================================")

    ws_src = wb_source['بنك القسم الإعدادي']
    cur_lesson = None
    cur_mcqs = []
    cur_tfs = []
    lessons_data = []

    for r in range(2, ws_src.max_row + 1):
        q_type = ws_src.cell(r, 1).value
        q_text = ws_src.cell(r, 2).value
        if not q_text:
            continue
        q_exp = ws_src.cell(r, 3).value
        q_diff = ws_src.cell(r, 4).value
        q_ans = ws_src.cell(r, 5).value
        q_optA = ws_src.cell(r, 6).value
        q_optB = ws_src.cell(r, 7).value
        q_optC = ws_src.cell(r, 8).value
        q_optD = ws_src.cell(r, 9).value
        q_lesson = ws_src.cell(r, 10).value
        
        if q_lesson != cur_lesson:
            if cur_lesson is not None:
                lessons_data.append((cur_lesson, cur_mcqs, cur_tfs))
            cur_lesson = q_lesson
            cur_mcqs = []
            cur_tfs = []
            
        item = {
            'text': q_text,
            'exp': q_exp,
            'diff': q_diff,
            'ans': q_ans,
            'optA': q_optA,
            'optB': q_optB,
            'optC': q_optC,
            'optD': q_optD,
            'lesson': q_lesson
        }
        if str(q_type).startswith('1') or 'اختيار' in str(q_type):
            cur_mcqs.append(item)
        else:
            cur_tfs.append(item)
            
    if cur_lesson is not None:
        lessons_data.append((cur_lesson, cur_mcqs, cur_tfs))

    print(f"Loaded {len(lessons_data)} lessons for القسم الإعدادي.")

    shutil.copy2(template_path, output_path)
    wb_out = openpyxl.load_workbook(output_path)

    # 2. Update Lookups sheet
    ws_lookups = wb_out['Lookups']
    lesson_names = [ld[0] for ld in lessons_data]
    
    for r in range(2, ws_lookups.max_row + 10):
        ws_lookups.cell(r, 4).value = None

    for idx, les_name in enumerate(lesson_names, start=2):
        ws_lookups.cell(idx, 4, les_name)

    max_les_row = len(lesson_names) + 1
    wb_out.defined_names['L_4'] = DefinedName('L_4', attr_text=f"Lookups!$D$2:$D${max_les_row}")

    # 3. Populate Questions sheet
    ws_q = wb_out['Questions']
    while ws_q.max_row > 1:
        ws_q.delete_rows(2)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    regular_font = Font(name='Calibri', size=11, bold=False)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)

    current_row = 2

    # Stats
    stat_types = {}
    stat_diffs = {}
    stat_diffs_mcq = {}
    stat_diffs_tf = {}

    for les_idx, (les_name, mcqs, tfs) in enumerate(lessons_data):
        if les_idx == 12:
            mcq_diffs, tf_diffs = L13_PATTERN
        elif les_idx == 13:
            mcq_diffs, tf_diffs = L14_PATTERN
        else:
            mcq_diffs, tf_diffs = CYCLE[les_idx % 4]

        target_mcq_count = len(mcq_diffs)
        target_tf_count = len(tf_diffs)

        lesson_mcqs = mcqs[:target_mcq_count]

        needed_extra_tf = target_tf_count - len(tfs)
        extra_tfs = []
        for extra_i in range(needed_extra_tf):
            donor_mcq = mcqs[target_mcq_count + extra_i]
            make_true = (extra_i % 2 == 0)
            extra_tfs.append(convert_mcq_to_tf(donor_mcq, make_true=make_true))

        lesson_tfs = tfs + extra_tfs
        assert len(lesson_mcqs) == target_mcq_count
        assert len(lesson_tfs) == target_tf_count
        assert len(lesson_mcqs) + len(lesson_tfs) == 25

        # Write MCQs
        for idx in range(target_mcq_count):
            item = lesson_mcqs[idx]
            diff = mcq_diffs[idx]
            q_type = '1 - اختيار من متعدد'
            
            ans = str(item['ans']).strip()
            if ans not in ['A', 'B', 'C', 'D']:
                if ans.startswith('الخيار أ') or ans == 'أ': ans = 'A'
                elif ans.startswith('الخيار ب') or ans == 'ب': ans = 'B'
                elif ans.startswith('الخيار ج') or ans == 'ج': ans = 'C'
                elif ans.startswith('الخيار د') or ans == 'د': ans = 'D'

            row_data = [
                q_type,
                item['text'],
                item['exp'],
                diff,
                ans,
                item['optA'],
                item['optB'],
                item['optC'],
                item['optD'],
                les_name
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws_q.cell(current_row, c_idx, val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = center_align if c_idx in [1, 4, 5] else right_align

            current_row += 1
            stat_types[q_type] = stat_types.get(q_type, 0) + 1
            stat_diffs[diff] = stat_diffs.get(diff, 0) + 1
            stat_diffs_mcq[diff] = stat_diffs_mcq.get(diff, 0) + 1

        # Write TFs
        for idx in range(target_tf_count):
            item = lesson_tfs[idx]
            diff = tf_diffs[idx]
            q_type = '2 - صح/خطأ'

            ans = str(item['ans']).strip()
            if ans in ['صواب', 'صح', 'True', 'true', 'A', 'الخيار أ (A)']:
                norm_ans = 'True'
            elif ans in ['خطأ', 'خطا', 'False', 'false', 'B', 'الخيار ب (B)']:
                norm_ans = 'False'
            else:
                norm_ans = 'True'

            row_data = [
                q_type,
                item['text'],
                item['exp'],
                diff,
                norm_ans,
                'True',
                'False',
                None,
                None,
                les_name
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws_q.cell(current_row, c_idx, val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = center_align if c_idx in [1, 4, 5, 6, 7] else right_align

            current_row += 1
            stat_types[q_type] = stat_types.get(q_type, 0) + 1
            stat_diffs[diff] = stat_diffs.get(diff, 0) + 1
            stat_diffs_tf[diff] = stat_diffs_tf.get(diff, 0) + 1

    max_q_row = current_row - 1
    ws_q.data_validations.dataValidation.clear()

    dv_qtype = DataValidation(type="list", formula1="L_1", allow_blank=True)
    ws_q.add_data_validation(dv_qtype)
    dv_qtype.add(f"A2:A{max_q_row + 500}")

    dv_diff = DataValidation(type="list", formula1="L_2", allow_blank=True)
    ws_q.add_data_validation(dv_diff)
    dv_diff.add(f"D2:D{max_q_row + 500}")

    dv_ans = DataValidation(type="list", formula1="L_3", allow_blank=True)
    ws_q.add_data_validation(dv_ans)
    dv_ans.add(f"E2:E{max_q_row + 500}")

    dv_les = DataValidation(type="list", formula1="L_4", allow_blank=True)
    ws_q.add_data_validation(dv_les)
    dv_les.add(f"J2:J{max_q_row + 500}")

    wb_out.save(output_path)
    print(f"Saved {output_path} with {max_q_row - 1} questions ({len(lessons_data)} lessons * 25 questions).")
    total_q = max_q_row - 1
    print(f"Question Types:")
    for k, v in stat_types.items():
        print(f"  {k}: {v} ({v/total_q*100:.2f}%)")
    print(f"Overall Difficulties:")
    for k in sorted(stat_diffs.keys()):
        v = stat_diffs[k]
        print(f"  {k}: {v} ({v/total_q*100:.2f}%)")

process_prep()
