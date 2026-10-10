import openpyxl
import json
import collections
import sys

sys.stdout.reconfigure(encoding='utf-8')

for tag, fname in [('المتوسط', 'بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'), ('النهائي', 'بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx')]:
    fpath = r'C:\Users\MaximuM-Tech\Downloads\\' + fname
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb['Questions']
    seen = {}
    dups_by_lesson = collections.defaultdict(list)
    for r in range(2, ws.max_row + 1):
        txt = str(ws.cell(r, 2).value).strip() if ws.cell(r, 2).value else ''
        les = ws.cell(r, 10).value
        q_type = ws.cell(r, 1).value
        diff = ws.cell(r, 4).value
        if txt not in seen:
            seen[txt] = (r, les)
        else:
            dups_by_lesson[les].append((r, q_type, diff, txt))
            
    print(f'=== {tag}: {len(dups_by_lesson)} lessons need replacements ===')
    for les, items in dups_by_lesson.items():
        mcq_c = sum(1 for x in items if 'متعدد' in str(x[1]))
        tf_c = sum(1 for x in items if 'صح' in str(x[1]))
        print(f'  L: \"{les}\" -> MCQ={mcq_c}, TF={tf_c}, Total={len(items)}')
