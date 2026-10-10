import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment, Border, Side
import shutil
import sys

sys.stdout.reconfigure(encoding='utf-8')

template_path = r'C:\Users\MaximuM-Tech\Downloads\نموذج تسجيل الاسئلة (1).xlsx'
source_lms_path = r'C:\Users\MaximuM-Tech\Downloads\برنامج تدريب تخصص الضبع الاسود + بنوك_محدث_LMS.xlsx'

wb_source = openpyxl.load_workbook(source_lms_path, data_only=True)

# Perfect 4-lesson repeating pattern
PATTERNS = [
    # L1 (idx % 4 == 0)
    (
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*3 + ['5 - تفوق']*2,
        ['1 - سهل']*3 + ['2 - متوسط']*3 + ['3 - صعب']*2 + ['4 - صعب جدا']*1 + ['5 - تفوق']*1
    ),
    # L2 (idx % 4 == 1)
    (
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*3 + ['4 - صعب جدا']*2 + ['5 - تفوق']*2,
        ['1 - سهل']*3 + ['2 - متوسط']*3 + ['3 - صعب']*1 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1
    ),
    # L3 (idx % 4 == 2)
    (
        ['1 - سهل']*5 + ['2 - متوسط']*5 + ['3 - صعب']*2 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1,
        ['1 - سهل']*3 + ['2 - متوسط']*3 + ['3 - صعب']*1 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1
    ),
    # L4 (idx % 4 == 3)
    (
        ['1 - سهل']*5 + ['2 - متوسط']*5 + ['3 - صعب']*2 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1,
        ['1 - سهل']*3 + ['2 - متوسط']*3 + ['3 - صعب']*2 + ['4 - صعب جدا']*1 + ['5 - تفوق']*1
    )
]

def process_section_25q_exact(sec_name, source_sheets, output_path):
    print(f"\n=======================================================")
    print(f"Processing Section (25 Q/lesson - Exact 100%): {sec_name}")
    print(f"Source sheets: {source_sheets}")
    print(f"Output path: {output_path}")
    print(f"=======================================================")

    # 1. Read all questions from source sheets, grouped by lesson
    lessons_data = []
    
    for s_name in source_sheets:
        ws_src = wb_source[s_name]
        cur_lesson = None
        cur_mcqs = []
        cur_tfs = []
        
        for r in range(2, ws_src.max_row + 1):
            q_type = ws_src.cell(r, 1).value
            q_text = ws_src.cell(r, 2).value
            if not q_text:
                continue
            q_exp = ws_src.cell(r, 3).value
            q_diff = ws_src.cell(r, 4).value
            q_ans = ws_src.cell(r, 5).value
            q_optA = ws_src.cell(r, 6).value
            q_optB = ws_src.cell(r, 7).value
            q_optC = ws_src.cell(r, 8).value
            q_optD = ws_src.cell(r, 9).value
            q_lesson = ws_src.cell(r, 10).value
            
            if q_lesson != cur_lesson:
                if cur_lesson is not None:
                    lessons_data.append((cur_lesson, cur_mcqs, cur_tfs))
                cur_lesson = q_lesson
                cur_mcqs = []
                cur_tfs = []
                
            item = {
                'text': q_text,
                'exp': q_exp,
                'diff': q_diff,
                'ans': q_ans,
                'optA': q_optA,
                'optB': q_optB,
                'optC': q_optC,
                'optD': q_optD,
                'lesson': q_lesson
            }
            if str(q_type).startswith('1') or 'اختيار' in str(q_type):
                cur_mcqs.append(item)
            else:
                cur_tfs.append(item)
                
        if cur_lesson is not None:
            lessons_data.append((cur_lesson, cur_mcqs, cur_tfs))
            
    print(f"Loaded {len(lessons_data)} lessons for {sec_name}.")

    # Copy template to target
    shutil.copy2(template_path, output_path)
    wb_out = openpyxl.load_workbook(output_path)

    # 2. Update Lookups sheet
    ws_lookups = wb_out['Lookups']
    lesson_names = [ld[0] for ld in lessons_data]
    
    # Clear existing lessons in Col D
    for r in range(2, ws_lookups.max_row + 10):
        ws_lookups.cell(r, 4).value = None

    # Populate unique lesson names in order
    for idx, les_name in enumerate(lesson_names, start=2):
        ws_lookups.cell(idx, 4, les_name)

    max_les_row = len(lesson_names) + 1
    wb_out.defined_names['L_4'] = DefinedName('L_4', attr_text=f"Lookups!$D$2:$D${max_les_row}")

    # 3. Populate Questions sheet
    ws_q = wb_out['Questions']
    while ws_q.max_row > 1:
        ws_q.delete_rows(2)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    regular_font = Font(name='Calibri', size=11, bold=False)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)

    current_row = 2

    # Stats
    stat_types = {}
    stat_diffs = {}
    stat_diffs_mcq = {}
    stat_diffs_tf = {}

    for les_idx, (les_name, mcqs, tfs) in enumerate(lessons_data):
        assert len(mcqs) == 15, f"Lesson {les_name} has {len(mcqs)} MCQs"
        assert len(tfs) == 10, f"Lesson {les_name} has {len(tfs)} TFs"

        mcq_diffs, tf_diffs = PATTERNS[les_idx % 4]

        # Write 15 MCQs
        for idx in range(15):
            item = mcqs[idx]
            diff = mcq_diffs[idx]
            q_type = '1 - اختيار من متعدد'
            
            ans = str(item['ans']).strip()
            if ans not in ['A', 'B', 'C', 'D']:
                if ans.startswith('الخيار أ') or ans == 'أ': ans = 'A'
                elif ans.startswith('الخيار ب') or ans == 'ب': ans = 'B'
                elif ans.startswith('الخيار ج') or ans == 'ج': ans = 'C'
                elif ans.startswith('الخيار د') or ans == 'د': ans = 'D'

            row_data = [
                q_type,
                item['text'],
                item['exp'],
                diff,
                ans,
                item['optA'],
                item['optB'],
                item['optC'],
                item['optD'],
                les_name
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws_q.cell(current_row, c_idx, val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = center_align if c_idx in [1, 4, 5] else right_align

            current_row += 1
            stat_types[q_type] = stat_types.get(q_type, 0) + 1
            stat_diffs[diff] = stat_diffs.get(diff, 0) + 1
            stat_diffs_mcq[diff] = stat_diffs_mcq.get(diff, 0) + 1

        # Write 10 TFs
        for idx in range(10):
            item = tfs[idx]
            diff = tf_diffs[idx]
            q_type = '2 - صح/خطأ'

            ans = str(item['ans']).strip()
            if ans in ['صواب', 'صح', 'True', 'true', 'A', 'الخيار أ (A)']:
                norm_ans = 'True'
            elif ans in ['خطأ', 'خطا', 'False', 'false', 'B', 'الخيار ب (B)']:
                norm_ans = 'False'
            else:
                norm_ans = 'True'

            row_data = [
                q_type,
                item['text'],
                item['exp'],
                diff,
                norm_ans,
                'True',
                'False',
                None,
                None,
                les_name
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws_q.cell(current_row, c_idx, val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = center_align if c_idx in [1, 4, 5, 6, 7] else right_align

            current_row += 1
            stat_types[q_type] = stat_types.get(q_type, 0) + 1
            stat_diffs[diff] = stat_diffs.get(diff, 0) + 1
            stat_diffs_tf[diff] = stat_diffs_tf.get(diff, 0) + 1

    max_q_row = current_row - 1
    ws_q.data_validations.dataValidation.clear()

    dv_qtype = DataValidation(type="list", formula1="L_1", allow_blank=True)
    ws_q.add_data_validation(dv_qtype)
    dv_qtype.add(f"A2:A{max_q_row + 500}")

    dv_diff = DataValidation(type="list", formula1="L_2", allow_blank=True)
    ws_q.add_data_validation(dv_diff)
    dv_diff.add(f"D2:D{max_q_row + 500}")

    dv_ans = DataValidation(type="list", formula1="L_3", allow_blank=True)
    ws_q.add_data_validation(dv_ans)
    dv_ans.add(f"E2:E{max_q_row + 500}")

    dv_les = DataValidation(type="list", formula1="L_4", allow_blank=True)
    ws_q.add_data_validation(dv_les)
    dv_les.add(f"J2:J{max_q_row + 500}")

    wb_out.save(output_path)
    print(f"Saved {output_path} with {max_q_row - 1} questions ({len(lessons_data)} lessons * 25 questions).")
    print(f"Question Types: {stat_types}")
    print(f"Overall Difficulties:")
    total_q = max_q_row - 1
    for k in sorted(stat_diffs.keys()):
        v = stat_diffs[k]
        print(f"  {k}: {v} ({v/total_q*100:.2f}%)")
    print(f"MCQ Difficulties:")
    total_mcq = stat_types['1 - اختيار من متعدد']
    for k in sorted(stat_diffs_mcq.keys()):
        v = stat_diffs_mcq[k]
        print(f"  {k}: {v} ({v/total_mcq*100:.2f}%)")
    print(f"TF Difficulties:")
    total_tf = stat_types['2 - صح/خطأ']
    for k in sorted(stat_diffs_tf.keys()):
        v = stat_diffs_tf[k]
        print(f"  {k}: {v} ({v/total_tf*100:.2f}%)")

# Execute generation
med_sheets = ['بنك المتوسط - ترم أول', 'بنك المتوسط - ترم ثاني']
med_out = r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'
process_section_25q_exact('القسم المتوسط', med_sheets, med_out)

fin_sheets = ['بنك النهائي - ترم أول', 'بنك النهائي - ترم ثاني']
fin_out = r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx'
process_section_25q_exact('القسم النهائي', fin_sheets, fin_out)
