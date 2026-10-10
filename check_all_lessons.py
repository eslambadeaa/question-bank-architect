import openpyxl
import sys
import docx
import re

sys.stdout.reconfigure(encoding='utf-8')

# 1. Load syllabus lessons
wb = openpyxl.load_workbook(r'C:\Users\MaximuM-Tech\Downloads\بنوك\معدة الضبع الاسود.xlsx', data_only=True)
ws = wb['الجدول الزمني وفهرس الدروس']

lessons = []
for r in range(3, ws.max_row + 1):
    idx = ws.cell(r, 1).value
    if idx is not None and isinstance(idx, (int, float)):
        title = str(ws.cell(r, 2).value or '').strip()
        ref = str(ws.cell(r, 6).value or '').strip()
        lessons.append({'id': int(idx), 'title': title, 'ref': ref})

print(f"Loaded {len(lessons)} lessons.")

# 2. Extract full text paragraphs with page numbers
def get_paras_with_pages(docx_path):
    doc = docx.Document(docx_path)
    cur_p = 1
    paras = []
    for i, p in enumerate(doc.paragraphs, 1):
        p_start = cur_p
        for r in p.runs:
            xml = r._r.xml
            if 'lastRenderedPageBreak' in xml or 'w:type="page"' in xml:
                cur_p += 1
        p_end = cur_p
        t = p.text.strip()
        if t:
            paras.append({'idx': i, 'start_page': p_start, 'end_page': p_end, 'text': t})
    return paras

dhab_paras = get_paras_with_pages(r'C:\Users\MaximuM-Tech\Downloads\مرجع الضبع الاسود.docx')
ejla_paras = get_paras_with_pages(r'C:\Users\MaximuM-Tech\Downloads\مرجع الايجلا.docx')

print(f"Dhab: {len(dhab_paras)} paragraphs, pages up to {dhab_paras[-1]['end_page']}")
print(f"Ejla: {len(ejla_paras)} paragraphs, pages up to {ejla_paras[-1]['end_page']}")

# Print sections:
# Lessons 1..28 (Prep) -> Dhab
# Lessons 29..84 (Med) -> Dhab
# Lessons 85..124 (Fin) -> Ejla

print("\n--- PREPARATORY (1..28) ---")
for l in lessons[:28]:
    print(f"{l['id']:2d}: {l['title']}")

print("\n--- MEDIUM (29..84) ---")
for l in lessons[28:84]:
    print(f"{l['id']:2d}: {l['title']}")

print("\n--- FINAL (85..124) ---")
for l in lessons[84:124]:
    print(f"{l['id']:2d}: {l['title']}")
