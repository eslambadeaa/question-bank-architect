import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

src_file = r'C:\Users\MaximuM-Tech\Downloads\برنامج تدريب تخصص الضبع الاسود + بنوك_محدث_LMS.xlsx'
med_file = r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'
fin_file = r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx'

wb_src = openpyxl.load_workbook(src_file, data_only=True)

def verify_bank(section_name, syllabus_sheet, bank_path, expected_lessons_count):
    print(f"\n=======================================================")
    print(f"VERIFYING: {section_name}")
    print(f"Bank File: {bank_path}")
    print(f"=======================================================")

    # 1. Extract lessons from approved Syllabus sheet
    ws_syl = wb_src[syllabus_sheet]
    syl_lessons = []
    for r in range(5, ws_syl.max_row + 1):
        lec = ws_syl.cell(r, 3).value
        name = ws_syl.cell(r, 4).value
        if lec in ['ف1', 'ف3'] and name:
            syl_lessons.append(name.strip())

    print(f"Syllabus ({syllabus_sheet}) lessons count: {len(syl_lessons)}")
    assert len(syl_lessons) == expected_lessons_count, f"Expected {expected_lessons_count}, got {len(syl_lessons)}"

    # 2. Inspect Bank file
    wb_bank = openpyxl.load_workbook(bank_path, data_only=True)
    ws_q = wb_bank['Questions']
    ws_l = wb_bank['Lookups']

    # 2.1 Check Lookups lessons order
    lookups_lessons = []
    for r in range(2, ws_l.max_row + 1):
        v = ws_l.cell(r, 4).value
        if v:
            lookups_lessons.append(v.strip())

    print(f"Lookups lessons count: {len(lookups_lessons)}")
    assert len(lookups_lessons) == expected_lessons_count

    # 2.2 Check Questions rows and lessons sequence
    q_lessons_sequence = []
    lesson_q_counts = {}
    type_counts = {}
    diff_counts = {}
    row_errors = []

    cur_les = None
    for r in range(2, ws_q.max_row + 1):
        q_type = ws_q.cell(r, 1).value
        q_text = ws_q.cell(r, 2).value
        q_diff = ws_q.cell(r, 4).value
        q_ans = ws_q.cell(r, 5).value
        q_optA = ws_q.cell(r, 6).value
        q_optB = ws_q.cell(r, 7).value
        q_les = ws_q.cell(r, 10).value

        if not q_text:
            continue

        q_les = q_les.strip() if q_les else ''
        if q_les != cur_les:
            q_lessons_sequence.append(q_les)
            cur_les = q_les

        lesson_q_counts[q_les] = lesson_q_counts.get(q_les, 0) + 1
        type_counts[q_type] = type_counts.get(q_type, 0) + 1
        diff_counts[q_diff] = diff_counts.get(q_diff, 0) + 1

        # Check answers
        if str(q_type).startswith('1'):
            if q_ans not in ['A', 'B', 'C', 'D']:
                row_errors.append(f"Row {r}: MCQ invalid ans {q_ans}")
        else:
            if q_ans not in ['True', 'False']:
                row_errors.append(f"Row {r}: TF invalid ans {q_ans}")

    print(f"Total questions in bank: {ws_q.max_row - 1}")
    assert ws_q.max_row - 1 == expected_lessons_count * 25
    assert len(q_lessons_sequence) == expected_lessons_count

    # Check 1-to-1 match with Syllabus
    mismatch = False
    for idx, (s_name, b_name, l_name) in enumerate(zip(syl_lessons, q_lessons_sequence, lookups_lessons), start=1):
        if s_name != b_name or s_name != l_name:
            print(f"MISMATCH at Lesson {idx}!")
            print(f"  Syllabus: {s_name}")
            print(f"  Bank:     {b_name}")
            print(f"  Lookups:  {l_name}")
            mismatch = True

    if not mismatch:
        print(">> 100% PERFECT 1-to-1 LESSON MATCH between Syllabus, Bank, and Lookups!")

    # Check question counts per lesson
    all_25 = all(c == 25 for c in lesson_q_counts.values())
    print(f">> Every lesson has exactly 25 questions: {all_25}")

    # Check Question Types
    tot = ws_q.max_row - 1
    print(f">> Question Types:")
    for k, v in type_counts.items():
        print(f"     {k}: {v} ({v/tot*100:.2f}%)")
    assert type_counts['1 - اختيار من متعدد'] == tot // 2
    assert type_counts['2 - صح/خطأ'] == tot // 2

    # Check Difficulties
    print(f">> Difficulties:")
    for k in sorted(diff_counts.keys()):
        v = diff_counts[k]
        print(f"     {k}: {v} ({v/tot*100:.2f}%)")

    # Check Data Validation
    dvs = ws_q.data_validations.dataValidation
    print(f">> Data Validations count: {len(dvs)}")
    for dv in dvs:
        print(f"     Formula: {dv.formula1}, Sqref: {dv.sqref}")

    print(f">> Row errors count: {len(row_errors)}")
    if row_errors:
        print(f"Sample errors: {row_errors[:5]}")
    else:
        print(">> ZERO ROW ERRORS FOUND!")

print("STARTING FULL AUDIT...")
verify_bank("القسم المتوسط", "برنامج تدريب - القسم المتوسط", med_file, 52)
verify_bank("القسم النهائي", "برنامج تدريب - القسم النهائي", fin_file, 44)
print("\nAUDIT FINISHED.")
