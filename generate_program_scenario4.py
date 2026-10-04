"""
Generate Scenario 4 Institutional Training Program Workbook
Specialty: الضبع الاسود
Source Question Bank: C:\\Users\\MaximuM-Tech\\Downloads\\بنوك\\معدة الضبع الاسود.xlsx
"""

import os
import sys
from copy import copy
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

ARABIC_DAYS = [
    "الأول", "الثاني", "الثالث", "الرابع", "الخامس", "السادس",
    "السابع", "الثامن", "التاسع", "العاشر", "الحادي عشر", "الثانى عشر",
    "الثالث عشر", "الرابع عشر", "الخامس عشر", "السادس عشر"
]

def create_thin_border():
    thin = Side(border_style='thin', color='A0A0A0')
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def apply_rtl(ws):
    try:
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = True
    except Exception:
        pass

def calc_page_range(idx, total_items, total_pages=26):
    start_p = 1 + int((idx - 1) * max(1, total_pages - 1) / total_items)
    end_p = max(start_p, 1 + int(idx * max(1, total_pages - 1) / total_items))
    return start_p, end_p

def main():
    src_path = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "معدة الضبع الاسود.xlsx")
    if not os.path.exists(src_path):
        print(f"Error: Source file not found at {src_path}")
        return

    src_wb = openpyxl.load_workbook(src_path, data_only=True)
    ws_syl = src_wb['الجدول الزمني وفهرس الدروس']

    all_items = []
    for r in range(3, ws_syl.max_row + 1):
        idx = ws_syl.cell(r, 1).value
        if idx is not None and isinstance(idx, (int, float)):
            all_items.append({
                'id': int(idx),
                'title': ws_syl.cell(r, 2).value,
                'hours': ws_syl.cell(r, 3).value or 2,
                'questions': ws_syl.cell(r, 5).value or 25,
                'reference': ws_syl.cell(r, 6).value or 'مرجع الضبع الاسود.docx'
            })

    print(f"Extracted {len(all_items)} syllabus items from source.")

    sections_config = [
        {
            'sheet_name': 'برنامج محاضرات - الإعدادي',
            'title_sec': 'الإعدادي',
            'items': all_items[:28],
            'total_pages': 33,
            'days_count': 7,
            'midterm_day': 4,
            'hours': 56.0,
            'questions': 700
        },
        {
            'sheet_name': 'برنامج محاضرات - المتوسط',
            'title_sec': 'المتوسط',
            'items': all_items[28:84],
            'total_pages': 33,
            'days_count': 14,
            'midterm_day': 7,
            'hours': 112.0,
            'questions': 1400
        },
        {
            'sheet_name': 'برنامج محاضرات - النهائي',
            'title_sec': 'النهائي',
            'items': all_items[84:124],
            'total_pages': 33,
            'days_count': 10,
            'midterm_day': 5,
            'hours': 80.0,
            'questions': 1000
        }
    ]

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active) # Remove default sheet

    border_cell = create_thin_border()
    font_title = Font(name='Times New Roman', size=22, bold=True)
    font_header = Font(name='Times New Roman', size=15, bold=True)
    font_data = Font(name='Times New Roman', size=13, bold=False)
    font_exam = Font(name='Times New Roman', size=15, bold=True)
    font_total = Font(name='Times New Roman', size=15, bold=True)

    header_fill = PatternFill(start_color='E8EEF5', end_color='E8EEF5', fill_type='solid')
    exam_fill = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
    total_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')

    for cfg in sections_config:
        ws = out_wb.create_sheet(title=cfg['sheet_name'])
        apply_rtl(ws)

        # Row 1: Title
        ws.merge_cells('A1:H1')
        t_cell = ws['A1']
        t_cell.value = f"برنامج محاضرات تخصص ( الضبع الاسود ) للقسم ( {cfg['title_sec']} )"
        t_cell.font = font_title
        t_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 38
        ws.row_dimensions[2].height = 10

        # Header Rows 3 and 4
        ws.merge_cells('A3:A4')
        ws.merge_cells('B3:B4')
        ws.merge_cells('C3:C4')
        ws.merge_cells('D3:D4')
        ws.merge_cells('E3:E4')
        ws.merge_cells('F3:G3')
        ws.merge_cells('H3:H4')

        headers_3 = {
            'A3': 'الترم',
            'B3': 'اليوم',
            'C3': 'المحاضرة',
            'D3': 'اسم الموضوع',
            'E3': 'عدد الساعات',
            'F3': 'الصفحة في المرجع الموحد',
            'H3': 'عدد الاسئلة'
        }
        for cell_id, text in headers_3.items():
            c = ws[cell_id]
            c.value = text
            c.font = font_header
            c.fill = header_fill
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = border_cell

        ws['F4'].value = 'من'
        ws['G4'].value = 'الي'
        for c_id in ['F4', 'G4']:
            c = ws[c_id]
            c.font = font_header
            c.fill = header_fill
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = border_cell

        for r in [3, 4]:
            ws.row_dimensions[r].height = 25
            for col in range(1, 9):
                ws.cell(r, col).border = border_cell
                if ws.cell(r, col).fill.fill_type is None:
                    ws.cell(r, col).fill = header_fill

        current_row = 5
        items = cfg['items']
        days_count = cfg['days_count']
        midterm_day = cfg['midterm_day']
        item_idx = 0
        term_start_row = current_row

        for day_num in range(1, days_count + 1):
            day_label = ARABIC_DAYS[day_num - 1] if day_num <= len(ARABIC_DAYS) else f'اليوم {day_num}'
            day_start_row = current_row

            # 4 periods per day
            for period_idx in range(1, 5):
                it = items[item_idx] if item_idx < len(items) else None
                p_label = f'ف{period_idx}'
                lec_title = it['title'] if it else f'موضوع تدريبي {item_idx + 1}'
                hrs = 2.0
                p_from, p_to = calc_page_range(item_idx + 1, len(items), cfg['total_pages'])
                q_cnt = 25

                ws.cell(current_row, 3, value=p_label).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(current_row, 4, value=lec_title).alignment = Alignment(horizontal='right', vertical='center')
                ws.cell(current_row, 5, value=hrs).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(current_row, 6, value=p_from).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(current_row, 7, value=p_to).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(current_row, 8, value=q_cnt).alignment = Alignment(horizontal='center', vertical='center')

                for col in range(1, 9):
                    c = ws.cell(current_row, col)
                    c.font = font_data
                    c.border = border_cell

                ws.row_dimensions[current_row].height = 24
                current_row += 1
                item_idx += 1

            # Merge day column (Col B) for the 4 rows
            ws.merge_cells(start_row=day_start_row, start_column=2, end_row=day_start_row + 3, end_column=2)
            day_cell = ws.cell(day_start_row, 2, value=day_label)
            day_cell.font = font_header
            day_cell.alignment = Alignment(horizontal='center', vertical='center')
            for r_b in range(day_start_row, day_start_row + 4):
                ws.cell(r_b, 2).border = border_cell

            # Insert Midterm Exam row after midterm_day
            if day_num == midterm_day:
                ws.row_dimensions[current_row].height = 26
                ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=8)
                m_cell = ws.cell(current_row, 2, value='امتحان منتصف الترم')
                m_cell.font = font_exam
                m_cell.alignment = Alignment(horizontal='center', vertical='center')
                for col in range(1, 9):
                    c = ws.cell(current_row, col)
                    c.fill = exam_fill
                    c.border = border_cell
                current_row += 1

        # Final Exam row
        ws.row_dimensions[current_row].height = 28
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=4)
        f_cell = ws.cell(current_row, 2, value='امتحان ختامى الترم')
        f_cell.font = font_total
        f_cell.alignment = Alignment(horizontal='center', vertical='center')

        # Hours sum
        tot_hrs = cfg['hours']
        c_hrs = ws.cell(current_row, 5, value=tot_hrs)
        c_hrs.font = font_total
        c_hrs.alignment = Alignment(horizontal='center', vertical='center')

        # Questions sum
        tot_qs = cfg['questions']
        c_qs = ws.cell(current_row, 8, value=tot_qs)
        c_qs.font = font_total
        c_qs.alignment = Alignment(horizontal='center', vertical='center')

        for col in range(1, 9):
            c = ws.cell(current_row, col)
            c.fill = total_fill
            c.border = border_cell

        # Merge Term column A
        ws.merge_cells(start_row=term_start_row, start_column=1, end_row=current_row, end_column=1)
        a_term = ws.cell(term_start_row, 1, value='الترم الأول')
        a_term.font = font_header
        a_term.alignment = Alignment(horizontal='center', vertical='center')
        for r_a in range(term_start_row, current_row + 1):
            ws.cell(r_a, 1).border = border_cell

        # Column widths
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 14
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 52
        ws.column_dimensions['E'].width = 16
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 12
        ws.column_dimensions['H'].width = 16

    # Copy Question Bank sheets from source into out_wb
    for sname in ['بنك القسم الإعدادي', 'بنك القسم المتوسط', 'بنك القسم النهائي']:
        if sname in src_wb.sheetnames:
            src_ws = src_wb[sname]
            dest_ws = out_wb.create_sheet(title=sname)
            apply_rtl(dest_ws)
            for row in src_ws.iter_rows():
                for cell in row:
                    dest_cell = dest_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                    if cell.has_style:
                        if cell.font:
                            dest_cell.font = copy(cell.font)
                        if cell.alignment:
                            dest_cell.alignment = copy(cell.alignment)
                        if cell.fill:
                            dest_cell.fill = copy(cell.fill)
                        if cell.border:
                            dest_cell.border = copy(cell.border)
            for col_letter, col_dim in src_ws.column_dimensions.items():
                dest_ws.column_dimensions[col_letter].width = col_dim.width

    out_file_path = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "برنامج تدريب تخصص الضبع الاسود.xlsx")
    out_wb.save(out_file_path)
    print(f"Successfully generated {out_file_path}!")
    print(f"Sheets in generated file: {out_wb.sheetnames}")

if __name__ == '__main__':
    main()
