import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

wb = openpyxl.Workbook()
ws_lookups = wb.active
ws_lookups.title = 'Lookups'

lookups_data = [
    ['أنواع الأسئلة', 'مستويات الصعوبة', 'نماذج الإجابة الصحيحة', 'دروس المقرر'],
    ['1 - اختيار من متعدد', '1 - سهل', 'A', None],
    ['2 - صح/خطأ', '2 - متوسط', 'B', None],
    ['3 - مقالي', '3 - صعب', 'C', None],
    ['4 - مطابقة', '4 - صعب جدا', 'D', None],
    ['5 - أكمل الفراغ', '5 - تفوق', 'A;B', None],
    ['6 - ترتيب', None, 'A;C', None],
    ['7 - متعدد الإجابات', None, 'True', None],
    [None, None, 'False', None],
]
for r_idx, row in enumerate(lookups_data, start=1):
    for c_idx, val in enumerate(row, start=1):
        ws_lookups.cell(r_idx, c_idx, val)

wb.defined_names['L_1'] = DefinedName('L_1', attr_text="Lookups!$A$2:$A$8")
wb.defined_names['L_2'] = DefinedName('L_2', attr_text="Lookups!$B$2:$B$6")
wb.defined_names['L_3'] = DefinedName('L_3', attr_text="Lookups!$C$2:$C$9")

ws_test = wb.create_sheet(title='Test')
ws_test.cell(1, 1, 'نوع السؤال')
ws_test.cell(2, 1, '1 - اختيار من متعدد')

dv_qtype = DataValidation(type='list', formula1='L_1', allow_blank=True)
dv_qtype.error = 'اختر قيمة من القائمة.'
dv_qtype.errorTitle = 'قيمة غير صحيحة'
dv_qtype.showErrorMessage = True
dv_qtype.add('A2:A100')
ws_test.add_data_validation(dv_qtype)

wb.save(r'e:\bank\test_dv.xlsx')
print('Successfully saved test_dv.xlsx')
