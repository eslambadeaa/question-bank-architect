# -*- coding: utf-8 -*-
"""
بناء برنامجي التدريب لتخصص المدفعية المضادة للطائرات (مد م ط)
طباعة المرجع الموحد لقائد السرية:
- برنامج القسم الإعدادي: 16 درساً (ص 1 إلى 47) - ترم واحد - 32 س نظري + 32 س عملي
- برنامج القسم النهائي: 48 درساً (ص 1 إلى 149) - ترمين (الرشاش ترم أول 1-47، والمدفع ترم ثاني 48-149)
يتضمن عمود أرقام الصفحات بالمرجع الموحد (من - إلى) وعمود الترم الدراسي
"""

import os
import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from artillery_lessons_data import ARTILLERY_PREP_LESSONS, ARTILLERY_FINAL_LESSONS

if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')

def style_artillery_program(ws, title, stage_name, lessons, is_multi_term=False):
    # الألوان المعيارية العسكرية الراقية
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    blue_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    light_blue = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    term_banner_fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")

    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    font_term = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
    font_data = Font(name="Calibri", size=10, bold=False, color="000000")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    # الصف 1: العنوان الرئيسي
    ws.merge_cells("A1:I1")
    ws["A1"] = f"برنامج تدريب تخصص المدفعية المضادة للطائرات ( مد م ط ) - {stage_name}"
    ws["A1"].font = font_title
    ws["A1"].fill = navy_fill
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 40

    # الصف 2 & 3: ترويسة الجدول مع أعمدة الصفحات (من - إلى) والترم
    ws.merge_cells("A2:A3")
    ws["A2"] = "م"
    ws.merge_cells("B2:B3")
    ws["B2"] = "اسم الموضوع"
    ws.merge_cells("C2:C3")
    ws["C2"] = "الترم الدراسي"
    ws.merge_cells("D2:E2")
    ws["D2"] = "عدد الساعات"
    ws["D3"] = "نظرى"
    ws["E3"] = "عملي"
    ws.merge_cells("F2:G2")
    ws["F2"] = "الصفحة في المرجع الموحد"
    ws["F3"] = "من"
    ws["G3"] = "إلى"
    ws.merge_cells("H2:H3")
    ws["H2"] = "يصلح للامتحان المنظومة"
    ws.merge_cells("I2:I3")
    ws["I2"] = "عدد الأسئلة"

    header_cells = ["A2", "B2", "C2", "D2", "D3", "E3", "F2", "F3", "G3", "H2", "I2"]
    for c_ref in header_cells:
        cell = ws[c_ref]
        cell.font = font_header
        cell.fill = blue_fill
        cell.alignment = align_center

    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24

    # الصف 4: شريط المرجع التوضيحي
    ws.merge_cells("A4:I4")
    if is_multi_term:
        ws["A4"] = f"موضوعات برنامج تدريب ( {stage_name} ) - طباعة المرجع الموحد لقائد السرية (ص 1 إلى ص 149)"
    else:
        ws["A4"] = f"موضوعات برنامج تدريب ( {stage_name} ) - الرشاش 14.5 مم الرباعي (ص 1 إلى ص 47)"
    ws["A4"].font = font_section
    ws["A4"].fill = light_blue
    ws["A4"].alignment = align_center
    ws.row_dimensions[4].height = 28

    curr_r = 5
    current_term = None

    for idx, (code, topic, p_from, p_to, term_name) in enumerate(lessons, start=1):
        # إضافة فاصل إذا تغير الترم في القسم النهائي
        if is_multi_term and term_name != current_term:
            current_term = term_name
            ws.merge_cells(f"A{curr_r}:I{curr_r}")
            term_title = "الترم الأول: الرشاش 14.5 مم الرباعي المضاد للطائرات (ص 1 إلى ص 47)" if term_name == "الترم الأول" else "الترم الثاني: المدفع 23 مم الثنائي المضاد للطائرات (ص 48 إلى ص 149)"
            ws.cell(curr_r, 1, term_title).font = font_term
            ws.cell(curr_r, 1).fill = term_banner_fill
            ws.cell(curr_r, 1).alignment = align_center
            ws.row_dimensions[curr_r].height = 28
            curr_r += 1

        ws.cell(curr_r, 1, idx).alignment = align_center
        ws.cell(curr_r, 2, topic).alignment = align_right
        ws.cell(curr_r, 3, term_name).alignment = align_center
        ws.cell(curr_r, 4, 2).alignment = align_center # ساعتان نظري
        ws.cell(curr_r, 5, 2).alignment = align_center # ساعتان عملي
        ws.cell(curr_r, 6, p_from).alignment = align_center # ص من
        ws.cell(curr_r, 7, p_to).alignment = align_center # ص إلى
        ws.cell(curr_r, 8, "يصلح لعدد ساعات النظرى فقط" if idx == 1 or (is_multi_term and idx == 17) else "").alignment = align_center
        ws.cell(curr_r, 9, 25).alignment = align_center # 25 سؤالاً

        fill = zebra_fill if idx % 2 == 0 else PatternFill(fill_type=None)
        for c in range(1, 10):
            cell = ws.cell(curr_r, c)
            cell.font = font_data
            cell.border = thin_border
            if fill.fill_type: cell.fill = fill

        ws.row_dimensions[curr_r].height = 22
        curr_r += 1

    # صف الإجمالي العام
    ws.merge_cells(f"A{curr_r}:C{curr_r}")
    ws.cell(curr_r, 1, f"إجمالي {stage_name}").alignment = align_center
    ws.cell(curr_r, 4, len(lessons) * 2).alignment = align_center # إجمالي ساعات نظري
    ws.cell(curr_r, 5, len(lessons) * 2).alignment = align_center # إجمالي ساعات عملي
    ws.cell(curr_r, 6, lessons[0][2]).alignment = align_center # بداية أول صفحة
    ws.cell(curr_r, 7, lessons[-1][3]).alignment = align_center # نهاية آخر صفحة
    ws.cell(curr_r, 8, "").alignment = align_center
    ws.cell(curr_r, 9, len(lessons) * 25).alignment = align_center # إجمالي الأسئلة

    for c in range(1, 10):
        cell = ws.cell(curr_r, c)
        cell.font = font_bold
        cell.fill = yellow_fill
        cell.border = thin_border

    ws.row_dimensions[curr_r].height = 26

    # عرض الأعمدة
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 66
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 11
    ws.column_dimensions['E'].width = 11
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 10
    ws.column_dimensions['H'].width = 25
    ws.column_dimensions['I'].width = 14

    ws.views.sheetView[0].rightToLeft = True

def build_all_programs():
    out_dir = r"e:\bank"
    
    # 1. برنامج القسم الإعدادي (16 درساً - ترم واحد - ص 1 إلى ص 47)
    wb_prep = openpyxl.Workbook()
    ws_prep = wb_prep.active
    ws_prep.title = "التخصص"
    style_artillery_program(
        ws_prep,
        "برنامج تدريب تخصص المدفعية المضادة للطائرات",
        "القسم الإعدادي",
        ARTILLERY_PREP_LESSONS,
        is_multi_term=False
    )
    prep_path = os.path.join(out_dir, "برنامج_تدريب_تخصص_مدفعية_القسم_الإعدادي.xlsx")
    wb_prep.save(prep_path)
    print(f"Created Prep Program: {prep_path}")

    # 2. برنامج القسم النهائي (48 درساً - ترمين - ص 1 إلى ص 149)
    wb_fin = openpyxl.Workbook()
    ws_fin = wb_fin.active
    ws_fin.title = "التخصص"
    style_artillery_program(
        ws_fin,
        "برنامج تدريب تخصص المدفعية المضادة للطائرات",
        "القسم النهائي",
        ARTILLERY_FINAL_LESSONS,
        is_multi_term=True
    )
    fin_path = os.path.join(out_dir, "برنامج_تدريب_تخصص_مدفعية_القسم_النهائي.xlsx")
    wb_fin.save(fin_path)
    print(f"Created Final Program: {fin_path}")

if __name__ == "__main__":
    build_all_programs()
