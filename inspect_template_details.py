import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_template = r'C:\Users\MaximuM-Tech\Downloads\نموذج تسجيل الاسئلة (1).xlsx'
wb = openpyxl.load_workbook(f_template)

for name in wb.sheetnames:
    ws = wb[name]
    print(f'=== Sheet: {name} ===')
    print('Dimensions:', ws.dimensions)
    print('Row 1 height:', ws.row_dimensions[1].height)
    for c in range(1, 11):
        cell = ws.cell(1, c)
        col_letter = openpyxl.utils.get_column_letter(c)
        col_w = ws.column_dimensions[col_letter].width
        color = cell.fill.fgColor.value if cell.fill and hasattr(cell.fill.fgColor, 'value') else None
        print(f'Col {c} ({col_letter}, w={col_w}): val="{cell.value}", font={cell.font.name}, sz={cell.font.size}, b={cell.font.bold}, color={color}, align={cell.alignment.horizontal}/{cell.alignment.vertical}')

print("\n--- Defined Names ---")
for k, v in wb.defined_names.items():
    print(k, v.value)

print("\n--- Data Validations on Questions ---")
ws_q = wb['Questions']
for dv in ws_q.data_validations.dataValidation:
    print(f"DV formula: {dv.formula1}, type: {dv.type}, sqref: {dv.sqref}")
