# -*- coding: utf-8 -*-
"""
بناء برنامجي التدريب لتخصص الشيلكا (ZSU-23-4)
- برنامج القسم المتوسط: 52 درساً (مذكرة مدافع ص 1 إلى 125 ترم أول + رادار ص 126 إلى 203 ترم ثاني)
- برنامج القسم النهائي: 44 درساً (رادار تخصصي ص 126 إلى 203 ترم أول + حاسب وأجهزة اتزان ص 204 إلى 250 ترم ثاني)
يتضمن عمود أرقام الصفحات بالمرجع (من - إلى) وعمود الترم الدراسي
مع الالتزام الكامل بمعايير واشتراطات إدارة الامتحانات
"""

import os
import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

from shilka_lessons_data import SHILKA_INTERMEDIATE_LESSONS, SHILKA_FINAL_LESSONS

if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')

def style_shilka_program(ws, stage_name, lessons, manual_desc):
    # الألوان المعيارية العسكرية الراقية
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    blue_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    light_blue = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    term_banner_fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")

    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    font_term = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
    font_data = Font(name="Calibri", size=10, bold=False, color="000000")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    # الصف 1: العنوان الرئيسي
    ws.merge_cells("A1:I1")
    ws["A1"] = f"برنامج تدريب تخصص ( الشيلكا ZSU-23-4 ) - {stage_name}"
    ws["A1"].font = font_title
    ws["A1"].fill = navy_fill
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 40

    # الصف 2 & 3: ترويسة الجدول مع أرقام الصفحات (من - إلى)
    ws.merge_cells("A2:A3")
    ws["A2"] = "م"
    ws.merge_cells("B2:B3")
    ws["B2"] = "اسم الموضوع"
    ws.merge_cells("C2:C3")
    ws["C2"] = "الترم الدراسي"
    ws.merge_cells("D2:E2")
    ws["D2"] = "عدد الساعات"
    ws["D3"] = "نظرى"
    ws["E3"] = "عملي"
    ws.merge_cells("F2:G2")
    ws["F2"] = "الصفحة في المرجع"
    ws["F3"] = "من"
    ws["G3"] = "إلى"
    ws.merge_cells("H2:H3")
    ws["H2"] = "يصلح للامتحان المنظومة"
    ws.merge_cells("I2:I3")
    ws["I2"] = "عدد الأسئلة"

    header_cells = ["A2", "B2", "C2", "D2", "D3", "E3", "F2", "F3", "G3", "H2", "I2"]
    for c_ref in header_cells:
        cell = ws[c_ref]
        cell.font = font_header
        cell.fill = blue_fill
        cell.alignment = align_center

    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24

    # الصف 4: شريط المرجع التوضيحي
    ws.merge_cells("A4:I4")
    ws["A4"] = f"موضوعات برنامج تدريب ( {stage_name} ) - {manual_desc}"
    ws["A4"].font = font_section
    ws["A4"].fill = light_blue
    ws["A4"].alignment = align_center
    ws.row_dimensions[4].height = 28

    curr_r = 5
    current_term = None

    for idx, (code, topic, p_from, p_to, term_name) in enumerate(lessons, start=1):
        # إضافة فاصل إذا تغير الترم
        if term_name != current_term:
            current_term = term_name
            ws.merge_cells(f"A{curr_r}:I{curr_r}")
            if "المتوسط" in stage_name:
                term_title = "الترم الأول: المدافع 23 مم الرباعية وتجهيزاتها (ص 1 إلى ص 125)" if term_name == "الترم الأول" else "الترم الثاني: محطة الرادار ومنظومات الكشف والتتبع (ص 126 إلى ص 203)"
            else:
                term_title = "الترم الأول: محطة الرادار التخصصية والحرب الإلكترونية (ص 126 إلى ص 203)" if term_name == "الترم الأول" else "الترم الثاني: جهاز الحاسب الإلكتروني وأجهزة الاتزان والمنظومة المتكاملة (ص 204 إلى ص 250)"

            ws.cell(curr_r, 1, term_title).font = font_term
            ws.cell(curr_r, 1).fill = term_banner_fill
            ws.cell(curr_r, 1).alignment = align_center
            ws.row_dimensions[curr_r].height = 28
            curr_r += 1

        ws.cell(curr_r, 1, idx).alignment = align_center
        ws.cell(curr_r, 2, topic).alignment = align_right
        ws.cell(curr_r, 3, term_name).alignment = align_center
        ws.cell(curr_r, 4, 2).alignment = align_center # ساعتان نظري
        ws.cell(curr_r, 5, 2).alignment = align_center # ساعتان عملي
        ws.cell(curr_r, 6, p_from).alignment = align_center # ص من
        ws.cell(curr_r, 7, p_to).alignment = align_center # ص إلى
        ws.cell(curr_r, 8, "يصلح لعدد ساعات النظرى فقط" if idx == 1 or (term_name == "الترم الثاني" and (idx == 33 or idx == 15)) else "").alignment = align_center
        ws.cell(curr_r, 9, 25).alignment = align_center # 25 سؤالاً

        fill = zebra_fill if idx % 2 == 0 else PatternFill(fill_type=None)
        for c in range(1, 10):
            cell = ws.cell(curr_r, c)
            cell.font = font_data
            cell.border = thin_border
            if fill.fill_type: cell.fill = fill

        ws.row_dimensions[curr_r].height = 22
        curr_r += 1

    # صف الإجمالي العام
    ws.merge_cells(f"A{curr_r}:C{curr_r}")
    ws.cell(curr_r, 1, f"إجمالي {stage_name}").alignment = align_center
    ws.cell(curr_r, 4, len(lessons) * 2).alignment = align_center # إجمالي ساعات نظري
    ws.cell(curr_r, 5, len(lessons) * 2).alignment = align_center # إجمالي ساعات عملي
    ws.cell(curr_r, 6, lessons[0][2]).alignment = align_center # أول صفحة
    ws.cell(curr_r, 7, lessons[-1][3]).alignment = align_center # آخر صفحة
    ws.cell(curr_r, 8, "").alignment = align_center
    ws.cell(curr_r, 9, len(lessons) * 25).alignment = align_center # إجمالي الأسئلة

    for c in range(1, 10):
        cell = ws.cell(curr_r, c)
        cell.font = font_bold
        cell.fill = yellow_fill
        cell.border = thin_border

    ws.row_dimensions[curr_r].height = 26

    # أبعاد الأعمدة
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 62
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 11
    ws.column_dimensions['E'].width = 11
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 14
    ws.column_dimensions['H'].width = 28
    ws.column_dimensions['I'].width = 14

    ws.views.sheetView[0].rightToLeft = True

def build_training_programs():
    out_dir = r"e:\bank"
    
    # 1. برنامج تدريب القسم المتوسط
    wb_med = openpyxl.Workbook()
    ws_med = wb_med.active
    ws_med.title = "التخصص"
    style_shilka_program(
        ws_med,
        "القسم المتوسط",
        SHILKA_INTERMEDIATE_LESSONS,
        manual_desc="مذكرة الشيلكا: الجزء الأول المدافع (ص 1-125) والجزء الثاني الرادار (ص 126-203)"
    )
    med_path = os.path.join(out_dir, "برنامج_تدريب_تخصص_شيلكا_القسم_المتوسط.xlsx")
    wb_med.save(med_path)
    print(f"Created Shilka Intermediate Program: {med_path}")

    # 2. برنامج تدريب القسم النهائي
    wb_fin = openpyxl.Workbook()
    ws_fin = wb_fin.active
    ws_fin.title = "التخصص"
    style_shilka_program(
        ws_fin,
        "القسم النهائي",
        SHILKA_FINAL_LESSONS,
        manual_desc="مذكرة الشيلكا: الجزء الثاني الرادار التخصصي (ص 126-203) والجزء الثالث الحاسب والاتزان (ص 204-250)"
    )
    fin_path = os.path.join(out_dir, "برنامج_تدريب_تخصص_شيلكا_القسم_النهائي.xlsx")
    wb_fin.save(fin_path)
    print(f"Created Shilka Final Program: {fin_path}")

if __name__ == "__main__":
    build_training_programs()
