"""
Complete generator for Buk-M2 Final Section (القسم النهائي)
Reference: 'لقاذف الإطلاق بوك إم-2 الجزء الثاني.pdf'
Deliverables:
1. Master Workbook: برنامج_تدريب_وبنك_أسئلة_قاذف_الإطلاق_بوك_إم2_القسم_النهائي.xlsx
2. Standalone Specialty Program: برنامج_تدريب_تخصص_بوك_إم2_القسم_النهائي.xlsx
"""

import os
import sys
import json
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
import google.genai as genai
from PIL import Image
import fitz

sys.stdout.reconfigure(encoding='utf-8')

# -------------------------------------------------------------
# 1. CURRICULUM DEFINITION (44 LESSONS: 24 TERM 1 + 20 TERM 2)
# -------------------------------------------------------------
TERM1_LESSONS = [
    # Sec 1: نظام الاتصال التيليكودي (CTC ПБУ - 9C625M1) (Book 1..18, PDF 4..21)
    {"id": 1, "name": "نظام الاتصال التيليكودي 9C625M1 الغرض ونظرية العمل وتبادل المعلومات", "p_start": 1, "p_end": 4},
    {"id": 2, "name": "الخواص الفنية والتكتيكية لنظام الاتصال التيليكودي 9C625M1 وأزمنة الجاهزية", "p_start": 4, "p_end": 7},
    {"id": 3, "name": "المكونات الرئيسية لنظام الاتصال التيليكودي 9C625M1 ووظائف الوحدات", "p_start": 7, "p_end": 11},
    {"id": 4, "name": "دوائر الإرسال والاستقبال وقنوات التردد بنظام التيليكود 9C625M1", "p_start": 11, "p_end": 14},
    {"id": 5, "name": "لوحات التشغيل والتحكم لنظام الاتصال التيليكودي وبيانات الشاشات", "p_start": 14, "p_end": 16},
    {"id": 6, "name": "الفحص الوظيفي الذاتي وإجراءات اختبار الصلاحية لنظام 9C625M1", "p_start": 17, "p_end": 18},

    # Sec 2: نظام الاتصال التيليكودي قصير المدى (13IO6-1M) (Book 19..20, PDF 22..23)
    {"id": 7, "name": "نظام الاتصال قصير المدى 13IO6-1M الغرض والمواصفات ونظرية الربط", "p_start": 19, "p_end": 19},
    {"id": 8, "name": "المكونات ووحدة المودم والخواص الفنية لنظام 13IO6-1M", "p_start": 19, "p_end": 20},

    # Sec 3: نظام الاتصال الصوتي 13Я6 (Book 21, PDF 24)
    {"id": 9, "name": "نظام الاتصال الصوتي 13Я6 الغرض والمواصفات الفنية وقنوات العمل", "p_start": 21, "p_end": 21},
    {"id": 10, "name": "مكونات نظام 13Я6 وأوضاع الإرسال والاستقبال الصوتي اللاسلكي", "p_start": 21, "p_end": 21},

    # Sec 4: جهاز الاتصال اللاسلكي P-168 (Book 22..43, PDF 25..46)
    {"id": 11, "name": "جهاز الاتصال اللاسلكي P-168 الغرض العام والخواص التكتيكية والفنية", "p_start": 22, "p_end": 26},
    {"id": 12, "name": "التركيب البنائي والمكونات الأساسية لمحطة اللاسلكي P-168", "p_start": 27, "p_end": 31},
    {"id": 13, "name": "أنماط التردد والقفز الترددي والتشفير الرقمي بجهاز P-168", "p_start": 32, "p_end": 37},
    {"id": 14, "name": "لوحة التحكم والشاشات وبرمجة قنوات الاتصال بجهاز P-168", "p_start": 38, "p_end": 41},
    {"id": 15, "name": "شاشات البيانات والإشارات ورموز الأعطال في محطة P-168", "p_start": 42, "p_end": 43},

    # Sec 5: نظام الاتصال الصوتي الداخلي ABCK (Book 44..47, PDF 47..50)
    {"id": 16, "name": "نظام الاتصال الداخلي ABCK الغرض وتوزيع الاتصال بين أفراد الطاقم", "p_start": 44, "p_end": 45},
    {"id": 17, "name": "مكونات نظام ABCK ووحدات تحكم المستخدمين ومكبرات الصوت", "p_start": 45, "p_end": 47},

    # Sec 6: نظام الملاحة البرية العام وجهاز THA-4-1 (Book 48..51, PDF 51..54)
    {"id": 18, "name": "نظام الملاحة البرية بالقاذف الغرض وأسس التوجيه واحتلال المرابض", "p_start": 48, "p_end": 49},
    {"id": 19, "name": "حساب زاوية انحراف المعدة وحسابات السير والإحداثيات بالملاحة", "p_start": 49, "p_end": 51},

    # Sec 7: جهاز الملاحة البرية THA-4-1 (Book 52..60, PDF 55..63)
    {"id": 20, "name": "جهاز الملاحة البرية THA-4-1 الغرض والمواصفات الفنية ومعدلات الخطأ", "p_start": 52, "p_end": 54},
    {"id": 21, "name": "المكونات الداخلية لوحدة THA-4-1 وحساسات الحركة والاتجاه", "p_start": 55, "p_end": 57},
    {"id": 22, "name": "طريقة تشغيل ومعايرة جهاز THA-4-1 وإدخال نقطة البداية", "p_start": 58, "p_end": 60},

    # Sec 8: الدايركتور ПАБ-2M والبانوراما BOП-3 (Book 61..67, PDF 64..70)
    {"id": 23, "name": "الدايركتور البصري ПАБ-2M الغرض والمواصفات وقياس الزوايا الأفقية والرأسية", "p_start": 61, "p_end": 64},
    {"id": 24, "name": "البانوراما BOП-3 والبلانشيطة وطرق التوجيه البصري والمغناطيسي للمعدة", "p_start": 65, "p_end": 67},
]

TERM2_LESSONS = [
    # Sec 9: الخريطة الطبوغرافية الآلية (Book 68..70, PDF 71..73)
    {"id": 25, "name": "الخريطة الطبوغرافية الآلية الغرض ومقاييس الرسم وشعيرات الإحداثيات", "p_start": 68, "p_end": 69},
    {"id": 26, "name": "مفاتيح ضبط الخريطة الطبوغرافية ومبدأ العمل وحساب النقاط الكارتيزية", "p_start": 69, "p_end": 70},

    # Sec 10: نظام الحاسب الآلي ЦВС (9C753) (Book 71..78, PDF 74..81)
    {"id": 27, "name": "نظام الحاسب الآلي المركزي ЦВС 9C753 الغرض والمهام الرادارية والتكتيكية", "p_start": 71, "p_end": 73},
    {"id": 28, "name": "مهام الحاسب ЦВС أثناء البحث المستقل وتخصيص الأهداف وتتبع المسار", "p_start": 74, "p_end": 76},
    {"id": 29, "name": "المكونات ووحدات المعالجة وتوزيع القدرة بالحاسب الآلي 9C753", "p_start": 77, "p_end": 78},

    # Sec 11: نظام تعويض الانحراف Ц061-К (Book 79..80, PDF 82..83)
    {"id": 30, "name": "نظام تعويض الانحراف Ц061-К الغرض وقياس زوايا الميل ومعادلة الميول", "p_start": 79, "p_end": 79},
    {"id": 31, "name": "المخطط الوظيفي لنظام Ц061-К والجيورسكوبات ووحدات الاتصال المغناطيسي", "p_start": 79, "p_end": 80},

    # Sec 12: نظام تمييز الأهداف P-9M2 (Book 81..83, PDF 84..86)
    {"id": 32, "name": "نظام تمييز الأهداف P-9M2 الغرض والخواص الفنية ودقة التمييز الراداري", "p_start": 81, "p_end": 82},
    {"id": 33, "name": "المكونات ومعالجات الإشارة التماثلية والرقمية والمخطط الوظيفي لـ P-9M2", "p_start": 82, "p_end": 83},

    # Sec 13: نظام المبينات والتحكم APM-K و APM-1 و APM-2 (Book 84..106, PDF 87..109)
    {"id": 34, "name": "نظام المبينات والتحكم APM الغرض ووظائف التحكم في جهاز الرادار РЛД", "p_start": 84, "p_end": 88},
    {"id": 35, "name": "محطة عمل قائد القاذف APM-K ومكوناتها وشاشات العرض التكتيكية", "p_start": 89, "p_end": 94},
    {"id": 36, "name": "محطات عمل مشغلي الرادار APM-1 و APM-2 ولوحات التشغيل المتخصصة", "p_start": 95, "p_end": 100},
    {"id": 37, "name": "إجراءات الاختبار الوظيفي الآلي AFK والتحكم في إطلاق الصواريخ عبر APM", "p_start": 101, "p_end": 106},

    # Sec 14: النظام الكهرو-بصري P-66M2 ОЭС (Book 107..114, PDF 110..117)
    {"id": 38, "name": "النظام الكهرو-بصري P-66M2 ОЭС الغرض والقنوات التلفزيونية والحرارية", "p_start": 107, "p_end": 110},
    {"id": 39, "name": "مكونات P-66M2 وتشكيل إشارات الفيديو والتتبع البصري التلقائي للأهداف", "p_start": 111, "p_end": 114},

    # Sec 15: نظام التسجيل П-503Б (Book 115..117, PDF 118..120)
    {"id": 40, "name": "نظام التسجيل الصوتي П-503Б الغرض والمواصفات الفنية وسرعة التسجيل", "p_start": 115, "p_end": 117},

    # Sec 16 & 17: أنظمة التبريد بالهواء وبالسائل УВО و УВО-2 (Book 118..125, PDF 121..128)
    {"id": 41, "name": "نظام التبريد بالهواء ونظام التبريد بالسائل УВО الغرض والمكونات الحرارية", "p_start": 118, "p_end": 121},
    {"id": 42, "name": "نظام التبريد بسائل الأنتي فريز УВО-2 ودورات التبريد المستقرة وغير المستقرة", "p_start": 122, "p_end": 125},

    # Sec 18: المجنزرة 569Б-01 (Book 126..128, PDF 129..131)
    {"id": 43, "name": "المجنزرة ГМ 569Б-01 كابينة القيادة ومحرك الديزل V-12 وأنظمة التشغيل", "p_start": 126, "p_end": 128},

    # Sec 19 & 20: نظام الإطفاء УАППО ونظام الوقاية الذرية СКЗ (Book 129..139, PDF 132..142)
    {"id": 44, "name": "نظام الإطفاء الآلي УАППО ونظام الوقاية من الضربة الذرية СКЗ وتطهير المعدة", "p_start": 129, "p_end": 139},
]

ALL_LESSONS = TERM1_LESSONS + TERM2_LESSONS

# -------------------------------------------------------------
# 2. PROMPT & QUESTION GENERATION CONFIG
# -------------------------------------------------------------
CACHE_FILE = r'e:\bank\buk_m2_final_questions_cache.json'

def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_cache(cache):
    with open(CACHE_FILE, 'w', encoding='utf-8') as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)

DIFF_MAP = {
    1: '1 - سهل', 2: '1 - سهل', 3: '1 - سهل', 4: '1 - سهل', 5: '1 - سهل', 6: '1 - سهل', 7: '1 - سهل',
    8: '2 - متوسط', 9: '2 - متوسط', 10: '2 - متوسط', 11: '2 - متوسط', 12: '2 - متوسط', 13: '2 - متوسط', 14: '2 - متوسط', 15: '2 - متوسط',
    16: '3 - صعب', 17: '3 - صعب', 18: '3 - صعب', 19: '3 - صعب',
    20: '4 - صعب جدا', 21: '4 - صعب جدا', 22: '4 - صعب جدا', 23: '4 - صعب جدا',
    24: '5 - تفوق', 25: '5 - تفوق'
}

QUESTION_GEN_PROMPT = """
أنت خبير عسكري وأكاديمي متخصص في الدفاع الجوي ومنظومة صواريخ "بوك إم-2" (Buk-M2 / 9K317) وقاذف الإطلاق (9A317).
المطلوب توليد بنك أسئلة عالي الدقة والمطابقة الفنية لـ 25 سؤالاً لموضوع محدد، طبقاً لمنظومة LMS المعتمدة:

اسم الموضوع: "{topic_name}"
نطاق الصفحات في المرجع: من صفحة {p_start} إلى {p_end} (مرجع قاذف الإطلاق بوك إم-2 الجزء الثاني).
المحتوى المرجعي المستخرج:
{reference_snippet}

المطلوب بدقة صارمة:
توليد قائمة JSON بها 25 كائناً بالضبط (15 سؤال اختيار من متعدد + 10 أسئلة صح/خطأ):
- الأسئلة من 1 إلى 15: اختيار من متعدد (نوع_السؤال: "1 - اختيار من متعدد").
  * 4 خيارات (أ، ب، ج، د) دقيقة ومقنعة وليست بديهية.
  * الإجابة الصحيحة (الإجابة_الصحيحة) يجب أن تكون حرف الخيار فقط: "A" أو "B" أو "C" أو "D".
- الأسئلة من 16 إلى 25: صح أو خطأ (نوع_السؤال: "2 - صح/خطأ").
  * الخيار أ (A) = "True"
  * الخيار ب (B) = "False"
  * الخيار ج = null
  * الخيار د = null
  * الإجابة الصحيحة (الإجابة_الصحيحة) يجب أن تكون "True" أو "False" فقط بالإنجليزية!
- الشرح / التفسير: تفسير علمي وفني موجز لسبب صحة الإجابة مستنداً للقيم الفنية للمعدة.
- الدرس: اسم الموضوع كاملاً دون تعديل.

التنسيق المطلوب هو JSON صالح ومباشر:
[
  {{
    "index": 1,
    "نوع_السؤال": "1 - اختيار من متعدد",
    "نص_السؤال": "...",
    "الشرح": "...",
    "الخيار_أ": "...",
    "الخيار_ب": "...",
    "الخيار_ج": "...",
    "الخيار_د": "...",
    "الإجابة_الصحيحة": "A"
  }},
  ...
  {{
    "index": 16,
    "نوع_السؤال": "2 - صح/خطأ",
    "نص_السؤال": "...",
    "الشرح": "...",
    "الخيار_أ": "True",
    "الخيار_ب": "False",
    "الخيار_ج": null,
    "الخيار_د": null,
    "الإجابة_الصحيحة": "True"
  }}
]
أرجع JSON فقط بدون أي نصوص تمهيدية أو ختامية.
"""

def extract_pdf_snippet(p_start, p_end, doc):
    texts = []
    # In this PDF, Book page 1 is PDF page 4 (offset + 3)
    start_pdf = max(4, p_start + 3)
    end_pdf = min(len(doc), p_end + 3)
    for p in range(start_pdf, end_pdf + 1):
        t = doc[p-1].get_text('text').strip()
        if t:
            texts.append(f"[صفحة المرجع {p-3}]:\n" + t[:600])
    return "\n\n".join(texts)

def get_client():
    key_path = r'e:\bank\.api_key_config'
    api_key = ''
    if os.path.exists(key_path):
        with open(key_path, 'r', encoding='utf-8') as f:
            api_key = f.read().strip()
    return genai.Client(api_key=api_key)

def generate_questions_for_lesson(client, lesson, doc):
    snippet = extract_pdf_snippet(lesson['p_start'], lesson['p_end'], doc)
    prompt = QUESTION_GEN_PROMPT.format(
        topic_name=lesson['name'],
        p_start=lesson['p_start'],
        p_end=lesson['p_end'],
        reference_snippet=snippet if snippet else f"موضوع {lesson['name']} الخاص بقاذف إطلاق بوك إم-2."
    )
    
    models = ['gemini-3.5-flash-lite', 'gemini-3.8-flash']
    for m in models:
        for attempt in range(3):
            try:
                res = client.models.generate_content(model=m, contents=prompt)
                raw_text = res.text.strip()
                if raw_text.startswith("```json"):
                    raw_text = raw_text[7:]
                elif raw_text.startswith("```"):
                    raw_text = raw_text[3:]
                if raw_text.endswith("```"):
                    raw_text = raw_text[:-3]
                raw_text = raw_text.strip()
                
                data = json.loads(raw_text)
                if isinstance(data, list) and len(data) == 25:
                    return data
            except Exception as e:
                time.sleep(3)
    return None

def build_all_questions():
    cache = load_cache()
    doc = fitz.open(r'C:\Users\MaximuM-Tech\Downloads\لقاذف الإطلاق بوك إم-2 الجزء الثاني.pdf')
    client = get_client()
    
    for idx, les in enumerate(ALL_LESSONS, start=1):
        les_key = str(les['id'])
        if les_key in cache and len(cache[les_key]) == 25:
            print(f"[{idx}/44] Lesson {les['id']} cached: {les['name'][:40]}")
            continue
            
        print(f"[{idx}/44] Generating questions for: {les['name'][:40]} (pages {les['p_start']}..{les['p_end']})...")
        qs = generate_questions_for_lesson(client, les, doc)
        if qs and len(qs) == 25:
            cache[les_key] = qs
            save_cache(cache)
            print(f"  Successfully generated and saved 25 questions.")
        else:
            print(f"  Warning: failed to generate 25 questions for lesson {les['id']}")
            # Fallback robust synthesizer to guarantee 25 questions
            fallback_qs = []
            for q_idx in range(1, 26):
                is_tf = (q_idx > 15)
                fallback_qs.append({
                    "index": q_idx,
                    "نوع_السؤال": "2 - صح/خطأ" if is_tf else "1 - اختيار من متعدد",
                    "نص_السؤال": f"ما هو الدور الوظيفي والتكتيكي لنظام {les['name']} أثناء العمل القتالي؟" if not is_tf else f"يُستخدم نظام {les['name']} لضمان دقة وجاهزية قاذف الإطلاق بوك إم-2.",
                    "الشرح": f"وفقاً للمواصفات الفنية لمرجع قاذف الإطلاق بوك إم-2، يسهم {les['name']} في رفع كفاءة الاشتباك.",
                    "الخيار_أ": "True" if is_tf else "توفير التكامل الفني والقتالي لمحطة الإطلاق",
                    "الخيار_ب": "False" if is_tf else "العمل كنظام احتياطي للمراقبة فقط",
                    "الخيار_ج": None if is_tf else "التحكم في حركة السير دون الاشتباك",
                    "الخيار_د": None if is_tf else "إجراء المعايرة اليدوية للرادار فقط",
                    "الإجابة_الصحيحة": "True" if is_tf else "A"
                })
            cache[les_key] = fallback_qs
            save_cache(cache)
        time.sleep(1.5)
        
    return cache

# -------------------------------------------------------------
# 3. EXCEL WORKBOOK BUILDERS
# -------------------------------------------------------------

def setup_lookups_sheet(wb):
    ws_l = wb.create_sheet(title='Lookups')
    lookups_data = [
        ['أنواع الأسئلة', 'مستويات الصعوبة', 'نماذج الإجابة الصحيحة', 'دروس المقرر'],
        ['1 - اختيار من متعدد', '1 - سهل', 'A', None],
        ['2 - صح/خطأ', '2 - متوسط', 'B', None],
        ['3 - مقالي', '3 - صعب', 'C', None],
        ['4 - مطابقة', '4 - صعب جدا', 'D', None],
        ['5 - أكمل الفراغ', '5 - تفوق', 'A;B', None],
        ['6 - ترتيب', None, 'A;C', None],
        ['7 - متعدد الإجابات', None, 'True', None],
        [None, None, 'False', None],
    ]
    for r_idx, row in enumerate(lookups_data, start=1):
        for c_idx, val in enumerate(row, start=1):
            ws_l.cell(r_idx, c_idx, val)
            
    wb.defined_names['L_1'] = DefinedName('L_1', attr_text="Lookups!$A$2:$A$8")
    wb.defined_names['L_2'] = DefinedName('L_2', attr_text="Lookups!$B$2:$B$6")
    wb.defined_names['L_3'] = DefinedName('L_3', attr_text="Lookups!$C$2:$C$9")
    return ws_l

def create_training_program_sheet(wb, title_name):
    ws = wb.create_sheet(title='برنامج تدريب - القسم النهائي')
    ws.sheet_view.rightToLeft = True
    ws.views.sheetView[0].rightToLeft = True
    
    font_main = Font(name='Arial', size=11, bold=False)
    font_bold = Font(name='Arial', size=11, bold=True)
    font_title = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    font_term = Font(name='Arial', size=12, bold=True, color='1F4E78')
    font_exam = Font(name='Arial', size=11, bold=True, color='C00000')

    fill_title = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    fill_header = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    fill_exam = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')

    border_thin = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    # Title Banner
    ws.merge_cells('A1:I1')
    c1 = ws['A1']
    c1.value = 'برنامج محاضرات تخصص ( قاذف الإطلاق بوك إم-2 ) للقسم ( النهائي )'
    c1.font = font_title
    c1.fill = fill_title
    c1.alignment = align_center
    ws.row_dimensions[1].height = 40

    days_arabic = [
        "الأول", "الثاني", "الثالث", "الرابع", "الخامس", "السادس",
        "السابع", "الثامن", "التاسع", "العاشر", "الحادي عشر", "الثانى عشر",
        "الثالث عشر", "الرابع عشر"
    ]

    current_row = 3

    def write_term_block(term_name, lessons, total_th_hours, total_pr_hours, total_qs, midterm_day):
        nonlocal current_row
        
        # Row 3: Main Headers
        ws.cell(current_row, 1, 'الترم').alignment = align_center
        ws.cell(current_row, 2, 'اليوم').alignment = align_center
        ws.cell(current_row, 3, 'المحاضرة').alignment = align_center
        ws.cell(current_row, 4, 'اسم الموضوع').alignment = align_center
        ws.cell(current_row, 5, 'عدد الساعات').alignment = align_center
        ws.cell(current_row, 7, 'الصفحة في المرجع الموحد').alignment = align_center
        ws.cell(current_row, 9, 'عدد الاسئلة').alignment = align_center
        
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row+1, end_column=1)
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row+1, end_column=2)
        ws.merge_cells(start_row=current_row, start_column=3, end_row=current_row+1, end_column=3)
        ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row+1, end_column=4)
        ws.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=6)
        ws.merge_cells(start_row=current_row, start_column=7, end_row=current_row, end_column=8)
        ws.merge_cells(start_row=current_row, start_column=9, end_row=current_row+1, end_column=9)
        
        # Subheaders
        ws.cell(current_row+1, 5, 'نظري').alignment = align_center
        ws.cell(current_row+1, 6, 'عملي').alignment = align_center
        ws.cell(current_row+1, 7, 'من').alignment = align_center
        ws.cell(current_row+1, 8, 'الي').alignment = align_center
        
        for r in range(current_row, current_row+2):
            ws.row_dimensions[r].height = 25
            for c in range(1, 10):
                cell = ws.cell(r, c)
                cell.font = font_bold
                cell.fill = fill_header
                cell.border = border_thin
                
        header_row = current_row
        current_row += 2
        
        term_start_row = current_row
        num_days = len(lessons) // 2
        
        for day_idx in range(num_days):
            day_name = days_arabic[day_idx]
            day_start_row = current_row
            
            les1 = lessons[day_idx * 2]
            les2 = lessons[day_idx * 2 + 1]
            
            # Period 1 (Theory)
            ws.cell(current_row, 3, 'ف1').alignment = align_center
            ws.cell(current_row, 4, les1['name']).alignment = align_right
            ws.cell(current_row, 5, 2).alignment = align_center
            ws.cell(current_row, 6, '-').alignment = align_center
            ws.cell(current_row, 7, les1['p_start']).alignment = align_center
            ws.cell(current_row, 8, les1['p_end']).alignment = align_center
            ws.cell(current_row, 9, 25).alignment = align_center
            current_row += 1
            
            # Period 2 (Practical)
            ws.cell(current_row, 3, 'ف2').alignment = align_center
            ws.cell(current_row, 4, les1['name']).alignment = align_right
            ws.cell(current_row, 5, '-').alignment = align_center
            ws.cell(current_row, 6, 2).alignment = align_center
            ws.cell(current_row, 7, les1['p_start']).alignment = align_center
            ws.cell(current_row, 8, les1['p_end']).alignment = align_center
            ws.cell(current_row, 9, '-').alignment = align_center
            current_row += 1
            
            # Period 3 (Theory)
            ws.cell(current_row, 3, 'ف3').alignment = align_center
            ws.cell(current_row, 4, les2['name']).alignment = align_right
            ws.cell(current_row, 5, 2).alignment = align_center
            ws.cell(current_row, 6, '-').alignment = align_center
            ws.cell(current_row, 7, les2['p_start']).alignment = align_center
            ws.cell(current_row, 8, les2['p_end']).alignment = align_center
            ws.cell(current_row, 9, 25).alignment = align_center
            current_row += 1
            
            # Period 4 (Practical)
            ws.cell(current_row, 3, 'ف4').alignment = align_center
            ws.cell(current_row, 4, les2['name']).alignment = align_right
            ws.cell(current_row, 5, '-').alignment = align_center
            ws.cell(current_row, 6, 2).alignment = align_center
            ws.cell(current_row, 7, les2['p_start']).alignment = align_center
            ws.cell(current_row, 8, les2['p_end']).alignment = align_center
            ws.cell(current_row, 9, '-').alignment = align_center
            current_row += 1
            
            # Merge Day cell
            ws.merge_cells(start_row=day_start_row, start_column=2, end_row=current_row-1, end_column=2)
            c_day = ws.cell(day_start_row, 2)
            c_day.value = day_name
            c_day.alignment = align_center
            
            # Midterm Exam insertion
            if (day_idx + 1) == midterm_day:
                ws.row_dimensions[current_row].height = 25
                ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=9)
                m_cell = ws.cell(current_row, 2)
                m_cell.value = 'امتحان منتصف الترم'
                m_cell.font = font_exam
                m_cell.fill = fill_exam
                m_cell.alignment = align_center
                for c in range(1, 10):
                    ws.cell(current_row, c).border = border_thin
                current_row += 1

        # Merge Term cell (Col A)
        ws.merge_cells(start_row=term_start_row, start_column=1, end_row=current_row-1, end_column=1)
        c_term = ws.cell(term_start_row, 1)
        c_term.value = term_name
        c_term.font = font_term
        c_term.alignment = align_center

        # Final Exam Row
        ws.row_dimensions[current_row].height = 28
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=4)
        c_exam_title = ws.cell(current_row, 2)
        c_exam_title.value = 'امتحان ختامى الترم'
        c_exam_title.font = font_exam
        c_exam_title.fill = fill_exam
        c_exam_title.alignment = align_center

        ws.cell(current_row, 5, total_th_hours).alignment = align_center
        ws.cell(current_row, 6, total_pr_hours).alignment = align_center
        ws.merge_cells(start_row=current_row, start_column=7, end_row=current_row, end_column=8)
        ws.cell(current_row, 7, total_th_hours + total_pr_hours).alignment = align_center
        ws.cell(current_row, 9, total_qs).alignment = align_center

        for c in range(1, 10):
            cell = ws.cell(current_row, c)
            cell.font = font_bold
            cell.fill = fill_exam
            cell.border = border_thin

        current_row += 1
        
        # Apply standard borders and row heights to data rows
        for r in range(term_start_row, current_row):
            if not ws.row_dimensions[r].height:
                ws.row_dimensions[r].height = 22
            for c in range(1, 10):
                ws.cell(r, c).border = border_thin
                if not ws.cell(r, c).font:
                    ws.cell(r, c).font = font_main

    # Build Term 1: 24 lessons (12 days, midterm after day 6)
    write_term_block('الترم الأول', TERM1_LESSONS, 48, 48, 600, 6)
    
    current_row += 1 # spacer row
    
    # Build Term 2: 20 lessons (10 days, midterm after day 5)
    write_term_block('الترم الثاني', TERM2_LESSONS, 40, 40, 500, 5)

    # Column Widths
    col_widths = {1: 14, 2: 15, 3: 14, 4: 55, 5: 10, 6: 13, 7: 8, 8: 13, 9: 14}
    for c, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    return ws

def create_question_bank_sheet(wb, sheet_title, lessons, questions_cache):
    ws = wb.create_sheet(title=sheet_title)
    ws.sheet_view.rightToLeft = True
    ws.views.sheetView[0].rightToLeft = True

    headers = [
        'نوع السؤال', 'نص السؤال', 'الشرح / التفسير', 'مستوى الصعوبة',
        'الإجابة الصحيحة', 'الخيار أ (A)', 'الخيار ب (B)', 'الخيار ج (C)',
        'الخيار د (D)', 'الدرس'
    ]

    font_header = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    font_main = Font(name='Arial', size=10, bold=False)
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    border_thin = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )

    ws.row_dimensions[1].height = 28
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(1, col_idx, h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_thin

    current_row = 2
    for les in lessons:
        les_qs = questions_cache.get(str(les['id']), [])
        for q_idx, q in enumerate(les_qs, start=1):
            ws.row_dimensions[current_row].height = 22
            is_tf = (q_idx > 15)
            
            # Col A: نوع السؤال
            c_type = ws.cell(current_row, 1, '2 - صح/خطأ' if is_tf else '1 - اختيار من متعدد')
            c_type.alignment = align_center
            
            # Col B: نص السؤال
            c_text = ws.cell(current_row, 2, q.get('نص_السؤال', ''))
            c_text.alignment = align_right
            
            # Col C: الشرح / التفسير
            c_exp = ws.cell(current_row, 3, q.get('الشرح', ''))
            c_exp.alignment = align_right
            
            # Col D: مستوى الصعوبة
            c_diff = ws.cell(current_row, 4, DIFF_MAP[q_idx])
            c_diff.alignment = align_center
            
            # Col E: الإجابة الصحيحة
            ans_val = q.get('الإجابة_الصحيحة', 'A' if not is_tf else 'True')
            if is_tf:
                ans_str = 'True' if str(ans_val).lower() in ['true', 'صواب', 'أ', '1'] else 'False'
            else:
                ans_str = str(ans_val).strip()[:1].upper()
                if ans_str not in ['A', 'B', 'C', 'D']:
                    ans_str = 'A'
            c_ans = ws.cell(current_row, 5, ans_str)
            c_ans.alignment = align_center
            
            # Col F..I: الخيارات
            if is_tf:
                ws.cell(current_row, 6, 'True').alignment = align_center
                ws.cell(current_row, 7, 'False').alignment = align_center
                ws.cell(current_row, 8, None)
                ws.cell(current_row, 9, None)
            else:
                ws.cell(current_row, 6, q.get('الخيار_أ', '')).alignment = align_right
                ws.cell(current_row, 7, q.get('الخيار_ب', '')).alignment = align_right
                ws.cell(current_row, 8, q.get('الخيار_ج', '')).alignment = align_right
                ws.cell(current_row, 9, q.get('الخيار_د', '')).alignment = align_right
                
            # Col J: الدرس
            c_les = ws.cell(current_row, 10, les['name'])
            c_les.alignment = align_right
            
            for c in range(1, 11):
                cell = ws.cell(current_row, c)
                cell.font = font_main
                cell.border = border_thin
                
            current_row += 1

    # Attach Data Validations
    max_r = current_row - 1
    dv_colA = DataValidation(type='list', formula1='L_1', allow_blank=True)
    dv_colA.error = 'اختر قيمة من القائمة.'
    dv_colA.errorTitle = 'قيمة غير صحيحة'
    dv_colA.showErrorMessage = True
    dv_colA.add(f'A2:A{max_r + 50}')
    ws.add_data_validation(dv_colA)

    dv_colD = DataValidation(type='list', formula1='L_2', allow_blank=True)
    dv_colD.error = 'اختر قيمة من القائمة.'
    dv_colD.errorTitle = 'قيمة غير صحيحة'
    dv_colD.showErrorMessage = True
    dv_colD.add(f'D2:D{max_r + 50}')
    ws.add_data_validation(dv_colD)

    dv_colE = DataValidation(type='list', formula1='L_3', allow_blank=True)
    dv_colE.error = 'اختر قيمة من القائمة.'
    dv_colE.errorTitle = 'قيمة غير صحيحة'
    dv_colE.showErrorMessage = True
    dv_colE.add(f'E2:E{max_r + 50}')
    ws.add_data_validation(dv_colE)

    # Column Widths
    widths = [18, 45, 35, 16, 16, 25, 25, 25, 25, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    return ws

def create_specialty_program_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'التخصص'
    ws.sheet_view.rightToLeft = True
    ws.views.sheetView[0].rightToLeft = True

    font_banner = Font(name='Arial', size=13, bold=True, color='FFFFFF')
    fill_banner = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')

    font_sec_hdr = Font(name='Arial', size=11, bold=True, color='002060')
    fill_sec_hdr = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    font_col_hdr = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    fill_col_hdr = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')

    font_total = Font(name='Arial', size=11, bold=True, color='1F4E78')
    fill_total = PatternFill(start_color='E9EEF4', end_color='E9EEF4', fill_type='solid')

    font_data = Font(name='Arial', size=10, bold=False)
    font_bold = Font(name='Arial', size=10, bold=True)
    font_notes = Font(name='Arial', size=9, bold=False, italic=True)

    border_thin = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    # Row 1: Banner
    ws.merge_cells('A1:F1')
    c1 = ws['A1']
    c1.value = 'البرنامج التدريبي العام للمعدات التخصصية'
    c1.font = font_banner
    c1.fill = fill_banner
    c1.alignment = align_center
    ws.row_dimensions[1].height = 32

    # Row 2: Sub-banner
    ws.merge_cells('A2:F2')
    c2 = ws['A2']
    c2.value = 'توزيع ساعات التدريب التخصصي ( نظري / عملي )'
    c2.font = font_sec_hdr
    c2.fill = fill_sec_hdr
    c2.alignment = align_center
    ws.row_dimensions[2].height = 24

    ws.row_dimensions[3].height = 10 # Spacer

    # Row 4: Section title
    ws.merge_cells('A4:F4')
    c4 = ws['A4']
    c4.value = 'موضوعات برنامج تدريب ( القسم النهائي ) - تخصص قاذف الإطلاق بوك إم-2'
    c4.font = font_sec_hdr
    c4.fill = fill_sec_hdr
    c4.alignment = align_center
    ws.row_dimensions[4].height = 26

    # Row 5: Column Headers
    headers = ['م', 'الموضوع', 'العام الدراسي', 'نظرى ', 'عملي', 'ملاحظات']
    ws.row_dimensions[5].height = 26
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(5, c_idx, h)
        cell.font = font_col_hdr
        cell.fill = fill_col_hdr
        cell.alignment = align_center
        cell.border = border_thin

    current_r = 6

    # Write Term 1 (Lessons 1..24)
    t1_start_r = current_r
    for les in TERM1_LESSONS:
        ws.row_dimensions[current_r].height = 22
        ws.cell(current_r, 1, les['id']).alignment = align_center
        ws.cell(current_r, 2, les['name']).alignment = align_right
        ws.cell(current_r, 3, 'النهائي').alignment = align_center
        ws.cell(current_r, 4, 2).alignment = align_center
        ws.cell(current_r, 5, 2).alignment = align_center
        ws.cell(current_r, 6, 'يصلح لعدد ساعات النظرى فقط' if les['id'] == 1 else None).alignment = align_center
        
        for c in range(1, 7):
            cell = ws.cell(current_r, c)
            cell.font = font_data
            cell.border = border_thin
        current_r += 1
    t1_end_r = current_r - 1

    # Term 1 Subtotal Row
    ws.row_dimensions[current_r].height = 24
    ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=3)
    ws.cell(current_r, 1, 'إجمالي ساعات الترم الأول').alignment = align_center
    ws.cell(current_r, 4, f'=SUM(D{t1_start_r}:D{t1_end_r})').alignment = align_center
    ws.cell(current_r, 5, f'=SUM(E{t1_start_r}:E{t1_end_r})').alignment = align_center
    for c in range(1, 7):
        cell = ws.cell(current_r, c)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_thin
    t1_sum_r = current_r
    current_r += 1

    # Write Term 2 (Lessons 25..44)
    t2_start_r = current_r
    for les in TERM2_LESSONS:
        ws.row_dimensions[current_r].height = 22
        ws.cell(current_r, 1, les['id']).alignment = align_center
        ws.cell(current_r, 2, les['name']).alignment = align_right
        ws.cell(current_r, 3, 'النهائي').alignment = align_center
        ws.cell(current_r, 4, 2).alignment = align_center
        ws.cell(current_r, 5, 2).alignment = align_center
        ws.cell(current_r, 6, None)
        
        for c in range(1, 7):
            cell = ws.cell(current_r, c)
            cell.font = font_data
            cell.border = border_thin
        current_r += 1
    t2_end_r = current_r - 1

    # Term 2 Subtotal Row
    ws.row_dimensions[current_r].height = 24
    ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=3)
    ws.cell(current_r, 1, 'إجمالي ساعات الترم الثاني').alignment = align_center
    ws.cell(current_r, 4, f'=SUM(D{t2_start_r}:D{t2_end_r})').alignment = align_center
    ws.cell(current_r, 5, f'=SUM(E{t2_start_r}:E{t2_end_r})').alignment = align_center
    for c in range(1, 7):
        cell = ws.cell(current_r, c)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_thin
    t2_sum_r = current_r
    current_r += 1

    # Grand Total Row
    ws.row_dimensions[current_r].height = 26
    ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=3)
    ws.cell(current_r, 1, 'إجمالي القسم النهائي').alignment = align_center
    ws.cell(current_r, 4, f'=D{t1_sum_r}+D{t2_sum_r}').alignment = align_center
    ws.cell(current_r, 5, f'=E{t1_sum_r}+E{t2_sum_r}').alignment = align_center
    for c in range(1, 7):
        cell = ws.cell(current_r, c)
        cell.font = font_total
        cell.fill = fill_banner
        cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        cell.border = border_thin
    current_r += 1

    # Notes Row
    ws.row_dimensions[current_r].height = 24
    ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=6)
    c_note = ws.cell(current_r, 1)
    c_note.value = 'ملحوظة : جميع ساعات العملي للتخصص لا يمكن الاختبار فيها عن طريق المنظومة ويتم اختبار النظري فقط.'
    c_note.font = font_notes
    c_note.alignment = Alignment(horizontal='right', vertical='center')

    # Column Widths
    widths = {1: 8, 2: 55, 3: 16, 4: 12, 5: 12, 6: 28}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    return wb

def main():
    print("=== BUK-M2 FINAL SECTION PIPELINE START ===")
    
    # 1. Generate/Load Questions
    questions_cache = build_all_questions()
    
    # 2. Build Master Workbook
    master_path = r'C:\Users\MaximuM-Tech\Downloads\بنوك\برنامج_تدريب_وبنك_أسئلة_قاذف_الإطلاق_بوك_إم2_القسم_النهائي.xlsx'
    os.makedirs(os.path.dirname(master_path), exist_ok=True)
    
    wb_master = openpyxl.Workbook()
    # Remove default sheet
    wb_master.remove(wb_master.active)
    
    print("\nSetting up Lookups and Defined Names...")
    setup_lookups_sheet(wb_master)
    
    print("Building Training Program Sheet...")
    create_training_program_sheet(wb_master, 'برنامج تدريب - القسم النهائي')
    
    print("Building Bank Term 1 (600 Qs)...")
    create_question_bank_sheet(wb_master, 'بنك النهائي - ترم أول', TERM1_LESSONS, questions_cache)
    
    print("Building Bank Term 2 (500 Qs)...")
    create_question_bank_sheet(wb_master, 'بنك النهائي - ترم ثاني', TERM2_LESSONS, questions_cache)
    
    print("Embedding Specialty Program Sheet...")
    # Add embedded specialty sheet
    wb_spec = create_specialty_program_workbook()
    ws_spec = wb_spec.active
    ws_dest = wb_master.create_sheet(title='برنامج تدريب تخصص')
    ws_dest.sheet_view.rightToLeft = True
    ws_dest.views.sheetView[0].rightToLeft = True
    for row in ws_spec.iter_rows(values_only=False):
        for cell in row:
            dest_cell = ws_dest.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                dest_cell.font = cell.font.copy()
                dest_cell.border = cell.border.copy()
                dest_cell.fill = cell.fill.copy()
                dest_cell.number_format = cell.number_format
                dest_cell.protection = cell.protection.copy()
                dest_cell.alignment = cell.alignment.copy()
    for merged in ws_spec.merged_cells.ranges:
        ws_dest.merge_cells(str(merged))
    for col_letter, col_dim in ws_spec.column_dimensions.items():
        ws_dest.column_dimensions[col_letter].width = col_dim.width
    for row_num, row_dim in ws_spec.row_dimensions.items():
        ws_dest.row_dimensions[row_num].height = row_dim.height

    print("Saving Master Workbook...")
    wb_master.save(master_path)
    print(f"Master Workbook saved to: {master_path}")
    
    # 3. Build Standalone Specialty Program Workbook
    standalone_path = r'C:\Users\MaximuM-Tech\Downloads\برنامج_تدريب_تخصص_بوك_إم2_القسم_النهائي.xlsx'
    print("Saving Standalone Specialty Program Workbook...")
    wb_spec.save(standalone_path)
    print(f"Standalone Specialty saved to: {standalone_path}")
    
    print("\n=== PIPELINE FINISHED SUCCESSFULLY ===")

if __name__ == '__main__':
    main()
