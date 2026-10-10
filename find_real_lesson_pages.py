import docx
import openpyxl
import sys
import os
import re

sys.stdout.reconfigure(encoding='utf-8')

def extract_paragraphs_with_pages(docx_path):
    doc = docx.Document(docx_path)
    current_page = 1
    para_list = []
    
    for p_idx, p in enumerate(doc.paragraphs, start=1):
        text = p.text.strip()
        p_start_page = current_page
        
        for r in p.runs:
            xml = r._r.xml
            if 'w:lastRenderedPageBreak' in xml or 'w:type="page"' in xml:
                current_page += 1
                
        p_end_page = current_page
        if text:
            para_list.append({
                'para_idx': p_idx,
                'text': text,
                'start_page': p_start_page,
                'end_page': p_end_page
            })
    return para_list, current_page

# Load source files
dhab_paras, dhab_total = extract_paragraphs_with_pages(r'C:\Users\MaximuM-Tech\Downloads\مرجع الضبع الاسود.docx')
ejla_paras, ejla_total = extract_paragraphs_with_pages(r'C:\Users\MaximuM-Tech\Downloads\مرجع الايجلا.docx')

print(f"مرجع الضبع الأسود: {len(dhab_paras)} فقرة نصية عبر {dhab_total} صفحة.")
print(f"مرجع الإيجلا: {len(ejla_paras)} فقرة نصية عبر {ejla_total} صفحة.")

# Load questions from the question bank to get rich search context for each lesson
src_wb = openpyxl.load_workbook(r'C:\Users\MaximuM-Tech\Downloads\بنوك\معدة الضبع الاسود.xlsx', data_only=True)

def get_lesson_questions_map(sheet_name):
    ws = src_wb[sheet_name]
    l_map = {}
    for r in range(3, ws.max_row + 1):
        q_text = str(ws.cell(r, 2).value or '')
        expl = str(ws.cell(r, 3).value or '')
        ans = str(ws.cell(r, 8).value or '')
        lesson = str(ws.cell(r, 10).value or '')
        if lesson:
            if lesson not in l_map:
                l_map[lesson] = []
            l_map[lesson].append(f"{q_text} {expl} {ans}")
    return l_map

prep_q_map = get_lesson_questions_map('بنك القسم الإعدادي')
med_q_map = get_lesson_questions_map('بنك القسم المتوسط')
fin_q_map = get_lesson_questions_map('بنك القسم النهائي')

# Match each lesson
def match_lesson_to_pages(lesson_title, questions_text_list, para_list):
    clean_title = re.sub(r'[\(\)\[\]\.\:\-]', ' ', lesson_title).strip()
    title_words = [w for w in clean_title.split() if len(w) > 2 and w not in ['المعدة', 'العام', 'عامة', 'الخواص', 'الفنية', 'طريقة', 'عمل', 'مكونات', 'جهاز', 'نظام', 'دائرة']]
    
    # Check direct occurrence in paragraphs
    best_para = None
    best_score = -1
    
    combined_q_text = " ".join(questions_text_list[:5]) if questions_text_list else ""
    q_words = [w for w in re.split(r'\s+', combined_q_text) if len(w) > 3 and w not in ['معدة', 'الضبع', 'الاسود', 'الايجلا', 'صاروخ', 'الصاروخ', 'الهدف']]
    
    for p in para_list:
        ptxt = p['text']
        score = 0
        # Check title match
        if lesson_title in ptxt or ptxt in lesson_title:
            score += 50
        for tw in title_words:
            if tw in ptxt:
                score += 10
        # Check questions terms match
        for qw in q_words[:30]:
            if qw in ptxt:
                score += 1
                
        if score > best_score:
            best_score = score
            best_para = p
            
    if best_para:
        return best_para['start_page'], max(best_para['start_page'], best_para['end_page'])
    return 1, 1

# Load syllabus items
ws_syl = src_wb['الجدول الزمني وفهرس الدروس']
syl_items = []
for r in range(3, ws_syl.max_row + 1):
    idx = ws_syl.cell(r, 1).value
    if idx is not None and isinstance(idx, (int, float)):
        syl_items.append((int(idx), str(ws_syl.cell(r, 2).value).strip()))

print("\n--- مطابقة دروس القسم الإعدادي مع صفحات مرجع الضبع الأسود الحقيقية ---")
prep_results = []
last_page = 1
for idx, title in syl_items[:28]:
    q_texts = prep_q_map.get(title, [])
    sp, ep = match_lesson_to_pages(title, q_texts, dhab_paras)
    # Ensure logical sequential flow
    sp = max(sp, last_page)
    ep = max(sp, ep)
    last_page = sp
    prep_results.append((idx, title, sp, ep))
    print(f"درس {idx:2d}: من {sp:2d} الي {ep:2d} | {title[:50]}")

print("\n--- مطابقة دروس القسم النهائي مع صفحات مرجع الإيجلا الحقيقية ---")
fin_results = []
last_page = 1
for idx, title in syl_items[84:124]:
    rel_id = idx - 84
    q_texts = fin_q_map.get(title, [])
    sp, ep = match_lesson_to_pages(title, q_texts, ejla_paras)
    sp = max(sp, last_page)
    ep = max(sp, ep)
    last_page = sp
    fin_results.append((rel_id, title, sp, ep))
    print(f"درس {rel_id:2d}: من {sp:2d} الي {ep:2d} | {title[:50]}")
