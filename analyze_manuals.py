import docx
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Check training program topics
excel_path = r'C:\Users\MaximuM-Tech\Downloads\برنامج تدريب تخصص الضبع الاسود + بنوك_محدث_LMS.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

for sheet_name in ['برنامج تدريب - القسم المتوسط', 'برنامج تدريب - القسم النهائي']:
    ws = wb[sheet_name]
    print('========================================')
    print('SHEET:', sheet_name)
    print('========================================')
    topics = []
    for r in range(5, ws.max_row + 1):
        lec = ws.cell(r, 3).value
        name = ws.cell(r, 4).value
        ps = ws.cell(r, 7).value
        pe = ws.cell(r, 8).value
        if lec in ['ف1', 'ف3']:
            topics.append({
                'num': len(topics) + 1,
                'name': name,
                'ps': ps,
                'pe': pe
            })
    for t in topics:
        print(f"{t['num']:2d} | {t['name']} | من {t['ps']} إلى {t['pe']}")
