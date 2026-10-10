import openpyxl, sys, collections
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx', data_only=True)
ws = wb['Questions']

seen = {}
dups = collections.defaultdict(list)

for r in range(2, ws.max_row + 1):
    txt = str(ws.cell(r, 2).value).strip() if ws.cell(r, 2).value else ''
    les = ws.cell(r, 10).value
    qtype = ws.cell(r, 1).value
    diff = ws.cell(r, 4).value
    ans = ws.cell(r, 5).value
    optA = ws.cell(r, 6).value
    optB = ws.cell(r, 7).value
    optC = ws.cell(r, 8).value
    optD = ws.cell(r, 9).value
    
    if txt not in seen:
        seen[txt] = (r, les)
    else:
        dups[les].append({
            'row': r,
            'orig_row': seen[txt][0],
            'orig_les': seen[txt][1],
            'type': qtype,
            'diff': diff,
            'ans': ans,
            'text': txt,
            'opts': [optA, optB, optC, optD]
        })

print(f"=== MEDIUM SECTION DUPLICATES ({sum(len(v) for v in dups.values())} rows) ===")
for les, items in dups.items():
    print(f"\nLesson: '{les}' -> {len(items)} duplicates")
    for it in items[:3]:
        print(f"  Row {it['row']} (copied from Row {it['orig_row']} in '{it['orig_les']}'):")
        print(f"    [{it['type']} | {it['diff']} | ans={it['ans']}] {it['text'][:80]}...")
