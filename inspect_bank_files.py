import openpyxl, sys, glob
sys.stdout.reconfigure(encoding='utf-8')

for pattern in [r'C:\Users\MaximuM-Tech\Downloads\بنك*.xlsx']:
    for f in glob.glob(pattern):
        try:
            wb = openpyxl.load_workbook(f, read_only=True)
            print(f"File: {f}")
            for s in wb.sheetnames:
                ws = wb[s]
                rows = list(ws.iter_rows(values_only=True, max_row=5))
                print(f"  Sheet: {s}, rows={ws.max_row}, sample={rows[0][:5] if rows else 'empty'}")
        except Exception as e:
            print(f"Error {f}: {e}")
