import openpyxl
import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

def get_doc_paragraphs(path):
    doc = docx.Document(path)
    cur_p = 1
    paras = []
    for i, p in enumerate(doc.paragraphs, 1):
        p_start = cur_p
        for r in p.runs:
            xml = r._r.xml
            if 'lastRenderedPageBreak' in xml or 'w:type="page"' in xml:
                cur_p += 1
        p_end = cur_p
        t = p.text.strip()
        paras.append({
            'idx': i,
            'start_page': p_start,
            'end_page': p_end,
            'text': t
        })
    return paras

dhab_paras = get_doc_paragraphs(r'C:\Users\MaximuM-Tech\Downloads\مرجع الضبع الاسود.docx')
ejla_paras = get_doc_paragraphs(r'C:\Users\MaximuM-Tech\Downloads\مرجع الايجلا.docx')

# Load questions for each lesson from the bank
wb = openpyxl.load_workbook(r'C:\Users\MaximuM-Tech\Downloads\بنوك\معدة الضبع الاسود.xlsx', data_only=True)

def get_questions_by_sheet(sheet_name):
    ws = wb[sheet_name]
    res = {}
    for r in range(3, ws.max_row + 1):
        q = str(ws.cell(r, 2).value or '').strip()
        ans = str(ws.cell(r, 8).value or '').strip()
        lesson = str(ws.cell(r, 10).value or '').strip()
        if lesson:
            if lesson not in res:
                res[lesson] = []
            res[lesson].append((q, ans))
    return res

prep_q = get_questions_by_sheet('بنك القسم الإعدادي')
med_q = get_questions_by_sheet('بنك القسم المتوسط')
fin_q = get_questions_by_sheet('بنك القسم النهائي')

ws_syl = wb['الجدول الزمني وفهرس الدروس']
syl_lessons = []
for r in range(3, ws_syl.max_row + 1):
    idx = ws_syl.cell(r, 1).value
    if idx is not None and isinstance(idx, (int, float)):
        syl_lessons.append((int(idx), str(ws_syl.cell(r, 2).value or '').strip()))

def find_best_page_range(lesson_title, questions, paras):
    # Search for headings and content matching lesson title and questions
    keywords = set()
    # clean lesson title
    clean = re.sub(r'[\(\)\[\]\.\:\-،\/]', ' ', lesson_title)
    for w in clean.split():
        if len(w) > 2 and w not in ['المعدة', 'العام', 'عامة', 'الخواص', 'الفنية', 'طريقة', 'عمل', 'مكونات', 'جهاز', 'نظام', 'دائرة', 'الصاروخ']:
            keywords.add(w)
            
    # Also extract words from first 3 questions
    for q, ans in questions[:3]:
        q_clean = re.sub(r'[\(\)\[\]\.\:\-،\/؟\?]', ' ', q + ' ' + ans)
        for w in q_clean.split():
            if len(w) > 3 and w not in ['المعدة', 'الضبع', 'الاسود', 'الايجلا', 'الصاروخ', 'صاروخ', 'السؤال', 'الإجابة', 'الصحيحة']:
                keywords.add(w)

    scores_by_page = {}
    for p in paras:
        if not p['text']:
            continue
        p_score = 0
        txt = p['text']
        if lesson_title in txt:
            p_score += 50
        for kw in keywords:
            if kw in txt:
                p_score += 5
        
        pg = p['start_page']
        scores_by_page[pg] = scores_by_page.get(pg, 0) + p_score
        if p['end_page'] != pg:
            scores_by_page[p['end_page']] = scores_by_page.get(p['end_page'], 0) + p_score

    # Sort pages by score
    sorted_pages = sorted(scores_by_page.items(), key=lambda x: x[1], reverse=True)
    if sorted_pages and sorted_pages[0][1] > 0:
        # take top pages that are adjacent or best page
        best_p = sorted_pages[0][0]
        # check if adjacent page has high score
        p_start = best_p
        p_end = best_p
        if scores_by_page.get(best_p + 1, 0) >= sorted_pages[0][1] * 0.4:
            p_end = best_p + 1
        elif scores_by_page.get(best_p - 1, 0) >= sorted_pages[0][1] * 0.4:
            p_start = best_p - 1
        return p_start, p_end, sorted_pages[0][1]
    return 1, 1, 0

print("Testing matching...")
print("\n=== PREPARATORY (1..28) ===")
for idx, title in syl_lessons[:28]:
    qs = prep_q.get(title, [])
    sp, ep, sc = find_best_page_range(title, qs, dhab_paras)
    print(f"Lec {idx:2d}: P{sp:2d} - P{ep:2d} (Score {sc:3d}) | {title}")

print("\n=== MEDIUM (29..84) ===")
for idx, title in syl_lessons[28:84]:
    qs = med_q.get(title, [])
    sp, ep, sc = find_best_page_range(title, qs, dhab_paras)
    print(f"Lec {idx-28:2d}: P{sp:2d} - P{ep:2d} (Score {sc:3d}) | {title}")

print("\n=== FINAL (85..124) ===")
for idx, title in syl_lessons[84:124]:
    qs = fin_q.get(title, [])
    sp, ep, sc = find_best_page_range(title, qs, ejla_paras)
    print(f"Lec {idx-84:2d}: P{sp:2d} - P{ep:2d} (Score {sc:3d}) | {title}")

