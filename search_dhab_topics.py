import docx, re, sys
sys.stdout.reconfigure(encoding='utf-8')

for path, title in [
    (r'C:\Users\MaximuM-Tech\Downloads\تكتيك و استخدام.docx', 'تكتيك واستخدام'),
    (r'C:\Users\MaximuM-Tech\Downloads\قواعد خدمة الميدان صار محمول 29-11.docx', 'قواعد خدمة الميدان')
]:
    doc = docx.Document(path)
    print(f"=== {title}: {len(doc.paragraphs)} paragraphs ===")
    for p in doc.paragraphs:
        t = p.text.strip()
        if any(k in t for k in ['أوضاع الإطلاق', 'المحرك الدافع', 'التوربين', 'الجيروسكوب', 'العدسي', 'موازنة دوران']):
            if len(t) > 20:
                print(f"  [{t[:90]}]")
