import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font, Alignment, Border, Side
import shutil
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Perfect 4-lesson repeating pattern for 25 questions per lesson:
# 30% easy (7.5 -> 7 or 8), 30% medium (7.5 -> 7 or 8), 15% hard (3.75 -> 4), 15% very hard (3.75 -> 4), 10% excellence (2.5 -> 2 or 3)
# Total per 4 lessons = 30 easy, 30 medium, 15 hard, 15 very hard, 10 excellence (Exact 100%)
# Alternating: Odd lessons = 13 MCQ + 12 TF, Even lessons = 12 MCQ + 13 TF
PATTERNS = [
    # L1 (idx % 4 == 0) -> 13 MCQ + 12 TF
    (
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1,
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*1 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1
    ),
    # L2 (idx % 4 == 1) -> 12 MCQ + 13 TF
    (
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*1 + ['5 - تفوق']*1,
        ['1 - سهل']*3 + ['2 - متوسط']*3 + ['3 - صعب']*3 + ['4 - صعب جدا']*2 + ['5 - تفوق']*2
    ),
    # L3 (idx % 4 == 2) -> 13 MCQ + 12 TF
    (
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*2 + ['5 - تفوق']*1,
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*2 + ['4 - صعب جدا']*1 + ['5 - تفوق']*1
    ),
    # L4 (idx % 4 == 3) -> 12 MCQ + 13 TF
    (
        ['1 - سهل']*3 + ['2 - متوسط']*3 + ['3 - صعب']*2 + ['4 - صعب جدا']*2 + ['5 - تفوق']*2,
        ['1 - سهل']*4 + ['2 - متوسط']*4 + ['3 - صعب']*1 + ['4 - صعب جدا']*3 + ['5 - تفوق']*1
    )
]

template_path = r'C:\Users\MaximuM-Tech\Downloads\نموذج تسجيل الاسئلة (1).xlsx'
src_bank_path = r'C:\Users\MaximuM-Tech\Downloads\avenger_bank_from_drive.xlsx'
out_path = r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_تخصص_افنجر_LMS_محدث.xlsx'

print('Loading source files...')
wb_src = openpyxl.load_workbook(src_bank_path, data_only=True)
ws_src = wb_src['Questions']

with open('e:/bank/avenger_missing_questions_cache.json', 'r', encoding='utf-8') as f:
    missing_cache = json.load(f)

# Extract original 32 lessons in order
lessons_order = []
for r in range(2, 692):
    les = ws_src.cell(r, 10).value
    if les and les not in lessons_order:
        lessons_order.append(les)

print(f'Found {len(lessons_order)} unique lessons in source bank.')

# Group existing questions by lesson
existing_by_lesson = {les: {'mcq': [], 'tf': []} for les in lessons_order}
for r in range(2, 692):
    q_type = ws_src.cell(r, 1).value
    q_text = ws_src.cell(r, 2).value
    if not q_text:
        continue
    exp = ws_src.cell(r, 3).value
    diff = ws_src.cell(r, 4).value
    ans = ws_src.cell(r, 5).value
    optA = ws_src.cell(r, 6).value
    optB = ws_src.cell(r, 7).value
    optC = ws_src.cell(r, 8).value
    optD = ws_src.cell(r, 9).value
    les = ws_src.cell(r, 10).value
    
    item = {
        'row': r,
        'text': q_text,
        'exp': None, # explicitly clear explanation
        'ans': ans,
        'optA': optA,
        'optB': optB,
        'optC': optC,
        'optD': optD,
        'lesson': les
    }
    if 'متعدد' in str(q_type):
        existing_by_lesson[les]['mcq'].append(item)
    else:
        existing_by_lesson[les]['tf'].append(item)

# Build unified 25 questions per lesson
final_lessons_data = []

for idx, les in enumerate(lessons_order, 1):
    mcq_list = existing_by_lesson[les]['mcq']
    tf_list = existing_by_lesson[les]['tf']
    
    req_mcq = 13 if idx % 2 == 1 else 12
    req_tf = 12 if idx % 2 == 1 else 13
    
    # Trim duplicates if lesson has excess (e.g. L4 has 15 MCQs, keep first 12; L32 has 13 MCQs, keep first 12)
    if len(mcq_list) > req_mcq:
        mcq_list = mcq_list[:req_mcq]
        
    # Append missing MCQs from cache
    cached_missing = missing_cache.get(str(idx), {})
    for m in cached_missing.get('mcq', []):
        if len(mcq_list) < req_mcq:
            mcq_list.append({
                'text': m['text'],
                'exp': None,
                'ans': m['ans'],
                'optA': m['opt_a'],
                'optB': m['opt_b'],
                'optC': m['opt_c'],
                'optD': m['opt_d'],
                'lesson': les
            })
            
    # Append missing TFs from cache
    for t in cached_missing.get('tf', []):
        if len(tf_list) < req_tf:
            tf_list.append({
                'text': t['text'],
                'exp': None,
                'ans': t['ans'],
                'optA': None,
                'optB': None,
                'optC': None,
                'optD': None,
                'lesson': les
            })
            
    assert len(mcq_list) == req_mcq, f'L{idx} ({les}): expected {req_mcq} MCQs, got {len(mcq_list)}'
    assert len(tf_list) == req_tf, f'L{idx} ({les}): expected {req_tf} TFs, got {len(tf_list)}'
    
    final_lessons_data.append((les, mcq_list, tf_list))

print(f'Successfully assembled {len(final_lessons_data)} lessons x 25 questions = 800 questions.')

# Create output workbook from clean template
shutil.copy2(template_path, out_path)
wb_out = openpyxl.load_workbook(out_path)

# 1. Update Lookups sheet
ws_lookups = wb_out['Lookups']
for r in range(2, ws_lookups.max_row + 10):
    ws_lookups.cell(r, 4).value = None

for idx, les_name in enumerate(lessons_order, start=2):
    ws_lookups.cell(idx, 4, les_name)

max_les_row = len(lessons_order) + 1
wb_out.defined_names['L_4'] = DefinedName('L_4', attr_text=f"Lookups!$D$2:$D${max_les_row}")

# 2. Populate Questions sheet
ws_q = wb_out['Questions']
while ws_q.max_row > 1:
    ws_q.delete_rows(2)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
regular_font = Font(name='Calibri', size=11, bold=False)
center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)

current_row = 2
stat_types = {}
stat_diffs = {}
stat_diffs_mcq = {}
stat_diffs_tf = {}

for les_idx, (les_name, mcqs, tfs) in enumerate(final_lessons_data):
    mcq_diffs, tf_diffs = PATTERNS[les_idx % 4]
    
    # Write MCQs
    for q_i, item in enumerate(mcqs):
        diff = mcq_diffs[q_i]
        q_type = '1 - اختيار من متعدد'
        ans = str(item['ans']).strip()
        if ans not in ['A', 'B', 'C', 'D']:
            if 'أ' in ans or 'A' in ans: ans = 'A'
            elif 'ب' in ans or 'B' in ans: ans = 'B'
            elif 'ج' in ans or 'C' in ans: ans = 'C'
            elif 'د' in ans or 'D' in ans: ans = 'D'
            else: ans = 'A'

        row_data = [
            q_type,
            item['text'],
            None, # Explanation cleared
            diff,
            ans,
            item['optA'],
            item['optB'],
            item['optC'],
            item['optD'],
            les_name
        ]
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws_q.cell(current_row, c_idx, val)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = center_align if c_idx in [1, 4, 5] else right_align

        current_row += 1
        stat_types[q_type] = stat_types.get(q_type, 0) + 1
        stat_diffs[diff] = stat_diffs.get(diff, 0) + 1
        stat_diffs_mcq[diff] = stat_diffs_mcq.get(diff, 0) + 1

    # Write TFs
    for q_i, item in enumerate(tfs):
        diff = tf_diffs[q_i]
        q_type = '2 - صح/خطأ'
        ans_raw = str(item['ans']).strip().lower()
        if ans_raw in ['true', 'صواب', 'صح', 'a', 't']:
            norm_ans = 'True'
        else:
            norm_ans = 'False'

        row_data = [
            q_type,
            item['text'],
            None, # Explanation cleared
            diff,
            norm_ans,
            None, # Col F strictly cleared
            None, # Col G strictly cleared
            None, # Col H strictly cleared
            None, # Col I strictly cleared
            les_name
        ]
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws_q.cell(current_row, c_idx, val)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = center_align if c_idx in [1, 4, 5] else right_align

        current_row += 1
        stat_types[q_type] = stat_types.get(q_type, 0) + 1
        stat_diffs[diff] = stat_diffs.get(diff, 0) + 1
        stat_diffs_tf[diff] = stat_diffs_tf.get(diff, 0) + 1

# Configure Data Validations
max_q_row = current_row - 1
ws_q.data_validations.dataValidation.clear()

dv_qtype = DataValidation(type="list", formula1="L_1", allow_blank=True)
ws_q.add_data_validation(dv_qtype)
dv_qtype.add(f"A2:A{max_q_row + 200}")

dv_diff = DataValidation(type="list", formula1="L_2", allow_blank=True)
ws_q.add_data_validation(dv_diff)
dv_diff.add(f"D2:D{max_q_row + 200}")

dv_ans = DataValidation(type="list", formula1="L_3", allow_blank=True)
ws_q.add_data_validation(dv_ans)
dv_ans.add(f"E2:E{max_q_row + 200}")

dv_les = DataValidation(type="list", formula1="L_4", allow_blank=True)
ws_q.add_data_validation(dv_les)
dv_les.add(f"J2:J{max_q_row + 200}")

wb_out.save(out_path)
print(f'Successfully saved compliant bank to: {out_path}')
print(f'Total questions: {max_q_row - 1}')
print(f'Question Types: {stat_types}')
print('Overall Difficulties:')
tot_q = max_q_row - 1
for k in sorted(stat_diffs.keys()):
    v = stat_diffs[k]
    print(f'  {k}: {v} ({v/tot_q*100:.2f}%)')
print('MCQ Difficulties:')
tot_mcq = stat_types['1 - اختيار من متعدد']
for k in sorted(stat_diffs_mcq.keys()):
    v = stat_diffs_mcq[k]
    print(f'  {k}: {v} ({v/tot_mcq*100:.2f}%)')
print('TF Difficulties:')
tot_tf = stat_types['2 - صح/خطأ']
for k in sorted(stat_diffs_tf.keys()):
    v = stat_diffs_tf[k]
    print(f'  {k}: {v} ({v/tot_tf*100:.2f}%)')
