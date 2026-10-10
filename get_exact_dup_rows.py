import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

for tag, fname in [('Medium', 'بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'), ('Final', 'بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx')]:
    wb = openpyxl.load_workbook(r'C:\Users\MaximuM-Tech\Downloads\\' + fname, data_only=True)
    ws = wb['Questions']
    seen = {}
    dups_info = []
    for r in range(2, ws.max_row + 1):
        txt = str(ws.cell(r, 2).value).strip() if ws.cell(r, 2).value else ''
        les = ws.cell(r, 10).value
        qtype = ws.cell(r, 1).value
        diff = ws.cell(r, 4).value
        ans = ws.cell(r, 5).value
        if txt not in seen:
            seen[txt] = r
        else:
            dups_info.append((r, les, qtype, diff, ans, seen[txt], txt))
            
    print(f"=== {tag} Duplicates: {len(dups_info)} ===")
    lessons = {}
    for it in dups_info:
        lessons.setdefault(it[1], []).append(it)
    for les, lst in lessons.items():
        print(f"Lesson: {les} ({len(lst)} rows): rows {[x[0] for x in lst]}")
