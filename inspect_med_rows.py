import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx', data_only=True)
ws = wb['Questions']

rows_to_check = [
    33, 156,
    *range(77, 99),
    252, 253, 254, 255, 256, 257, 258, 259, 260, 261, 262, 263, 265, 266, 267, 268, 269, 270, 271, 272, 273, 274,
    *range(327, 349),
    *range(402, 427),
    *range(577, 599),
    1052, 1053, 1054, 1055, 1056, 1057, 1058, 1059, 1060, 1061, 1062, 1063, 1065, 1066, 1067, 1068, 1069, 1070, 1071, 1072, 1073, 1074
]

print(f"Total rows to inspect in Medium: {len(rows_to_check)}")
for r in rows_to_check[:10]:
    print(f"Row {r}: Type={ws.cell(r,1).value} | Diff={ws.cell(r,4).value} | Les={ws.cell(r,10).value}")
