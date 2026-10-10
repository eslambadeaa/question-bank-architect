import docx
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Let's see the pages of Black Hyena docx
doc_h = docx.Document(r'C:\Users\MaximuM-Tech\Downloads\ضبع اسود\مرجع الضبع الاسود.docx')
pages_h = {1: []}
cp = 1
for p in doc_h.paragraphs:
    for r in p.runs:
        xml = r._r.xml
        if 'w:lastRenderedPageBreak' in xml or ('w:br' in xml and 'type="page"' in xml):
            cp += 1
            pages_h[cp] = []
    t = p.text.strip()
    if t:
        pages_h[cp].append(t)

print("=== Black Hyena Manual Pages (Total:", len(pages_h), ") ===")
for p_num in range(1, len(pages_h)+1):
    first_lines = pages_h[p_num][:2]
    print(f"Page {p_num:02d}: {' // '.join(first_lines)[:100]}")

# Let's see Igla docx
doc_i = docx.Document(r'C:\Users\MaximuM-Tech\Downloads\ضبع اسود\مرجع الايجلا.docx')
pages_i = {1: []}
cp = 1
for p in doc_i.paragraphs:
    for r in p.runs:
        xml = r._r.xml
        if 'w:lastRenderedPageBreak' in xml or ('w:br' in xml and 'type="page"' in xml):
            cp += 1
            pages_i[cp] = []
    t = p.text.strip()
    if t:
        pages_i[cp].append(t)

print("\n=== Igla Manual Pages (Total:", len(pages_i), ") ===")
for p_num in range(1, len(pages_i)+1):
    first_lines = pages_i[p_num][:2]
    print(f"Page {p_num:02d}: {' // '.join(first_lines)[:100]}")
