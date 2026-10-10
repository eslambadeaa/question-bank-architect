import openpyxl
import collections
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

def norm(text):
    if not text: return ''
    t = str(text).strip()
    t = re.sub(r'[\s\.\?\:\!\-_،]+', ' ', t)
    return t.strip()

for fname in ['بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx', 'بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx']:
    fpath = r'C:\Users\MaximuM-Tech\Downloads\\' + fname
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb['Questions']
    print('='*60)
    print(f'Checking: {fname} (Rows: {ws.max_row})')
    
    exact_texts = collections.defaultdict(list)
    norm_texts = collections.defaultdict(list)
    
    for r in range(2, ws.max_row + 1):
        txt = ws.cell(r, 2).value
        if txt:
            exact_texts[str(txt).strip()].append(r)
            norm_texts[norm(txt)].append(r)
            
    exact_dups = {k: v for k, v in exact_texts.items() if len(v) > 1}
    norm_dups = {k: v for k, v in norm_texts.items() if len(v) > 1}
    print(f'  Exact duplicates (Highlight Duplicates in Excel): {len(exact_dups)} texts affecting {sum(len(v) for v in exact_dups.values())} rows')
    print(f'  Normalized duplicates: {len(norm_dups)} texts affecting {sum(len(v) for v in norm_dups.values())} rows')
    if exact_dups:
        print('  Top duplicate examples:')
        for k, v in list(exact_dups.items())[:5]:
            les_list = [ws.cell(r, 10).value for r in v]
            print(f'    - \"{k[:45]}...\" in rows {v} | Lessons: {set(les_list)}')
