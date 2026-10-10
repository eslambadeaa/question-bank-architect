import openpyxl
import collections
import sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'C:\Users\MaximuM-Tech\Downloads\avenger_bank_from_drive.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)
ws = wb['Questions']

print('Total rows in Questions:', ws.max_row)
rows = []
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 15)]
    if vals[1]: # if question text exists
        rows.append((r, vals))

print(f'Populated question rows: {len(rows)}')

# Group by lesson (Col 10)
by_lesson = collections.OrderedDict()
for r, vals in rows:
    les = str(vals[9]).strip() if vals[9] is not None else 'None'
    if les not in by_lesson:
        by_lesson[les] = []
    by_lesson[les].append((r, vals))

print(f'Number of lessons in Questions: {len(by_lesson)}')
for les, q_list in by_lesson.items():
    mcq = [x for x in q_list if 'متعدد' in str(x[1][0])]
    tf = [x for x in q_list if 'صح' in str(x[1][0])]
    diffs = collections.Counter([x[1][3] for x in q_list])
    print(f'  الدرس: \"{les}\" -> {len(q_list)} سؤال (MCQ: {len(mcq)}, TF: {len(tf)}) | الصعوبات: {dict(diffs)}')

# Lookups sheet lessons
ws_lookups = wb['Lookups']
print('\nLessons in Lookups sheet (Col D):')
lookup_lessons = []
for r in range(2, ws_lookups.max_row + 1):
    val = ws_lookups.cell(r, 4).value
    if val:
        lookup_lessons.append(str(val).strip())
print(f'Total lookup lessons: {len(lookup_lessons)}')
for i, l in enumerate(lookup_lessons, 1):
    print(f'  {i}. {l}')
