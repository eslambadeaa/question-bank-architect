"""
Script to create standalone Index/TOC documents and updated reference manuals with TOC & page numbers:
With FULL Right-to-Left (RTL) layout: <w:bidiVisual/> on tables, <w:bidi/> on paragraphs, and Arabic alignment.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

# 1. LOAD DATA FROM APPROVED TRAINING PROGRAM
excel_path = r'C:\Users\MaximuM-Tech\Downloads\برنامج تدريب تخصص الضبع الاسود + بنوك_محدث_LMS.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

# Medium (Black Hyena)
ws_med = wb['برنامج تدريب - القسم المتوسط']
hyena_topics = []
for r in range(5, ws_med.max_row + 1):
    lec = ws_med.cell(r, 3).value
    name = ws_med.cell(r, 4).value
    ps = ws_med.cell(r, 7).value
    pe = ws_med.cell(r, 8).value
    if lec in ['ف1', 'ف3']:
        term = 'الترم الأول' if len(hyena_topics) < 24 else 'الترم الثاني'
        hyena_topics.append({
            'num': len(hyena_topics) + 1,
            'name': name,
            'term': term,
            'p_start': ps,
            'p_end': pe
        })

print(f'Loaded {len(hyena_topics)} topics for Black Hyena (Medium Section).')

# Final (Igla)
ws_fin = wb['برنامج تدريب - القسم النهائي']
igla_topics = []
for r in range(5, ws_fin.max_row + 1):
    lec = ws_fin.cell(r, 3).value
    name = ws_fin.cell(r, 4).value
    ps = ws_fin.cell(r, 7).value
    pe = ws_fin.cell(r, 8).value
    if lec in ['ف1', 'ف3']:
        term = 'الترم الأول' if len(igla_topics) < 24 else 'الترم الثاني'
        igla_topics.append({
            'num': len(igla_topics) + 1,
            'name': name,
            'term': term,
            'p_start': ps,
            'p_end': pe
        })

print(f'Loaded {len(igla_topics)} topics for Igla (Final Section).')

# RTL XML Helper Functions
def make_table_rtl(table):
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml('<w:bidiVisual %s/>' % nsdecls('w')))

def set_paragraph_bidi(p):
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml('<w:bidi %s/>' % nsdecls('w')))

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml('<w:shd %s w:fill="%s"/>' % (nsdecls('w'), fill_hex))
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml('<w:tcMar %s><w:top w:w="%s" w:type="dxa"/><w:bottom w:w="%s" w:type="dxa"/><w:left w:w="%s" w:type="dxa"/><w:right w:w="%s" w:type="dxa"/></w:tcMar>' % (nsdecls('w'), top, bottom, left, right))
    tcPr.append(tcMar)

def set_cell_borders(cell, top='single', bottom='single', left='single', right='single', color='D0D0D0', sz='4'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml('<w:tcBorders %s><w:top w:val="%s" w:sz="%s" w:space="0" w:color="%s"/><w:bottom w:val="%s" w:sz="%s" w:space="0" w:color="%s"/><w:left w:val="%s" w:sz="%s" w:space="0" w:color="%s"/><w:right w:val="%s" w:sz="%s" w:space="0" w:color="%s"/></w:tcBorders>' % (nsdecls('w'), top, sz, color, bottom, sz, color, left, sz, color, right, sz, color))
    tcPr.append(tcBorders)

def add_page_number_field(run):
    fldChar1 = parse_xml('<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml('<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml('<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml('<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_toc_table(doc, topics, title_text, system_name):
    # Title
    p_title = doc.add_paragraph()
    set_paragraph_bidi(p_title)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run(f'فهرس الموضوعات الفنية والتدريبية لمعدة {system_name}')
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(16)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    p_sub = doc.add_paragraph()
    set_paragraph_bidi(p_sub)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_s = p_sub.add_run(f'وفقاً لبرنامج التدريب المعتمد رسمياً للعام التدريبي - {title_text}')
    run_s.font.name = 'Arial'
    run_s.font.size = Pt(11)
    run_s.font.bold = True
    run_s.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # Table: 5 columns: م | اسم الموضوع التدريبي | الترم | الصفحة (من) | الصفحة (إلى)
    table = doc.add_table(rows=1, cols=5)
    make_table_rtl(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Widths in inches
    col_widths = [Inches(0.5), Inches(4.3), Inches(1.0), Inches(0.8), Inches(0.8)]

    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_titles = ['م', 'اسم الموضوع التدريبي في المرجع', 'الترم', 'من صفحة', 'إلى صفحة']
    for idx, (cell, h_text) in enumerate(zip(hdr_cells, hdr_titles)):
        cell.width = col_widths[idx]
        set_cell_background(cell, '1F4E78')
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        set_cell_borders(cell, color='1F4E78')
        p = cell.paragraphs[0]
        set_paragraph_bidi(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(h_text)
        run.font.name = 'Arial'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data Rows
    current_term = None
    for item in topics:
        # Check if term changed to insert a divider
        if item['term'] != current_term:
            current_term = item['term']
            row_term = table.add_row()
            # merge all 5 cells
            a = row_term.cells[0]
            b = row_term.cells[4]
            a.merge(b)
            set_cell_background(a, 'D9E1F2')
            set_cell_margins(a, top=100, bottom=100, left=150, right=150)
            set_cell_borders(a, color='B0C4DE')
            p = a.paragraphs[0]
            set_paragraph_bidi(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            r = p.add_run(f'--- موضوعات {current_term} ---')
            r.font.name = 'Arial'
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

        row = table.add_row()
        cells = row.cells
        for idx in range(5):
            cells[idx].width = col_widths[idx]
            bg_color = 'F9FAFB' if item['num'] % 2 == 0 else 'FFFFFF'
            set_cell_background(cells[idx], bg_color)
            set_cell_margins(cells[idx], top=80, bottom=80, left=120, right=120)
            set_cell_borders(cells[idx], color='E0E0E0')

        # Cell 0: Num
        p0 = cells[0].paragraphs[0]
        set_paragraph_bidi(p0)
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.space_before = Pt(0)
        r0 = p0.add_run(str(item['num']))
        r0.font.name = 'Arial'
        r0.font.size = Pt(10)
        r0.font.bold = True

        # Cell 1: Name
        p1 = cells[1].paragraphs[0]
        set_paragraph_bidi(p1)
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_before = Pt(0)
        r1 = p1.add_run(item['name'])
        r1.font.name = 'Arial'
        r1.font.size = Pt(10)

        # Cell 2: Term
        p2 = cells[2].paragraphs[0]
        set_paragraph_bidi(p2)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.space_before = Pt(0)
        r2 = p2.add_run(item['term'])
        r2.font.name = 'Arial'
        r2.font.size = Pt(9.5)

        # Cell 3: Start Page
        p3 = cells[3].paragraphs[0]
        set_paragraph_bidi(p3)
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(0)
        p3.paragraph_format.space_before = Pt(0)
        r3 = p3.add_run(str(item['p_start']))
        r3.font.name = 'Arial'
        r3.font.size = Pt(10)
        r3.font.bold = True
        r3.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

        # Cell 4: End Page
        p4 = cells[4].paragraphs[0]
        set_paragraph_bidi(p4)
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p4.paragraph_format.space_after = Pt(0)
        p4.paragraph_format.space_before = Pt(0)
        r4 = p4.add_run(str(item['p_end']))
        r4.font.name = 'Arial'
        r4.font.size = Pt(10)
        r4.font.bold = True
        r4.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

# 2. GENERATE STANDALONE INDEX DOCUMENTS
def create_standalone_index_doc(topics, system_name, subtitle, out_path):
    doc = docx.Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    # Header / Footer
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    set_paragraph_bidi(p_ftr)
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ftr_txt = p_ftr.add_run(f'فهرس مرجع {system_name}  |  صفحة ')
    r_ftr_txt.font.name = 'Arial'
    r_ftr_txt.font.size = Pt(9)
    r_ftr_txt.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
    r_page = p_ftr.add_run()
    r_page.font.name = 'Arial'
    r_page.font.size = Pt(9)
    r_page.font.bold = True
    add_page_number_field(r_page)

    add_toc_table(doc, topics, subtitle, system_name)
    try:
        doc.save(out_path)
        print(f'Saved standalone index with RTL: {out_path}')
    except PermissionError:
        base, ext = os.path.splitext(out_path)
        alt_path = f"{base}_RTL{ext}"
        doc.save(alt_path)
        print(f'Locked, saved alternative: {alt_path}')

# 3. GENERATE FULL REFERENCE MANUALS WITH TOC & PAGE NUMBERS
def create_manual_with_toc(orig_manual_path, topics, system_name, subtitle, out_path):
    doc_orig = docx.Document(orig_manual_path)
    doc_new = docx.Document()

    sec_new = doc_new.sections[0]
    sec_orig = doc_orig.sections[0]
    sec_new.top_margin = sec_orig.top_margin
    sec_new.bottom_margin = sec_orig.bottom_margin
    sec_new.left_margin = sec_orig.left_margin
    sec_new.right_margin = sec_orig.right_margin

    # Footer
    footer = sec_new.footer
    p_ftr = footer.paragraphs[0]
    set_paragraph_bidi(p_ftr)
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ftr_label = p_ftr.add_run(f'مرجع {system_name}  |  صفحة ')
    r_ftr_label.font.name = 'Arial'
    r_ftr_label.font.size = Pt(9.5)
    r_ftr_label.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    r_page = p_ftr.add_run()
    r_page.font.name = 'Arial'
    r_page.font.size = Pt(9.5)
    r_page.font.bold = True
    r_page.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    add_page_number_field(r_page)

    # Header
    header = sec_new.header
    p_hdr = header.paragraphs[0]
    set_paragraph_bidi(p_hdr)
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_hdr.add_run(f'المرجع الفني والتدريبي المعتمد - منظومة {system_name}')
    r_hdr.font.name = 'Arial'
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.color.rgb = RGBColor(0x8C, 0x8C, 0x8C)

    # 1. Add TOC Table with RTL at the very beginning
    add_toc_table(doc_new, topics, subtitle, system_name)

    # 2. Add Page Break after TOC
    doc_new.add_page_break()

    # 3. Append original document paragraphs
    for p in doc_orig.paragraphs:
        p_copy = doc_new.add_paragraph()
        set_paragraph_bidi(p_copy)
        p_copy.alignment = p.alignment
        p_copy.paragraph_format.space_before = p.paragraph_format.space_before
        p_copy.paragraph_format.space_after = p.paragraph_format.space_after
        p_copy.paragraph_format.line_spacing = p.paragraph_format.line_spacing

        for r in p.runs:
            r_copy = p_copy.add_run(r.text)
            r_copy.font.name = r.font.name if r.font.name else 'Arial'
            r_copy.font.size = r.font.size if r.font.size else Pt(12)
            r_copy.font.bold = r.font.bold
            r_copy.font.italic = r.font.italic
            r_copy.font.underline = r.font.underline
            if r.font.color and r.font.color.rgb:
                r_copy.font.color.rgb = r.font.color.rgb

    try:
        doc_new.save(out_path)
        print(f'Saved full manual with RTL TOC: {out_path}')
    except PermissionError:
        base, ext = os.path.splitext(out_path)
        alt_path = f"{base}_RTL{ext}"
        doc_new.save(alt_path)
        print(f'Locked, saved alternative: {alt_path}')


# EXECUTE ALL GENERATIONS
print('\n=== 1. Generating Standalone TOC Documents (RTL) ===')
toc_hyena_standalone = r'C:\Users\MaximuM-Tech\Downloads\فهرس_مرجع_الضبع_الاسود_القسم_المتوسط.docx'
create_standalone_index_doc(hyena_topics, 'الضبع الأسود (9K32M)', 'القسم المتوسط (52 موضوعاً)', toc_hyena_standalone)

toc_igla_standalone = r'C:\Users\MaximuM-Tech\Downloads\فهرس_مرجع_الايجلا_القسم_النهائي.docx'
create_standalone_index_doc(igla_topics, 'الإيجلا (9K38 Igla)', 'القسم النهائي (44 موضوعاً)', toc_igla_standalone)

print('\n=== 2. Generating Full Manuals with Integrated RTL TOC & Page Numbers ===')
orig_hyena = r'C:\Users\MaximuM-Tech\Downloads\ضبع اسود\مرجع الضبع الاسود.docx'
full_hyena = r'C:\Users\MaximuM-Tech\Downloads\مرجع_الضبع_الاسود_بالفهرس_وترقيم_الصفحات.docx'
create_manual_with_toc(orig_hyena, hyena_topics, 'الضبع الأسود (9K32M)', 'القسم المتوسط (52 موضوعاً)', full_hyena)

orig_igla = r'C:\Users\MaximuM-Tech\Downloads\ضبع اسود\مرجع الايجلا.docx'
full_igla = r'C:\Users\MaximuM-Tech\Downloads\مرجع_الايجلا_بالفهرس_وترقيم_الصفحات.docx'
create_manual_with_toc(orig_igla, igla_topics, 'الإيجلا (9K38 Igla)', 'القسم النهائي (44 موضوعاً)', full_igla)

print('\n=== ALL DELIVERABLES UPDATED WITH FULL RTL SUPPORT ===')
