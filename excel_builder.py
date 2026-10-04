"""
Excel Builder Module for Question Banks & Syllabus Schedules
Full RTL layout, custom difficulty color coding, and strictly formatted Syllabus Index (1 : 2 : 25).
حقوق الملكية الفكرية وتطوير النظام: إسلام عبد البديع
"""

import io
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Difficulty Color Mapping as strictly required
DIFFICULTY_COLORS = {
    "سهل": "E2EFDA",
    "متوسط": "FFF2CC",
    "صعب": "FCE4D6",
    "صعب جداً": "F8CBAD",
    "تفوق": "A9D08E"
}

# Neutral workbook headers
BRAND_NOTICE = "منظومة بنوك الأسئلة القياسية"
FOOTER_NOTICE = "تم إعداد وتنسيق هذا النموذج آلياً وفق المعايير القياسية للتقييم"

# Palette definitions
NAVY_HEADER = "1F4E78"
ACCENT_BLUE = "D9E1F2"
GRAY_ZEBRA = "F9FAFB"
BORDER_GRAY = "D3D3D3"
TOTAL_ROW_FILL = "EAEEF3"


def create_thin_border() -> Border:
    thin = Side(border_style="thin", color=BORDER_GRAY)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def apply_rtl_and_grid(ws):
    """Configures Right-to-Left orientation and enables visible gridlines."""
    try:
        if hasattr(ws, 'sheet_view'):
            ws.sheet_view.rightToLeft = True
            ws.sheet_view.showGridLines = True
    except Exception:
        pass
    try:
        if hasattr(ws, 'views') and len(ws.views.sheetView) > 0:
            ws.views.sheetView[0].rightToLeft = True
            ws.views.sheetView[0].showGridLines = True
    except Exception:
        pass


def auto_fit_columns(ws, max_cols: int, min_width: int = 12, max_width: int = 65):
    """Adjusts column widths dynamically for Arabic text."""
    for col_idx in range(1, max_cols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(3, ws.max_row + 1):
            cell_val = ws.cell(row=row, column=col_idx).value
            if cell_val is not None:
                lines = str(cell_val).split("\n")
                line_len = max(len(l) for l in lines) if lines else 0
                max_len = max(max_len, line_len)
        calculated_width = max(min_width, min(max_len + 4, max_width))
        ws.column_dimensions[col_letter].width = calculated_width


ARABIC_DAYS = [
    "الأول", "الثاني", "الثالث", "الرابع", "الخامس", "السادس",
    "السابع", "الثامن", "التاسع", "العاشر", "الحادي عشر", "الثانى عشر",
    "الثالث عشر", "الرابع عشر", "الخامس عشر", "السادس عشر"
]


def calc_page_range(idx: int, total_items: int, total_pages: int = 25):
    """Calculates realistic proportional page ranges for lectures."""
    if total_items <= 0:
        return 1, 1
    start_p = 1 + int((idx - 1) * max(1, total_pages - 1) / total_items)
    end_p = max(start_p, 1 + int(idx * max(1, total_pages - 1) / total_items))
    return start_p, end_p


def add_training_program_sheet(
    wb: openpyxl.Workbook,
    sheet_title: str,
    section_label: str,
    specialty_name: str,
    items: List[Dict[str, Any]],
    total_pages: int = 33,
    days_count: Optional[int] = None,
    midterm_day: Optional[int] = None,
    total_hours: Optional[float] = None,
    total_questions: Optional[int] = None
):
    """
    Builds an institutional training program sheet matching the official committee template:
    - 4 periods (ف1..ف4) per day with merged day cells.
    - Merged single topic column (اسم الموضوع).
    - Hours (2.0), reference pages (من / الي), and questions count (25).
    - Integrated Midterm and Final exam rows with total formulas/values.
    """
    ws = wb.create_sheet(title=sheet_title)
    apply_rtl_and_grid(ws)

    border_cell = create_thin_border()
    font_title = Font(name='Times New Roman', size=22, bold=True)
    font_header = Font(name='Times New Roman', size=15, bold=True)
    font_data = Font(name='Times New Roman', size=13, bold=False)
    font_exam = Font(name='Times New Roman', size=15, bold=True)
    font_total = Font(name='Times New Roman', size=15, bold=True)

    header_fill = PatternFill(start_color='E8EEF5', end_color='E8EEF5', fill_type='solid')
    exam_fill = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
    total_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')

    # Row 1: Official Specialty & Section Banner
    ws.merge_cells('A1:H1')
    t_cell = ws['A1']
    clean_sec = section_label.replace("القسم ", "").strip()
    t_cell.value = f"برنامج محاضرات تخصص ( {specialty_name} ) للقسم ( {clean_sec} )"
    t_cell.font = font_title
    t_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 10

    # Header Rows 3 and 4
    ws.merge_cells('A3:A4')
    ws.merge_cells('B3:B4')
    ws.merge_cells('C3:C4')
    ws.merge_cells('D3:D4')
    ws.merge_cells('E3:E4')
    ws.merge_cells('F3:G3')
    ws.merge_cells('H3:H4')

    headers_3 = {
        'A3': 'الترم',
        'B3': 'اليوم',
        'C3': 'المحاضرة',
        'D3': 'اسم الموضوع',
        'E3': 'عدد الساعات',
        'F3': 'الصفحة في المرجع الموحد',
        'H3': 'عدد الاسئلة'
    }
    for cell_id, text in headers_3.items():
        c = ws[cell_id]
        c.value = text
        c.font = font_header
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border_cell

    ws['F4'].value = 'من'
    ws['G4'].value = 'الي'
    for c_id in ['F4', 'G4']:
        c = ws[c_id]
        c.font = font_header
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border_cell

    for r in [3, 4]:
        ws.row_dimensions[r].height = 25
        for col in range(1, 9):
            ws.cell(r, col).border = border_cell
            if ws.cell(r, col).fill.fill_type is None:
                ws.cell(r, col).fill = header_fill

    n_items = len(items)
    if days_count is None:
        days_count = max(1, (n_items + 3) // 4)
    if midterm_day is None:
        midterm_day = max(1, days_count // 2)
    if total_hours is None:
        total_hours = n_items * 2.0
    if total_questions is None:
        total_questions = n_items * 25

    current_row = 5
    item_idx = 0
    term_start_row = current_row

    for day_num in range(1, days_count + 1):
        day_label = ARABIC_DAYS[day_num - 1] if day_num <= len(ARABIC_DAYS) else f'اليوم {day_num}'
        day_start_row = current_row

        for period_idx in range(1, 5):
            it = items[item_idx] if item_idx < n_items else None
            p_label = f'ف{period_idx}'
            lec_title = it.get('lesson_name') or it.get('title') if it else f'موضوع تدريبي {item_idx + 1}'
            hrs = 2.0
            if it and 'page_from' in it and 'page_to' in it and it['page_from'] is not None:
                p_from, p_to = it['page_from'], it['page_to']
            elif it and 'p_from' in it and 'p_to' in it and it['p_from'] is not None:
                p_from, p_to = it['p_from'], it['p_to']
            else:
                p_from, p_to = calc_page_range(item_idx + 1, n_items, total_pages)
            q_cnt = 25

            ws.cell(current_row, 3, value=p_label).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(current_row, 4, value=lec_title).alignment = Alignment(horizontal='right', vertical='center')
            ws.cell(current_row, 5, value=hrs).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(current_row, 6, value=p_from).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(current_row, 7, value=p_to).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(current_row, 8, value=q_cnt).alignment = Alignment(horizontal='center', vertical='center')

            for col in range(1, 9):
                c = ws.cell(current_row, col)
                c.font = font_data
                c.border = border_cell

            ws.row_dimensions[current_row].height = 24
            current_row += 1
            item_idx += 1

        # Merge day column (Col B) for the 4 rows
        ws.merge_cells(start_row=day_start_row, start_column=2, end_row=day_start_row + 3, end_column=2)
        day_cell = ws.cell(day_start_row, 2, value=day_label)
        day_cell.font = font_header
        day_cell.alignment = Alignment(horizontal='center', vertical='center')
        for r_b in range(day_start_row, day_start_row + 4):
            ws.cell(r_b, 2).border = border_cell

        # Midterm Exam row after midterm_day
        if day_num == midterm_day:
            ws.row_dimensions[current_row].height = 26
            ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=8)
            m_cell = ws.cell(current_row, 2, value='امتحان منتصف الترم')
            m_cell.font = font_exam
            m_cell.alignment = Alignment(horizontal='center', vertical='center')
            for col in range(1, 9):
                c = ws.cell(current_row, col)
                c.fill = exam_fill
                c.border = border_cell
            current_row += 1

    # Final Exam row
    ws.row_dimensions[current_row].height = 28
    ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=4)
    f_cell = ws.cell(current_row, 2, value='امتحان ختامى الترم')
    f_cell.font = font_total
    f_cell.alignment = Alignment(horizontal='center', vertical='center')

    # Total Hours
    c_hrs = ws.cell(current_row, 5, value=total_hours)
    c_hrs.font = font_total
    c_hrs.alignment = Alignment(horizontal='center', vertical='center')

    # Total Questions
    c_qs = ws.cell(current_row, 8, value=total_questions)
    c_qs.font = font_total
    c_qs.alignment = Alignment(horizontal='center', vertical='center')

    for col in range(1, 9):
        c = ws.cell(current_row, col)
        c.fill = total_fill
        c.border = border_cell

    # Merge Term column A
    ws.merge_cells(start_row=term_start_row, start_column=1, end_row=current_row, end_column=1)
    a_term = ws.cell(term_start_row, 1, value='الترم الأول')
    a_term.font = font_header
    a_term.alignment = Alignment(horizontal='center', vertical='center')
    for r_a in range(term_start_row, current_row + 1):
        ws.cell(r_a, 1).border = border_cell

    # Column widths
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 52
    ws.column_dimensions['E'].width = 16
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 16


def build_workbook(
    syllabus_data: Optional[Dict[str, Any]],
    question_banks: List[Dict[str, Any]],
    equipment_name: str,
    training_program_mode: bool = False,
    sections_setup: Optional[List[Dict[str, Any]]] = None
) -> io.BytesIO:
    """
    Builds the complete multi-sheet Excel workbook:
    - If training_program_mode is True (Scenario 4):
      Creates institutional training program sheets with (ف1..ف4), midterm and final exams.
    - Otherwise (Scenarios 1-3):
      Sheet 1: الجدول الزمني وفهرس الدروس (Strict Single-Row: 1 Lecture : 2 Hours : 25 Questions)
    - Sheets 2..N: بنك [اسم القسم] - [اسم المعدة]
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Intellectual Property Metadata
    wb.properties.creator = "Eslam Abdelbadea"
    wb.properties.lastModifiedBy = "Eslam Abdelbadea"
    wb.properties.title = f"بنك الأسئلة والجدول الزمني - {equipment_name}"
    wb.properties.subject = "LMS Question Bank & Syllabus Guide"
    wb.properties.description = f"{BRAND_NOTICE} | {FOOTER_NOTICE}"

    font_banner = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_banner = PatternFill(start_color="0D233A", end_color="0D233A", fill_type="solid")
    border_cell = create_thin_border()

    all_items = syllabus_data.get("syllabus_items", []) if syllabus_data else []

    # ----------------------------------------------------
    # SCENARIO 4: INSTITUTIONAL TRAINING PROGRAM TEMPLATE
    # ----------------------------------------------------
    if training_program_mode and all_items:
        if sections_setup and len(sections_setup) > 0:
            cur_offset = 0
            for sec in sections_setup:
                sec_name = sec["name"]
                sec_lec = sec.get("lectures", len(all_items))
                sec_eq = sec.get("equipment_name", equipment_name)
                sec_items = all_items[cur_offset : cur_offset + sec_lec]
                cur_offset += sec_lec

                # Days and midterm configuration
                days = max(1, (len(sec_items) + 3) // 4)
                midterm = 4 if days == 7 else (7 if days == 14 else max(1, days // 2))

                add_training_program_sheet(
                    wb=wb,
                    sheet_title=f"برنامج محاضرات - {sec_name}",
                    section_label=sec_name,
                    specialty_name=sec_eq,
                    items=sec_items,
                    total_pages=33,
                    days_count=days,
                    midterm_day=midterm,
                    total_hours=len(sec_items) * 2.0,
                    total_questions=len(sec_items) * 25
                )
        else:
            # Single section fallback
            days = max(1, (len(all_items) + 3) // 4)
            add_training_program_sheet(
                wb=wb,
                sheet_title="برنامج التدريب المعتمد",
                section_label="التخصص",
                specialty_name=equipment_name,
                items=all_items,
                total_pages=33,
                days_count=days,
                midterm_day=max(1, days // 2),
                total_hours=len(all_items) * 2.0,
                total_questions=len(all_items) * 25
            )

    # ----------------------------------------------------
    # STANDARD SYLLABUS SHEET (Scenarios 1-3)
    # ----------------------------------------------------
    elif syllabus_data and syllabus_data.get("syllabus_items"):
        ws_syl = wb.create_sheet(title="الجدول الزمني وفهرس الدروس")
        apply_rtl_and_grid(ws_syl)

        # Row 1: Header Banner
        ws_syl.merge_cells("A1:F1")
        banner_cell = ws_syl["A1"]
        banner_cell.value = f"الجدول الزمني وفهرس الدروس المعتمد | المنظومة: {equipment_name}"
        banner_cell.font = font_banner
        banner_cell.fill = fill_banner
        banner_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_syl.row_dimensions[1].height = 28

        # Row 2: Standard 6 Headers
        headers_syl = [
            "م",
            "عنوان الدرس / المكون الفني",
            "عدد الساعات المخصصة",
            "عدد المحاضرات",
            "عدد الأسئلة المستهدفة",
            "المرجع المعتمد"
        ]
        ws_syl.row_dimensions[2].height = 28
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")

        for col_idx, h_text in enumerate(headers_syl, start=1):
            c = ws_syl.cell(row=2, column=col_idx, value=h_text)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border_cell

        # Data Rows: strictly single row per lesson (1 : 2 : 25)
        items = syllabus_data.get("syllabus_items", [])
        current_row = 3
        data_font = Font(name="Calibri", size=10)

        for it in items:
            ws_syl.row_dimensions[current_row].height = 24
            row_vals = [
                it.get("id", current_row - 2),
                it.get("lesson_name", ""),
                2,   # STRICTLY 2 hours
                1,   # STRICTLY 1 lecture
                25,  # STRICTLY 25 questions
                it.get("reference", "الدليل الفني المعتمد")
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                c = ws_syl.cell(row=current_row, column=col_idx, value=val)
                c.font = data_font
                c.border = border_cell
                
                # Align numbers to center, titles to right
                if col_idx in [1, 3, 4, 5]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    
            current_row += 1

        # Summary / Totals Row
        ws_syl.row_dimensions[current_row].height = 26
        total_font = Font(name="Calibri", size=10, bold=True, color="000000")
        total_fill = PatternFill(start_color=TOTAL_ROW_FILL, end_color=TOTAL_ROW_FILL, fill_type="solid")

        c_blank = ws_syl.cell(row=current_row, column=1, value="")
        c_blank.fill = total_fill
        c_blank.border = border_cell

        c_total_lbl = ws_syl.cell(row=current_row, column=2, value="الإجمالي العام")
        c_total_lbl.font = total_font
        c_total_lbl.fill = total_fill
        c_total_lbl.alignment = Alignment(horizontal="center", vertical="center")
        c_total_lbl.border = border_cell

        # Hours sum formula (Col C / 3) = N * 2
        c_hrs = ws_syl.cell(row=current_row, column=3, value=f"=SUM(C3:C{current_row-1})")
        c_hrs.font = total_font
        c_hrs.fill = total_fill
        c_hrs.alignment = Alignment(horizontal="center", vertical="center")
        c_hrs.border = border_cell

        # Lectures sum formula (Col D / 4) = N * 1
        c_lec = ws_syl.cell(row=current_row, column=4, value=f"=SUM(D3:D{current_row-1})")
        c_lec.font = total_font
        c_lec.fill = total_fill
        c_lec.alignment = Alignment(horizontal="center", vertical="center")
        c_lec.border = border_cell

        # Questions sum formula (Col E / 5) = N * 25
        c_q = ws_syl.cell(row=current_row, column=5, value=f"=SUM(E3:E{current_row-1})")
        c_q.font = total_font
        c_q.fill = total_fill
        c_q.alignment = Alignment(horizontal="center", vertical="center")
        c_q.border = border_cell

        c_ref_total = ws_syl.cell(row=current_row, column=6, value="")
        c_ref_total.fill = total_fill
        c_ref_total.border = border_cell

        ws_syl.column_dimensions["A"].width = 6
        ws_syl.column_dimensions["B"].width = 46
        ws_syl.column_dimensions["C"].width = 20
        ws_syl.column_dimensions["D"].width = 16
        ws_syl.column_dimensions["E"].width = 22
        ws_syl.column_dimensions["F"].width = 30

    # ----------------------------------------------------
    # QUESTION BANK SHEETS: بنك [اسم القسم] - [اسم المعدة]
    # ----------------------------------------------------
    qb_headers = [
        "م",
        "نص السؤال",
        "الشرح / التفسير",
        "الخيار أ",
        "الخيار ب",
        "الخيار ج",
        "الخيار د",
        "الإجابة الصحيحة",
        "مستوى الصعوبة",
        "الدرس",
        "نوع السؤال"
    ]

    for qb in question_banks:
        sec_name = qb.get("section_name", "عام")
        sec_eq = qb.get("equipment_name") or equipment_name
        sheet_title = f"بنك {sec_name} - {sec_eq}"
        if len(sheet_title) > 31:
            sheet_title = f"بنك {sec_name}"[:31]

        ws_qb = wb.create_sheet(title=sheet_title)
        apply_rtl_and_grid(ws_qb)

        # Row 1: Header Banner
        ws_qb.merge_cells("A1:K1")
        b_cell = ws_qb["A1"]
        b_cell.value = f"بنك الأسئلة المعتمد | {sec_name} - {sec_eq}"
        b_cell.font = font_banner
        b_cell.fill = fill_banner
        b_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_qb.row_dimensions[1].height = 28

        # Row 2: Headers
        ws_qb.row_dimensions[2].height = 28
        h_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        h_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")

        for col_idx, h_name in enumerate(qb_headers, start=1):
            cell = ws_qb.cell(row=2, column=col_idx, value=h_name)
            cell.font = h_font
            cell.fill = h_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_cell

        # Data Rows
        questions = qb.get("questions", [])
        q_font = Font(name="Calibri", size=10)
        cur_r = 3

        for q in questions:
            ws_qb.row_dimensions[cur_r].height = 36
            diff = str(q.get("difficulty", "متوسط")).strip()
            diff_hex = DIFFICULTY_COLORS.get(diff, "FFF2CC")
            diff_fill = PatternFill(start_color=diff_hex, end_color=diff_hex, fill_type="solid")

            row_data = [
                q.get("id", cur_r - 2),
                q.get("question_text", ""),
                q.get("explanation", ""),
                q.get("option_a", ""),
                q.get("option_b", ""),
                q.get("option_c", ""),
                q.get("option_d", ""),
                q.get("correct_answer", ""),
                diff,
                q.get("lesson", ""),
                q.get("question_type", "")
            ]

            for c_idx, cell_value in enumerate(row_data, start=1):
                c = ws_qb.cell(row=cur_r, column=c_idx, value=cell_value)
                c.font = q_font
                c.border = border_cell

                if c_idx == 9:
                    c.fill = diff_fill
                    c.font = Font(name="Calibri", size=10, bold=True)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in [1, 11]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in [4, 5, 6, 7]:
                    c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
                elif c_idx == 8:
                    c.font = Font(name="Calibri", size=10, bold=True, color="004D20")
                    c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
                else:
                    c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

            cur_r += 1

        cur_r += 1
        ws_qb.merge_cells(start_row=cur_r, start_column=1, end_row=cur_r, end_column=11)
        foot = ws_qb.cell(row=cur_r, column=1, value=f"✨ {FOOTER_NOTICE}")
        foot.font = Font(name="Calibri", size=9, italic=True, color="595959")
        foot.alignment = Alignment(horizontal="center", vertical="center")

        ws_qb.column_dimensions["A"].width = 6
        ws_qb.column_dimensions["B"].width = 46
        ws_qb.column_dimensions["C"].width = 30
        ws_qb.column_dimensions["D"].width = 22
        ws_qb.column_dimensions["E"].width = 22
        ws_qb.column_dimensions["F"].width = 22
        ws_qb.column_dimensions["G"].width = 22
        ws_qb.column_dimensions["H"].width = 25
        ws_qb.column_dimensions["I"].width = 15
        ws_qb.column_dimensions["J"].width = 26
        ws_qb.column_dimensions["K"].width = 16

    if len(wb.sheetnames) == 0:
        ws_empty = wb.create_sheet(title="تنبيه")
        ws_empty["A1"] = "لم يتم توليد أي محتوى"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
