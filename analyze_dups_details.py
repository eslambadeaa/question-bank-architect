import openpyxl, sys, collections
sys.stdout.reconfigure(encoding='utf-8')

for tag, fname in [('متوسط', 'بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'), ('نهائي', 'بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx')]:
    fpath = r'C:\Users\MaximuM-Tech\Downloads\\' + fname
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb['Questions']
    seen = {}
    dups = collections.defaultdict(list)
    for r in range(2, ws.max_row + 1):
        txt = str(ws.cell(r, 2).value).strip() if ws.cell(r, 2).value else ''
        les = ws.cell(r, 10).value
        qtype = ws.cell(r, 1).value
        diff = ws.cell(r, 4).value
        ans = ws.cell(r, 5).value
        if txt not in seen:
            seen[txt] = (r, les)
        else:
            dups[les].append((r, qtype, diff, ans, txt))
    
    print(f'=== {tag}: {sum(len(v) for v in dups.values())} duplicates across {len(dups)} lessons ===')
    for les, items in dups.items():
        diff_counts = collections.Counter(x[2] for x in items)
        type_counts = collections.Counter(x[1] for x in items)
        print(f'Lesson: \"{les}\" (Need {len(items)}: Types={dict(type_counts)}, Diffs={dict(diff_counts)})')
        for it in items[:2]:
            print(f'   Row {it[0]}: [{it[1]} | {it[2]}] {it[4][:60]}...')
