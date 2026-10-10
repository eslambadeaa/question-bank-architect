import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'C:\Users\MaximuM-Tech\Downloads\برنامج تدريب تخصص الضبع الاسود + بنوك_محدث_LMS.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

def load_sheet(sheet_name):
    ws = wb[sheet_name]
    topics = []
    for r in range(5, ws.max_row + 1):
        lec = ws.cell(r, 3).value
        name = ws.cell(r, 4).value
        ps = ws.cell(r, 7).value
        pe = ws.cell(r, 8).value
        if lec in ['ف1', 'ف3']:
            topics.append({
                'num': len(topics) + 1,
                'name': name.strip() if name else '',
                'ps': int(ps) if ps else 0,
                'pe': int(pe) if pe else 0
            })
    return topics

hyena = load_sheet('برنامج تدريب - القسم المتوسط')
igla = load_sheet('برنامج تدريب - القسم النهائي')

print("Hyena loaded:", len(hyena))
print("Igla loaded:", len(igla))
