import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_path = r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx'
wb = openpyxl.load_workbook(f_path)
ws_q = wb['Questions']

rows_to_check = [953, 955, 956, 958, 959, 960, 962, 964, 1002] + list(range(1052, 1065))

for r in rows_to_check:
    print(f"Row {r}:")
    print(f"  Q: {ws_q.cell(r, 2).value}")
    print(f"  Exp: {ws_q.cell(r, 3).value}")
    print(f"  Ans: {ws_q.cell(r, 5).value}")
    print(f"  A: {ws_q.cell(r, 6).value}")
    print(f"  B: {ws_q.cell(r, 7).value}")
    print(f"  C: {ws_q.cell(r, 8).value}")
    print(f"  D: {ws_q.cell(r, 9).value}")
