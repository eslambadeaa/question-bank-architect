import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'C:\Users\MaximuM-Tech\Downloads\الكامل مرجع الإيجلا.docx')

print("Total paragraphs:", len(doc.paragraphs))
headings = []
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if p.style.name.startswith('Heading') or (len(txt) > 3 and len(txt) < 80 and ('الفصل' in txt or 'الباب' in txt or 'المبحث' in txt or 'القسم' in txt or 'أولاً' in txt or 'ثانياً' in txt)):
        headings.append((idx, txt))

print(f"Found {len(headings)} heading-like lines. First 30:")
for h in headings[:30]:
    print(f"  P{h[0]}: {h[1]}")
