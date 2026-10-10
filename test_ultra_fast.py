import openpyxl, time, os, sys
import win32com.client

sys.stdout.reconfigure(encoding='utf-8')

template_path = r'C:\Users\MaximuM-Tech\Downloads\نموذج تسجيل الاسئلة (1).xlsx'
source_med = r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'
test_out = r'C:\Users\MaximuM-Tech\Downloads\test_ultra_fast_med.xlsx'

wb_src = openpyxl.load_workbook(source_med, data_only=True)
ws_src = wb_src['Questions']
data_rows = list(ws_src.iter_rows(values_only=True))[1:] # 1300 rows

# Load clean template
wb_tmpl = openpyxl.load_workbook(template_path)
ws_q = wb_tmpl['Questions']

# Populate data in template without any custom styling/borders/wrapText
for r_idx, r_data in enumerate(data_rows, start=2):
    for c_idx, val in enumerate(r_data, start=1):
        ws_q.cell(r_idx, c_idx, val)

# Set exact data validation
from openpyxl.worksheet.datavalidation import DataValidation
ws_q.data_validations.dataValidation.clear()

dv_qtype = DataValidation(type="list", formula1="L_1", allow_blank=True)
ws_q.add_data_validation(dv_qtype)
dv_qtype.add(f"A2:A{len(data_rows)+1}")

dv_diff = DataValidation(type="list", formula1="L_2", allow_blank=True)
ws_q.add_data_validation(dv_diff)
dv_diff.add(f"D2:D{len(data_rows)+1}")

dv_ans = DataValidation(type="list", formula1="L_3", allow_blank=True)
ws_q.add_data_validation(dv_ans)
dv_ans.add(f"E2:E{len(data_rows)+1}")

# Populate Lookups Col D with lessons
ws_lookups = wb_tmpl['Lookups']
lessons_ordered = []
for r in data_rows:
    les = r[9]
    if les and les not in lessons_ordered:
        lessons_ordered.append(les)

for idx, les in enumerate(lessons_ordered, start=2):
    ws_lookups.cell(idx, 4, les)

dv_les = DataValidation(type="list", formula1="L_4", allow_blank=True)
ws_q.add_data_validation(dv_les)
dv_les.add(f"J2:J{len(data_rows)+1}")

from openpyxl.workbook.defined_name import DefinedName
wb_tmpl.defined_names['L_4'] = DefinedName('L_4', attr_text=f"Lookups!$D$2:$D${len(lessons_ordered)+1}")

wb_tmpl.save(test_out)
print(f"Saved {test_out}, size: {os.path.getsize(test_out)} bytes")

# Now resave with native Excel COM
excel = win32com.client.Dispatch("Excel.Application")
excel.Visible = False
excel.DisplayAlerts = False

t0 = time.time()
wb = excel.Workbooks.Open(test_out)
t1 = time.time()
print(f"Open time BEFORE native save: {t1 - t0:.2f} s")

wb.Save()
wb.Close(False)

t0 = time.time()
wb = excel.Workbooks.Open(test_out)
t1 = time.time()
print(f"Open time AFTER native save: {t1 - t0:.2f} s")
wb.Close(False)

excel.Quit()
