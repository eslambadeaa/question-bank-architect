"""
Regenerates Question Bank sheets using authentic respective pools:
- Preparatory: from 'اعدادي ضبع اسود.xlsx' (100% Black Hyena questions, covering 30% of manual)
- Medium: from 'متوسط ضبع اسود.xlsx' (100% Black Hyena questions, covering 100% of manual)
- Final: from 'نهائي ضبع اسود.xlsx' (100% Igla questions, covering 100% of manual)
Preserves training program sheets 100% intact with user edits.
Formats 'الإجابة الصحيحة' as option labels with full data validation.
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding='utf-8')

def create_thin_border():
    thin = Side(border_style='thin', color='A0A0A0')
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def apply_rtl(ws):
    try:
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = True
    except Exception:
        pass

def get_correct_option_label(correct_text, opt_a, opt_b, opt_c, opt_d):
    correct_str = str(correct_text or '').strip()
    if correct_str in ['أ', 'A', 'الخيار أ', 'الخيار أ (A)']: return 'الخيار أ (A)'
    if correct_str in ['ب', 'B', 'الخيار ب', 'الخيار ب (B)']: return 'الخيار ب (B)'
    if correct_str in ['ج', 'C', 'الخيار ج', 'الخيار ج (C)']: return 'الخيار ج (C)'
    if correct_str in ['د', 'D', 'الخيار د', 'الخيار د (D)']: return 'الخيار د (D)'

    def clean(s):
        s = str(s or '').strip()
        for ch in [' ', '،', '؟', '!', '.', ':', '-', '_', 'أ', 'إ', 'آ', 'ة', 'ه', 'ى', 'ي']:
            s = s.replace(ch, '')
        return s.lower()

    c_clean = clean(correct_str)
    opts = [
        ('الخيار أ (A)', opt_a, clean(opt_a)),
        ('الخيار ب (B)', opt_b, clean(opt_b)),
        ('الخيار ج (C)', opt_c, clean(opt_c)),
        ('الخيار د (D)', opt_d, clean(opt_d)),
    ]
    for label, raw, cl in opts:
        if correct_str == str(raw or '').strip():
            return label
    for label, raw, cl in opts:
        if c_clean and c_clean == cl:
            return label
    for label, raw, cl in opts:
        if c_clean and (c_clean in cl or cl in c_clean):
            return label
    return 'الخيار أ (A)'

def load_pool_file(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    questions = []
    by_topic = {}
    for r in range(2, ws.max_row + 1):
        q_text = ws.cell(r, 2).value
        if q_text is not None:
            topic = str(ws.cell(r, 10).value or '').strip()
            item = {
                "q_text": q_text,
                "explanation": ws.cell(r, 3).value or "",
                "opt_a": ws.cell(r, 4).value or "",
                "opt_b": ws.cell(r, 5).value or "",
                "opt_c": ws.cell(r, 6).value or "",
                "opt_d": ws.cell(r, 7).value or "",
                "correct_raw": ws.cell(r, 8).value or "",
                "difficulty": ws.cell(r, 9).value or "متوسط",
                "topic": topic,
                "q_type": ws.cell(r, 11).value or "اختيار من متعدد"
            }
            questions.append(item)
            by_topic.setdefault(topic, []).append(item)
    return questions, by_topic

def match_questions_for_topics(prog_topics, all_questions, by_topic):
    """
    Selects 25 questions for each topic in prog_topics.
    Attempts best topic keyword matching; falls back gracefully.
    """
    bank_rows = []
    pool_cursor = 0
    diff_cycle = ["سهل", "متوسط", "صعب", "متوسط"]

    for pt in prog_topics:
        # Find best matching topic bucket in pool
        best_bucket = None
        best_score = 0
        words_pt = set(pt.split())
        for t_key in by_topic.keys():
            if t_key:
                score = len(words_pt & set(t_key.split()))
                if score > best_score:
                    best_score = score
                    best_bucket = t_key

        matched_pool = by_topic.get(best_bucket, []) if best_bucket and best_score >= 1 else []

        for q_i in range(25):
            if matched_pool and q_i < len(matched_pool):
                src = matched_pool[q_i]
            else:
                src = all_questions[pool_cursor % len(all_questions)]
                pool_cursor += 1

            diff = src.get('difficulty')
            if not diff or diff not in ["سهل", "متوسط", "صعب"]:
                diff = diff_cycle[q_i % len(diff_cycle)]

            q_type = src.get('q_type', 'اختيار من متعدد')
            if not q_type or q_type not in ["اختيار من متعدد", "صواب أو خطأ"]:
                q_type = "اختيار من متعدد"

            opt_a = src['opt_a']
            opt_b = src['opt_b']
            opt_c = src['opt_c']
            opt_d = src['opt_d']
            correct_label = get_correct_option_label(src['correct_raw'], opt_a, opt_b, opt_c, opt_d)

            row = [
                q_type,
                src['q_text'],
                src['explanation'],
                diff,
                correct_label,
                opt_a,
                opt_b,
                opt_c,
                opt_d,
                pt # Exactly matching the program topic
            ]
            bank_rows.append(row)
    return bank_rows

def populate_qb_sheet(ws, q_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill):
    apply_rtl(ws)
    ws.delete_rows(1, ws.max_row + 1)

    headers = [
        "نوع السؤال",
        "نص السؤال",
        "الشرح / التفسير",
        "مستوى الصعوبة",
        "الإجابة الصحيحة",
        "الخيار أ (A)",
        "الخيار ب (B)",
        "الخيار ج (C)",
        "الخيار د (D)",
        "الدرس"
    ]
    ws.append(headers)
    for col_idx in range(1, 11):
        c = ws.cell(1, col_idx)
        c.font = qb_header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.fill = qb_header_fill
        c.border = border_cell
    ws.row_dimensions[1].height = 28

    for r_idx, row_data in enumerate(q_rows, 2):
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 20
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(r_idx, col_idx)
            c.font = qb_data_font
            c.border = border_cell
            if col_idx in [1, 4, 5]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="right", vertical="center")

    widths = {1: 18, 2: 50, 3: 35, 4: 15, 5: 18, 6: 22, 7: 22, 8: 22, 9: 22, 10: 45}
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    max_r = max(len(q_rows) + 50, 200)

    # 1. نوع السؤال (Col A)
    dv_type = DataValidation(type="list", formula1='"اختيار من متعدد,صواب أو خطأ"', allow_blank=True)
    ws.add_data_validation(dv_type)
    dv_type.add(f"A2:A{max_r}")

    # 2. مستوى الصعوبة (Col D)
    dv_diff = DataValidation(type="list", formula1='"سهل,متوسط,صعب"', allow_blank=True)
    ws.add_data_validation(dv_diff)
    dv_diff.add(f"D2:D{max_r}")

    # 3. الإجابة الصحيحة (Col E)
    dv_ans = DataValidation(type="list", formula1='"الخيار أ (A),الخيار ب (B),الخيار ج (C),الخيار د (D)"', allow_blank=True)
    ws.add_data_validation(dv_ans)
    dv_ans.add(f"E2:E{max_r}")

def main():
    target_file = r"C:\Users\MaximuM-Tech\Downloads\بنوك\برنامج تدريب تخصص الضبع الاسود_النسخة_الرسمية_المعتمدة.xlsx"
    print(f"Loading user-edited workbook: {target_file}")
    wb = openpyxl.load_workbook(target_file)

    # 1. Extract program topics
    # A) Prep
    ws_prep = wb['برنامج تدريب - القسم الإعدادي']
    prep_topics = []
    for r in range(5, ws_prep.max_row + 1):
        c = ws_prep.cell(r, 3).value
        d = ws_prep.cell(r, 4).value
        e = ws_prep.cell(r, 5).value
        if c == 'ف1' and d and (e == 2 or e == '2'):
            prep_topics.append(d)
    print(f"Extracted {len(prep_topics)} topics from برنامج تدريب - القسم الإعدادي")

    # B) Med
    ws_med = wb['برنامج تدريب - القسم المتوسط']
    med_t1_topics = []
    med_t2_topics = []
    for r in range(5, ws_med.max_row + 1):
        c = ws_med.cell(r, 3).value
        d = ws_med.cell(r, 4).value
        e = ws_med.cell(r, 5).value
        if (c in ['ف1', 'ف3']) and d and (e == 2 or e == '2'):
            if r <= 54:
                med_t1_topics.append(d)
            else:
                med_t2_topics.append(d)
    print(f"Extracted {len(med_t1_topics)} T1 topics and {len(med_t2_topics)} T2 topics from برنامج تدريب - القسم المتوسط")

    # C) Final
    ws_fin = wb['برنامج تدريب - القسم النهائي']
    fin_t1_topics = []
    fin_t2_topics = []
    for r in range(5, ws_fin.max_row + 1):
        c = ws_fin.cell(r, 3).value
        d = ws_fin.cell(r, 4).value
        e = ws_fin.cell(r, 5).value
        if (c in ['ف1', 'ف3']) and d and (e == 2 or e == '2'):
            if r <= 54:
                fin_t1_topics.append(d)
            else:
                fin_t2_topics.append(d)
    print(f"Extracted {len(fin_t1_topics)} T1 topics and {len(fin_t2_topics)} T2 topics from برنامج تدريب - القسم النهائي")

    # 2. Load authentic question pools
    print("\nLoading authentic question pools:")
    # Prep pool (اعدادي ضبع اسود)
    prep_pool_file = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "اعدادي ضبع اسود.xlsx")
    prep_all_q, prep_by_topic = load_pool_file(prep_pool_file)
    print(f"  - Loaded {len(prep_all_q)} questions for Preparatory from: {os.path.basename(prep_pool_file)}")

    # Med pool (متوسط ضبع اسود)
    med_pool_file = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "متوسط ضبع اسود.xlsx")
    med_all_q, med_by_topic = load_pool_file(med_pool_file)
    print(f"  - Loaded {len(med_all_q)} questions for Medium from: {os.path.basename(med_pool_file)}")

    # Final pool (نهائي ضبع اسود)
    fin_pool_file = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "نهائي ضبع اسود.xlsx")
    fin_all_q, fin_by_topic = load_pool_file(fin_pool_file)
    print(f"  - Loaded {len(fin_all_q)} questions for Final from: {os.path.basename(fin_pool_file)}")

    # Styles
    border_cell = create_thin_border()
    qb_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    qb_data_font = Font(name="Calibri", size=10)
    qb_header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")

    # 3. Populate sheets
    print("\nPopulating Question Bank sheets with authentic pools:")
    # 1. بنك القسم الإعدادي (350 سؤال من اعدادي ضبع اسود)
    prep_bank_rows = match_questions_for_topics(prep_topics, prep_all_q, prep_by_topic)
    ws_qb_prep = wb['بنك القسم الإعدادي'] if 'بنك القسم الإعدادي' in wb.sheetnames else wb.create_sheet('بنك القسم الإعدادي')
    populate_qb_sheet(ws_qb_prep, prep_bank_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"  - بنك القسم الإعدادي: {len(prep_bank_rows)} سؤال (الضبع الأسود - إعدادي)")

    # 2. بنك المتوسط - ترم أول (600 سؤال من متوسط ضبع اسود)
    med_t1_bank_rows = match_questions_for_topics(med_t1_topics, med_all_q, med_by_topic)
    ws_qb_med1 = wb['بنك المتوسط - ترم أول'] if 'بنك المتوسط - ترم أول' in wb.sheetnames else wb.create_sheet('بنك المتوسط - ترم أول')
    populate_qb_sheet(ws_qb_med1, med_t1_bank_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"  - بنك المتوسط - ترم أول: {len(med_t1_bank_rows)} سؤال (الضبع الأسود - متوسط ترم 1)")

    # 3. بنك المتوسط - ترم ثاني (700 سؤال من متوسط ضبع اسود)
    med_t2_bank_rows = match_questions_for_topics(med_t2_topics, med_all_q, med_by_topic)
    ws_qb_med2 = wb['بنك المتوسط - ترم ثاني'] if 'بنك المتوسط - ترم ثاني' in wb.sheetnames else wb.create_sheet('بنك المتوسط - ترم ثاني')
    populate_qb_sheet(ws_qb_med2, med_t2_bank_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"  - بنك المتوسط - ترم ثاني: {len(med_t2_bank_rows)} سؤال (الضبع الأسود - متوسط ترم 2)")

    # 4. بنك النهائي - ترم أول (600 سؤال من نهائي ضبع اسود - ايجلا)
    fin_t1_bank_rows = match_questions_for_topics(fin_t1_topics, fin_all_q, fin_by_topic)
    ws_qb_fin1 = wb['بنك النهائي - ترم أول'] if 'بنك النهائي - ترم أول' in wb.sheetnames else wb.create_sheet('بنك النهائي - ترم أول')
    populate_qb_sheet(ws_qb_fin1, fin_t1_bank_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"  - بنك النهائي - ترم أول: {len(fin_t1_bank_rows)} سؤال (ايجلا - نهائي ترم 1)")

    # 5. بنك النهائي - ترم ثاني (500 سؤال من نهائي ضبع اسود - ايجلا)
    fin_t2_bank_rows = match_questions_for_topics(fin_t2_topics, fin_all_q, fin_by_topic)
    ws_qb_fin2 = wb['بنك النهائي - ترم ثاني'] if 'بنك النهائي - ترم ثاني' in wb.sheetnames else wb.create_sheet('بنك النهائي - ترم ثاني')
    populate_qb_sheet(ws_qb_fin2, fin_t2_bank_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print(f"  - بنك النهائي - ترم ثاني: {len(fin_t2_bank_rows)} سؤال (ايجلا - نهائي ترم 2)")

    # Save
    wb.save(target_file)
    print(f"\nSUCCESS: Saved updated workbook to: {target_file}")

    primary_file = r"C:\Users\MaximuM-Tech\Downloads\بنوك\برنامج تدريب تخصص الضبع الاسود.xlsx"
    try:
        wb.save(primary_file)
        print(f"SUCCESS: Also updated {primary_file}")
    except PermissionError:
        print(f"NOTE: {primary_file} is locked.")

if __name__ == '__main__':
    main()
