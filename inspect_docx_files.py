import docx, sys, glob
sys.stdout.reconfigure(encoding='utf-8')

for f in glob.glob(r'C:\Users\MaximuM-Tech\Downloads\*.docx'):
    try:
        doc = docx.Document(f)
        texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables_text = []
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    if c.text.strip():
                        tables_text.append(c.text.strip())
        print(f"File: {f} | Paragraphs: {len(texts)} | Table cells: {len(tables_text)}")
        if texts:
            print("  First 3 paragraphs:", texts[:3])
    except Exception as e:
        print(f"File: {f} | Error: {e}")
