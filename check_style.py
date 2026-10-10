import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

for tag, fname in [('Medium', 'بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'), ('Final', 'بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx')]:
    wb = openpyxl.load_workbook(r'C:\Users\MaximuM-Tech\Downloads\\' + fname, data_only=True)
    ws = wb['Questions']
    print(f"=== SAMPLE QUESTIONS IN {tag} ===")
    for r in range(2, 20):
        print(f"Row {r}: [{ws.cell(r,1).value} | {ws.cell(r,4).value} | ans={ws.cell(r,5).value}]")
        print(f"  Q: {ws.cell(r,2).value}")
        if ws.cell(r,6).value:
            print(f"  A: {ws.cell(r,6).value} | B: {ws.cell(r,7).value} | C: {ws.cell(r,8).value} | D: {ws.cell(r,9).value}")
