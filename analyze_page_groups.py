import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'C:\Users\MaximuM-Tech\Downloads\برنامج تدريب تخصص الضبع الاسود + بنوك_محدث_LMS.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

for sheet_name in ['برنامج تدريب - القسم المتوسط', 'برنامج تدريب - القسم النهائي']:
    ws = wb[sheet_name]
    print(f"\n=================== {sheet_name} ===================")
    topics = []
    for r in range(5, ws.max_row + 1):
        lec = ws.cell(r, 3).value
        name = ws.cell(r, 4).value
        ps = ws.cell(r, 7).value
        pe = ws.cell(r, 8).value
        if lec in ['ف1', 'ف3']:
            topics.append({
                'num': len(topics) + 1,
                'name': name.strip() if name else '',
                'ps': int(ps) if ps else 0,
                'pe': int(pe) if pe else 0
            })
    
    # Group by (ps, pe)
    groups = {}
    for t in topics:
        key = (t['ps'], t['pe'])
        if key not in groups:
            groups[key] = []
        groups[key].append(t)
        
    print(f"Total topics: {len(topics)}, Unique (ps, pe) ranges: {len(groups)}")
    for k in sorted(groups.keys()):
        g_topics = groups[k]
        nums = [str(x['num']) for x in g_topics]
        names = " // ".join([x['name'] for x in g_topics])
        print(f"Pages {k[0]:2d} -> {k[1]:2d} (Topics {', '.join(nums)}): {names}")
