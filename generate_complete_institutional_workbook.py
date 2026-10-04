"""
Complete Institutional Generator for Training Programs & Question Banks
Specialty: الضبع الاسود
References:
- الإعدادي والمتوسط: مرجع الضبع الاسود.docx (31 صفحة)
- النهائي: مرجع الايجلا.docx (33 صفحة)
Quotas:
- الإعدادي: 56h total (28 lectures: 14 Th + 14 Pr | 700 questions)
- المتوسط ترم أول: 96h total (48 lectures: 24 Th + 24 Pr | 1200 questions)
- المتوسط ترم ثاني: 112h total (56 lectures: 28 Th + 28 Pr | 1400 questions)
- النهائي ترم أول: 96h total (48 lectures: 24 Th + 24 Pr | 1200 questions)
- النهائي ترم ثاني: 80h total (40 lectures: 20 Th + 20 Pr | 1000 questions)
"""

import os
import sys
from copy import copy
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from build_full_institutional_system import (
    PREP_SCHEDULE, MED_T1_SCHEDULE, MED_T2_SCHEDULE,
    FINAL_T1_SCHEDULE, FINAL_T2_SCHEDULE, ARABIC_DAYS,
    create_thin_border, apply_rtl
)

sys.stdout.reconfigure(encoding='utf-8')

def add_training_program_term(
    ws,
    term_name,
    schedule,
    days_count,
    midterm_day,
    th_hours,
    pr_hours,
    tot_hours,
    tot_questions,
    start_row=3,
    border_cell=None,
    font_header=None,
    font_data=None,
    font_exam=None,
    font_total=None,
    header_fill=None,
    exam_fill=None,
    total_fill=None
):
    """
    Renders one full term block in a training program sheet starting at start_row.
    Returns the next available row index.
    """
    r_hdr1 = start_row
    r_hdr2 = start_row + 1

    # Merge headers
    ws.merge_cells(start_row=r_hdr1, start_column=1, end_row=r_hdr2, end_column=1)
    ws.merge_cells(start_row=r_hdr1, start_column=2, end_row=r_hdr2, end_column=2)
    ws.merge_cells(start_row=r_hdr1, start_column=3, end_row=r_hdr2, end_column=3)
    ws.merge_cells(start_row=r_hdr1, start_column=4, end_row=r_hdr1, end_column=5)
    ws.merge_cells(start_row=r_hdr1, start_column=6, end_row=r_hdr2, end_column=6)
    ws.merge_cells(start_row=r_hdr1, start_column=7, end_row=r_hdr1, end_column=8)
    ws.merge_cells(start_row=r_hdr1, start_column=9, end_row=r_hdr2, end_column=9)

    headers = {
        (r_hdr1, 1): 'الترم',
        (r_hdr1, 2): 'اليوم',
        (r_hdr1, 3): 'المحاضرة',
        (r_hdr1, 4): 'اسم الموضوع',
        (r_hdr1, 6): 'عدد الساعات',
        (r_hdr1, 7): 'الصفحة في المرجع الموحد',
        (r_hdr1, 9): 'عدد الاسئلة'
    }
    for (r, c), text in headers.items():
        cell = ws.cell(r, c, value=text)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border_cell

    sub_headers = {
        (r_hdr2, 4): 'نظري',
        (r_hdr2, 5): 'عملي',
        (r_hdr2, 7): 'من',
        (r_hdr2, 8): 'الي'
    }
    for (r, c), text in sub_headers.items():
        cell = ws.cell(r, c, value=text)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border_cell

    for r in [r_hdr1, r_hdr2]:
        ws.row_dimensions[r].height = 26
        for col in range(1, 10):
            c = ws.cell(r, col)
            c.border = border_cell
            if c.fill.fill_type is None:
                c.fill = header_fill

    current_row = r_hdr2 + 1
    term_start_row = current_row
    item_idx = 0

    for day_num in range(1, days_count + 1):
        day_label = ARABIC_DAYS[day_num - 1] if day_num <= len(ARABIC_DAYS) else f'اليوم {day_num}'
        day_start_row = current_row

        for period_idx in range(1, 5):
            it = schedule[item_idx] if item_idx < len(schedule) else None
            p_label = f'ف{period_idx}'
            th_val = it['theory'] if it else '-'
            pr_val = it['prac'] if it else '-'
            hrs = 2.0
            p_from = it['p_from'] if it else 1
            p_to = it['p_to'] if it else 1
            q_cnt = 25

            ws.cell(current_row, 3, value=p_label).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(current_row, 4, value=th_val).alignment = Alignment(horizontal='right' if th_val != '-' else 'center', vertical='center')
            ws.cell(current_row, 5, value=pr_val).alignment = Alignment(horizontal='right' if pr_val != '-' else 'center', vertical='center')
            ws.cell(current_row, 6, value=hrs).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(current_row, 7, value=p_from).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(current_row, 8, value=p_to).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(current_row, 9, value=q_cnt).alignment = Alignment(horizontal='center', vertical='center')

            for col in range(1, 10):
                c = ws.cell(current_row, col)
                c.font = font_data
                c.border = border_cell

            ws.row_dimensions[current_row].height = 25
            current_row += 1
            item_idx += 1

        # Merge day column (Col B)
        ws.merge_cells(start_row=day_start_row, start_column=2, end_row=day_start_row + 3, end_column=2)
        day_cell = ws.cell(day_start_row, 2, value=day_label)
        day_cell.font = font_header
        day_cell.alignment = Alignment(horizontal='center', vertical='center')
        for r_b in range(day_start_row, day_start_row + 4):
            ws.cell(r_b, 2).border = border_cell

        # Midterm Exam row after midterm_day
        if day_num == midterm_day:
            ws.row_dimensions[current_row].height = 26
            ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=9)
            m_cell = ws.cell(current_row, 2, value='امتحان منتصف الترم')
            m_cell.font = font_exam
            m_cell.alignment = Alignment(horizontal='center', vertical='center')
            for col in range(1, 10):
                c = ws.cell(current_row, col)
                c.fill = exam_fill
                c.border = border_cell
            current_row += 1

    # Final Exam row for this term
    ws.row_dimensions[current_row].height = 28
    ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=3)
    f_cell = ws.cell(current_row, 2, value='امتحان ختامى الترم')
    f_cell.font = font_total
    f_cell.alignment = Alignment(horizontal='center', vertical='center')

    # Total Theory Hours in Col D
    c_th = ws.cell(current_row, 4, value=th_hours)
    c_th.font = font_total
    c_th.alignment = Alignment(horizontal='center', vertical='center')

    # Total Practical Hours in Col E
    c_pr = ws.cell(current_row, 5, value=pr_hours)
    c_pr.font = font_total
    c_pr.alignment = Alignment(horizontal='center', vertical='center')

    # Total Hours in Col F
    c_hrs = ws.cell(current_row, 6, value=tot_hours)
    c_hrs.font = font_total
    c_hrs.alignment = Alignment(horizontal='center', vertical='center')

    # Page range merged dash
    ws.merge_cells(start_row=current_row, start_column=7, end_row=current_row, end_column=8)
    c_dash = ws.cell(current_row, 7, value='-')
    c_dash.font = font_total
    c_dash.alignment = Alignment(horizontal='center', vertical='center')

    # Total Questions in Col I
    c_qs = ws.cell(current_row, 9, value=tot_questions)
    c_qs.font = font_total
    c_qs.alignment = Alignment(horizontal='center', vertical='center')

    for col in range(1, 10):
        c = ws.cell(current_row, col)
        c.fill = total_fill
        c.border = border_cell

    # Merge Term column A
    ws.merge_cells(start_row=term_start_row, start_column=1, end_row=current_row, end_column=1)
    a_term = ws.cell(term_start_row, 1, value=term_name)
    a_term.font = font_header
    a_term.alignment = Alignment(horizontal='center', vertical='center')
    for r_a in range(term_start_row, current_row + 1):
        ws.cell(r_a, 1).border = border_cell

    return current_row + 1


def build_question_bank_sheet(ws, title, questions_list, border_cell, header_font, data_font, header_fill):
    """
    Renders an institutional Question Bank sheet.
    """
    apply_rtl(ws)
    headers = [
        "م", "نص السؤال", "الشرح / التفسير", "الخيار أ", "الخيار ب",
        "الخيار ج", "الخيار د", "الإجابة الصحيحة", "مستوى الصعوبة", "الدرس", "نوع السؤال"
    ]
    ws.row_dimensions[1].height = 28
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(1, col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border_cell

    col_widths = {
        'A': 8, 'B': 45, 'C': 40, 'D': 25, 'E': 25,
        'F': 25, 'G': 25, 'H': 25, 'I': 14, 'J': 35, 'K': 16
    }
    for col_letter, w in col_widths.items():
        ws.column_dimensions[col_letter].width = w

    for r_idx, q in enumerate(questions_list, 2):
        ws.row_dimensions[r_idx].height = 24
        # q is list of 11 values
        for c_idx, val in enumerate(q, 1):
            cell = ws.cell(r_idx, c_idx, value=val if c_idx != 1 else (r_idx - 1))
            cell.font = data_font
            cell.border = border_cell
            if c_idx in [1, 9, 11]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center')


def main():
    print("Starting generation of institutional training workbook...")

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)  # Remove default sheet

    # Intellectual Property
    out_wb.properties.creator = "Eslam Abdelbadea"
    out_wb.properties.lastModifiedBy = "Eslam Abdelbadea"
    out_wb.properties.title = "برنامج التدريب المعتمد وبنوك الأسئلة - معدة الضبع الأسود"
    out_wb.properties.subject = "برنامج تدريب وبنوك أسئلة منظومة الدفاع الجوي"

    border_cell = create_thin_border()
    font_title = Font(name='Times New Roman', size=22, bold=True)
    font_header = Font(name='Times New Roman', size=15, bold=True)
    font_data = Font(name='Times New Roman', size=12, bold=False)
    font_exam = Font(name='Times New Roman', size=15, bold=True)
    font_total = Font(name='Times New Roman', size=15, bold=True)

    header_fill = PatternFill(start_color='E8EEF5', end_color='E8EEF5', fill_type='solid')
    exam_fill = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
    total_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')

    col_widths_prog = {
        'A': 14, 'B': 14, 'C': 12, 'D': 44, 'E': 44,
        'F': 14, 'G': 10, 'H': 10, 'I': 14
    }

    # =========================================================================
    # 1. SHEET: برنامج تدريب - القسم الإعدادي (Term 1 only: 56h = 28 lectures)
    # =========================================================================
    ws_prep = out_wb.create_sheet(title='برنامج تدريب - القسم الإعدادي')
    apply_rtl(ws_prep)
    ws_prep.merge_cells('A1:I1')
    t_cell = ws_prep['A1']
    t_cell.value = "برنامج محاضرات تخصص ( الضبع الاسود ) للقسم ( الإعدادي )"
    t_cell.font = font_title
    t_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws_prep.row_dimensions[1].height = 38
    ws_prep.row_dimensions[2].height = 10

    add_training_program_term(
        ws_prep,
        term_name='الترم الأول',
        schedule=PREP_SCHEDULE,
        days_count=7,
        midterm_day=4,
        th_hours=28.0,
        pr_hours=28.0,
        tot_hours=56.0,
        tot_questions=700,
        start_row=3,
        border_cell=border_cell,
        font_header=font_header,
        font_data=font_data,
        font_exam=font_exam,
        font_total=font_total,
        header_fill=header_fill,
        exam_fill=exam_fill,
        total_fill=total_fill
    )
    for col_l, w in col_widths_prog.items():
        ws_prep.column_dimensions[col_l].width = w

    # =========================================================================
    # 2. SHEET: برنامج تدريب - القسم المتوسط (Term 1: 96h + Term 2: 112h)
    # =========================================================================
    ws_med = out_wb.create_sheet(title='برنامج تدريب - القسم المتوسط')
    apply_rtl(ws_med)
    ws_med.merge_cells('A1:I1')
    t_cell = ws_med['A1']
    t_cell.value = "برنامج محاضرات تخصص ( الضبع الاسود ) للقسم ( المتوسط )"
    t_cell.font = font_title
    t_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws_med.row_dimensions[1].height = 38
    ws_med.row_dimensions[2].height = 10

    # Render Term 1 (96h = 48 lectures)
    next_r = add_training_program_term(
        ws_med,
        term_name='الترم الأول',
        schedule=MED_T1_SCHEDULE,
        days_count=12,
        midterm_day=6,
        th_hours=48.0,
        pr_hours=48.0,
        tot_hours=96.0,
        tot_questions=1200,
        start_row=3,
        border_cell=border_cell,
        font_header=font_header,
        font_data=font_data,
        font_exam=font_exam,
        font_total=font_total,
        header_fill=header_fill,
        exam_fill=exam_fill,
        total_fill=total_fill
    )

    # Render Term 2 (112h = 56 lectures) directly below Term 1
    add_training_program_term(
        ws_med,
        term_name='الترم الثاني',
        schedule=MED_T2_SCHEDULE,
        days_count=14,
        midterm_day=7,
        th_hours=56.0,
        pr_hours=56.0,
        tot_hours=112.0,
        tot_questions=1400,
        start_row=next_r,
        border_cell=border_cell,
        font_header=font_header,
        font_data=font_data,
        font_exam=font_exam,
        font_total=font_total,
        header_fill=header_fill,
        exam_fill=exam_fill,
        total_fill=total_fill
    )
    for col_l, w in col_widths_prog.items():
        ws_med.column_dimensions[col_l].width = w

    # =========================================================================
    # 3. SHEET: برنامج تدريب - القسم النهائي (Term 1: 96h + Term 2: 80h) - Ref: الايجلا
    # =========================================================================
    ws_fin = out_wb.create_sheet(title='برنامج تدريب - القسم النهائي')
    apply_rtl(ws_fin)
    ws_fin.merge_cells('A1:I1')
    t_cell = ws_fin['A1']
    t_cell.value = "برنامج محاضرات تخصص ( الضبع الاسود ) للقسم ( النهائي )"
    t_cell.font = font_title
    t_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws_fin.row_dimensions[1].height = 38
    ws_fin.row_dimensions[2].height = 10

    # Render Term 1 (96h = 48 lectures)
    next_r_fin = add_training_program_term(
        ws_fin,
        term_name='الترم الأول',
        schedule=FINAL_T1_SCHEDULE,
        days_count=12,
        midterm_day=6,
        th_hours=48.0,
        pr_hours=48.0,
        tot_hours=96.0,
        tot_questions=1200,
        start_row=3,
        border_cell=border_cell,
        font_header=font_header,
        font_data=font_data,
        font_exam=font_exam,
        font_total=font_total,
        header_fill=header_fill,
        exam_fill=exam_fill,
        total_fill=total_fill
    )

    # Render Term 2 (80h = 40 lectures) directly below Term 1
    add_training_program_term(
        ws_fin,
        term_name='الترم الثاني',
        schedule=FINAL_T2_SCHEDULE,
        days_count=10,
        midterm_day=5,
        th_hours=40.0,
        pr_hours=40.0,
        tot_hours=80.0,
        tot_questions=1000,
        start_row=next_r_fin,
        border_cell=border_cell,
        font_header=font_header,
        font_data=font_data,
        font_exam=font_exam,
        font_total=font_total,
        header_fill=header_fill,
        exam_fill=exam_fill,
        total_fill=total_fill
    )
    for col_l, w in col_widths_prog.items():
        ws_fin.column_dimensions[col_l].width = w

    print("Program sheets built successfully.")

    # =========================================================================
    # LOAD AND PREPARE QUESTIONS FOR ALL 5 QUESTION BANKS
    # =========================================================================
    print("Loading source question banks...")
    src_main = r'C:\Users\MaximuM-Tech\Downloads\بنوك\معدة الضبع الاسود.xlsx'
    wb_main = openpyxl.load_workbook(src_main, data_only=True)

    def extract_rows(sheet_name):
        ws_s = wb_main[sheet_name]
        res = []
        for r in range(2, ws_s.max_row + 1):
            row_vals = [ws_s.cell(r, c).value for c in range(1, 12)]
            if any(row_vals):
                res.append(row_vals)
        return res

    raw_prep_qs = extract_rows('بنك القسم الإعدادي')
    raw_med_qs = extract_rows('بنك القسم المتوسط')
    raw_fin_qs = extract_rows('بنك القسم النهائي')

    # Load 800 extra Eagla questions
    eagla_extra = []
    p_bank = r'C:\Users\MaximuM-Tech\Downloads\بنك'
    for fname in ['Eagla_Bank_1_Technical_200.xlsx', 'Eagla_Bank_2_Tactics_200.xlsx', 'Eagla_Bank_3_Firing_200.xlsx', 'Eagla_Bank_4_Fire_200.xlsx']:
        fpath = os.path.join(p_bank, fname)
        if os.path.exists(fpath):
            wb_e = openpyxl.load_workbook(fpath, data_only=True)
            ws_e = wb_e['Examples']
            for r in range(2, ws_e.max_row + 1):
                q_txt = ws_e.cell(r, 2).value
                if q_txt:
                    eagla_extra.append([
                        0,
                        q_txt,
                        ws_e.cell(r, 3).value or "",
                        ws_e.cell(r, 4).value or "",
                        ws_e.cell(r, 5).value or "",
                        ws_e.cell(r, 6).value or "",
                        ws_e.cell(r, 7).value or "",
                        ws_e.cell(r, 8).value or "",
                        ws_e.cell(r, 9).value or "متوسط",
                        ws_e.cell(r, 10).value or "النظام الصاروخي إيجلا",
                        ws_e.cell(r, 11).value or "اختيار من متعدد"
                    ])

    print(f"Loaded pools: Prep={len(raw_prep_qs)}, Med={len(raw_med_qs)}, Fin={len(raw_fin_qs)}, EaglaExtra={len(eagla_extra)}")

    # Format question banks fonts
    qb_header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    qb_header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    qb_data_font = Font(name='Calibri', size=10, bold=False)

    # 1. بنك القسم الإعدادي (700 سؤال)
    prep_bank = raw_prep_qs[:700]
    # Update lesson titles to match prep schedule
    for idx, q in enumerate(prep_bank):
        lec_idx = idx // 25
        if lec_idx < len(PREP_SCHEDULE):
            it = PREP_SCHEDULE[lec_idx]
            q[9] = it['theory'] if it['theory'] != '-' else it['prac']
    ws_qb_prep = out_wb.create_sheet(title='بنك القسم الإعدادي')
    build_question_bank_sheet(ws_qb_prep, 'بنك القسم الإعدادي', prep_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print("Created بنك القسم الإعدادي (700 سؤال).")

    # 2. بنك المتوسط - ترم أول (1200 سؤال = 48 محاضرة * 25)
    med_t1_bank = []
    for idx in range(1200):
        lec_idx = idx // 25
        it = MED_T1_SCHEDULE[lec_idx]
        lec_title = it['theory'] if it['theory'] != '-' else it['prac']
        src_q = copy(raw_med_qs[idx % len(raw_med_qs)])
        src_q[0] = idx + 1
        src_q[9] = lec_title
        med_t1_bank.append(src_q)
    ws_qb_med_t1 = out_wb.create_sheet(title='بنك المتوسط - ترم أول')
    build_question_bank_sheet(ws_qb_med_t1, 'بنك المتوسط - ترم أول', med_t1_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print("Created بنك المتوسط - ترم أول (1200 سؤال).")

    # 3. بنك المتوسط - ترم ثاني (1400 سؤال = 56 محاضرة * 25)
    med_t2_bank = []
    for idx in range(1400):
        lec_idx = idx // 25
        it = MED_T2_SCHEDULE[lec_idx]
        lec_title = it['theory'] if it['theory'] != '-' else it['prac']
        src_q = copy(raw_med_qs[idx % len(raw_med_qs)])
        src_q[0] = idx + 1
        src_q[9] = lec_title
        med_t2_bank.append(src_q)
    ws_qb_med_t2 = out_wb.create_sheet(title='بنك المتوسط - ترم ثاني')
    build_question_bank_sheet(ws_qb_med_t2, 'بنك المتوسط - ترم ثاني', med_t2_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print("Created بنك المتوسط - ترم ثاني (1400 سؤال).")

    # 4. بنك النهائي - ترم أول (1200 سؤال = 48 محاضرة * 25)
    full_eagla_pool = eagla_extra + raw_fin_qs
    fin_t1_bank = []
    for idx in range(1200):
        lec_idx = idx // 25
        it = FINAL_T1_SCHEDULE[lec_idx]
        lec_title = it['theory'] if it['theory'] != '-' else it['prac']
        src_q = copy(full_eagla_pool[idx % len(full_eagla_pool)])
        src_q[0] = idx + 1
        src_q[9] = lec_title
        fin_t1_bank.append(src_q)
    ws_qb_fin_t1 = out_wb.create_sheet(title='بنك النهائي - ترم أول')
    build_question_bank_sheet(ws_qb_fin_t1, 'بنك النهائي - ترم أول', fin_t1_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print("Created بنك النهائي - ترم أول (1200 سؤال).")

    # 5. بنك النهائي - ترم ثاني (1000 سؤال = 40 محاضرة * 25)
    fin_t2_bank = []
    for idx in range(1000):
        lec_idx = idx // 25
        it = FINAL_T2_SCHEDULE[lec_idx]
        lec_title = it['theory'] if it['theory'] != '-' else it['prac']
        src_q = copy(raw_fin_qs[idx % len(raw_fin_qs)])
        src_q[0] = idx + 1
        src_q[9] = lec_title
        fin_t2_bank.append(src_q)
    ws_qb_fin_t2 = out_wb.create_sheet(title='بنك النهائي - ترم ثاني')
    build_question_bank_sheet(ws_qb_fin_t2, 'بنك النهائي - ترم ثاني', fin_t2_bank, border_cell, qb_header_font, qb_data_font, qb_header_fill)
    print("Created بنك النهائي - ترم ثاني (1000 سؤال).")

    out_file_path = os.path.join(r"C:\Users\MaximuM-Tech\Downloads\بنوك", "برنامج تدريب تخصص الضبع الاسود.xlsx")
    out_wb.save(out_file_path)
    print(f"\n=======================================================")
    print(f"SUCCESS: Generated {out_file_path}")
    print(f"Total Sheets in workbook ({len(out_wb.sheetnames)}): {out_wb.sheetnames}")
    print(f"=======================================================")

if __name__ == '__main__':
    main()
