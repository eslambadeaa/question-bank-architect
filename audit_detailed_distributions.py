import openpyxl
import collections
import sys

sys.stdout.reconfigure(encoding='utf-8')

files = [
    ('القسم المتوسط', r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'),
    ('القسم النهائي', r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx')
]

for sec_name, fpath in files:
    print(f"\n=======================================================")
    print(f"تحليل تفصيلي لـ: {sec_name}")
    print(f"=======================================================")
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb['Questions']
    
    # Col 1: Type, Col 2: Text, Col 4: Diff, Col 5: Ans, Col 10: Lesson
    lessons = collections.defaultdict(list)
    for r in range(2, ws.max_row + 1):
        q_type = ws.cell(r, 1).value
        q_text = ws.cell(r, 2).value
        q_diff = ws.cell(r, 4).value
        q_ans = ws.cell(r, 5).value
        q_les = ws.cell(r, 10).value
        lessons[q_les].append((q_type, q_diff, q_text, r))
        
    print(f"عدد الدروس: {len(lessons)}")
    
    # فحص توزيع كل درس
    sample_printed = 0
    pattern_summary = collections.defaultdict(int)
    
    for les_name, items in lessons.items():
        mcqs = [it for it in items if 'متعدد' in str(it[0])]
        tfs = [it for it in items if 'صح' in str(it[0])]
        
        mcq_c = collections.Counter([it[1] for it in mcqs])
        tf_c = collections.Counter([it[1] for it in tfs])
        tot_c = collections.Counter([it[1] for it in items])
        
        # نسجل نمط توزيع الدرس
        pattern_key = (
            len(mcqs), len(tfs),
            tuple(sorted(tot_c.items()))
        )
        pattern_summary[pattern_key] += 1
        
        if sample_printed < 2:
            print(f"\nنموذج الدرس: {les_name}")
            print(f"  - عدد الاختيار: {len(mcqs)} سؤال")
            print(f"    توزيع صعوبة الاختيار: {dict(mcq_c)}")
            print(f"  - عدد الصح والخطأ: {len(tfs)} سؤال")
            print(f"    توزيع صعوبة الصح والخطأ: {dict(tf_c)}")
            print(f"  - إجمالي صعوبة الـ 25 سؤال: {dict(tot_c)}")
            sample_printed += 1

    print("\nملخص أنماط التوزيع عبر جميع الدروس في الملف:")
    for pat, count in pattern_summary.items():
        mcq_len, tf_len, tot_diff = pat
        print(f"  {count} درس بنمط: {mcq_len} اختيار + {tf_len} صح/خطأ | إجمالي الصعوبات: {dict(tot_diff)}")
