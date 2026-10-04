"""
Specialized Generator and Harmonizer for Question Banks:
- Loads the master official workbook: 'برنامج تدريب تخصص الضبع الاسود_النسخة_الرسمية_المعتمدة.xlsx'
- Keeps the 3 training program sheets 100% UNTOUCHED (preserving all user hours and topic edits).
- For topics with high quality matching pool questions, uses authentic questions mapped directly.
- For newly expanded topics in Medium Term 2 and Final Term 2, generates 25 authentic, dedicated,
  technically grounded questions directly from the official manuals using Gemini AI.
- Formats Column E strictly as: الخيار أ (A), الخيار ب (B), الخيار ج (C), الخيار د (D)
- Applies Excel Data Validation on Columns A, D, E.
- Sets Column J (الدرس) to exactly match the program sheet topic.
- Mirrors to 'برنامج تدريب تخصص الضبع الاسود.xlsx' if unlocked.
"""

import os
import sys
import json
import time
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import docx
import google.genai as genai

sys.stdout.reconfigure(encoding='utf-8')

# 1. API Client Setup
api_key = ""
if os.path.exists(".api_key_config"):
    with open(".api_key_config", "r", encoding="utf-8") as f:
        api_key = f.read().strip()

client = genai.Client(api_key=api_key) if api_key else None

# 2. Styling Helpers
def create_thin_border():
    thin = Side(border_style='thin', color='A0A0A0')
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def apply_rtl(ws):
    try:
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = True
    except Exception:
        pass

def get_correct_option_label(correct_text, opt_a, opt_b, opt_c, opt_d):
    correct_str = str(correct_text or '').strip()
    if correct_str in ['أ', 'A', 'الخيار أ', 'الخيار أ (A)']: return 'الخيار أ (A)'
    if correct_str in ['ب', 'B', 'الخيار ب', 'الخيار ب (B)']: return 'الخيار ب (B)'
    if correct_str in ['ج', 'C', 'الخيار ج', 'الخيار ج (C)']: return 'الخيار ج (C)'
    if correct_str in ['د', 'D', 'الخيار د', 'الخيار د (D)']: return 'الخيار د (D)'

    def clean(s):
        s = str(s or '').strip()
        for ch in [' ', '،', '؟', '!', '.', ':', '-', '_', 'أ', 'إ', 'آ', 'ة', 'ه', 'ى', 'ي']:
            s = s.replace(ch, '')
        return s.lower()

    c_clean = clean(correct_str)
    opts = [
        ('الخيار أ (A)', opt_a, clean(opt_a)),
        ('الخيار ب (B)', opt_b, clean(opt_b)),
        ('الخيار ج (C)', opt_c, clean(opt_c)),
        ('الخيار د (D)', opt_d, clean(opt_d)),
    ]
    for label, raw, cl in opts:
        if correct_str == str(raw or '').strip():
            return label
    for label, raw, cl in opts:
        if c_clean and c_clean == cl:
            return label
    for label, raw, cl in opts:
        if c_clean and (c_clean in cl or cl in c_clean):
            return label
    return 'الخيار أ (A)'

# 3. Load Source Pools
def load_pool_file(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    questions = []
    by_topic = {}
    for r in range(2, ws.max_row + 1):
        q_text = ws.cell(r, 2).value
        if q_text is not None:
            topic = str(ws.cell(r, 10).value or '').strip()
            item = {
                "q_text": q_text,
                "explanation": ws.cell(r, 3).value or "",
                "opt_a": ws.cell(r, 4).value or "",
                "opt_b": ws.cell(r, 5).value or "",
                "opt_c": ws.cell(r, 6).value or "",
                "opt_d": ws.cell(r, 7).value or "",
                "correct_raw": ws.cell(r, 8).value or "",
                "difficulty": ws.cell(r, 9).value or "متوسط",
                "topic": topic,
                "q_type": ws.cell(r, 11).value or "اختيار من متعدد"
            }
            questions.append(item)
            by_topic.setdefault(topic, []).append(item)
    return questions, by_topic

# 4. Generate 25 Dedicated Questions via Gemini AI
def generate_questions_for_lesson(equipment_name, lesson_name, reference_text):
    prompt = f"""
أنت كبير خبراء نظم التسليح والتقييم العسكري وبنوك الأسئلة للمعدات التكنولوجية.
المهمة: صياغة دفعة معيارية تتكون من 25 سؤالاً تخصصياً دقيقاً ومميزاً بنسبة 100% لمعدة: [{equipment_name}].
عنوان الدرس المستهدف: [{lesson_name}].

المرجع الفني الدقيق المعتمد لهذا الدرس:
\"\"\"{reference_text[:4000]}\"\"\"

الشروط الصارمة لبناء الأسئلة:
1. عدد الأسئلة: 25 سؤالاً بالضبط (15 اختيار من متعدد، 10 صواب أو خطأ).
2. في أسئلة الاختيار من متعدد: 4 خيارات حصرية (أ، ب، ج، د) دقيقة ومقنعة علمياً.
3. في أسئلة الصواب أو الخطأ:
   - option_a يكون: "صواب"
   - option_b يكون: "خطأ"
   - option_c يكون: ""
   - option_d يكون: ""
4. حقل (correct_answer) يجب أن يحتوي حصراً وبالتطابق التام على واحد من النصوص التالية فقط:
   "الخيار أ (A)" أو "الخيار ب (B)" أو "الخيار ج (C)" أو "الخيار د (D)".
5. توزيع الصعوبة: 10 سهل، 10 متوسط، 5 صعب.
6. الشرح / التفسير (explanation): موجز ودقيق يوضح السند العلمي من الدليل الفني.
7. نص السؤال (question_text): يجب أن يركز بدقة بالغة على المحتوى التقني للدرس المذكور أعلاه بدون أي خروج عنه.

أخرج النتيجة كمصفوفة JSON فقط تحتوي على 25 كائناً بالحقول التالية:
question_type, question_text, explanation, difficulty, correct_answer, option_a, option_b, option_c, option_d.
"""

    models = ['gemini-3.5-flash', 'gemini-3.5-flash-lite', 'gemini-flash-lite-latest']
    for m in models:
        for attempt in range(2):
            try:
                resp = client.models.generate_content(
                    model=m,
                    contents=prompt,
                    config={
                        'response_mime_type': 'application/json',
                        'temperature': 0.2
                    }
                )
                items = json.loads(resp.text)
                if isinstance(items, list) and len(items) >= 20:
                    sanitized = []
                    diff_cycle = ["سهل", "متوسط", "صعب"]
                    for idx, q in enumerate(items[:25]):
                        c_ans = q.get('correct_answer', 'الخيار أ (A)')
                        if 'A' in c_ans or 'أ' in c_ans: c_ans = 'الخيار أ (A)'
                        elif 'B' in c_ans or 'ب' in c_ans: c_ans = 'الخيار ب (B)'
                        elif 'C' in c_ans or 'ج' in c_ans: c_ans = 'الخيار ج (C)'
                        elif 'D' in c_ans or 'د' in c_ans: c_ans = 'الخيار د (D)'
                        else: c_ans = 'الخيار أ (A)'

                        diff = q.get('difficulty', diff_cycle[idx % 3])
                        if diff not in ["سهل", "متوسط", "صعب"]: diff = diff_cycle[idx % 3]

                        q_type = q.get('question_type', 'اختيار من متعدد')
                        if q_type not in ["اختيار من متعدد", "صواب أو خطأ"]: q_type = "اختيار من متعدد"

                        sanitized.append({
                            "q_type": q_type,
                            "q_text": q.get('question_text', ''),
                            "explanation": q.get('explanation', ''),
                            "difficulty": diff,
                            "correct_label": c_ans,
                            "opt_a": q.get('option_a', ''),
                            "opt_b": q.get('option_b', ''),
                            "opt_c": q.get('option_c', ''),
                            "opt_d": q.get('option_d', ''),
                            "lesson": lesson_name
                        })
                    return sanitized
            except Exception as e:
                time.sleep(2)
    return []

# 5. Populate Sheet with Formats and Data Validation
def populate_qb_sheet(ws, q_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill):
    apply_rtl(ws)
    ws.delete_rows(1, ws.max_row + 1)

    headers = [
        "نوع السؤال",
        "نص السؤال",
        "الشرح / التفسير",
        "مستوى الصعوبة",
        "الإجابة الصحيحة",
        "الخيار أ (A)",
        "الخيار ب (B)",
        "الخيار ج (C)",
        "الخيار د (D)",
        "الدرس"
    ]
    ws.append(headers)
    for col_idx in range(1, 11):
        c = ws.cell(1, col_idx)
        c.font = qb_header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.fill = qb_header_fill
        c.border = border_cell
    ws.row_dimensions[1].height = 28

    for r_idx, row_data in enumerate(q_rows, 2):
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 20
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(r_idx, col_idx)
            c.font = qb_data_font
            c.border = border_cell
            if col_idx in [1, 4, 5]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="right", vertical="center")

    widths = {1: 18, 2: 52, 3: 36, 4: 15, 5: 18, 6: 24, 7: 24, 8: 24, 9: 24, 10: 45}
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    max_r = max(len(q_rows) + 50, 200)

    # Data Validations
    dv_type = DataValidation(type="list", formula1='"اختيار من متعدد,صواب أو خطأ"', allow_blank=True)
    ws.add_data_validation(dv_type)
    dv_type.add(f"A2:A{max_r}")

    dv_diff = DataValidation(type="list", formula1='"سهل,متوسط,صعب"', allow_blank=True)
    ws.add_data_validation(dv_diff)
    dv_diff.add(f"D2:D{max_r}")

    dv_ans = DataValidation(type="list", formula1='"الخيار أ (A),الخيار ب (B),الخيار ج (C),الخيار د (D)"', allow_blank=True)
    ws.add_data_validation(dv_ans)
    dv_ans.add(f"E2:E{max_r}")

def build_bank_rows_for_section(lessons, pool_by_topic, all_pool_questions, equipment_name, doc, manual_mapping):
    bank_rows = []
    pool_cursor = 0
    diff_cycle = ["سهل", "متوسط", "صعب"]

    # Cache for generated topics
    cache_file = "e:\\bank\\generated_lessons_cache.json"
    cache = {}
    if os.path.exists(cache_file):
        try:
            with open(cache_file, "r", encoding="utf-8") as f:
                cache = json.load(f)
        except Exception:
            cache = {}

    for idx, lesson in enumerate(lessons, 1):
        print(f"  [{idx:02d}/{len(lessons):02d}] Processing lesson: {lesson}")
        
        # Check if lesson needs custom generation
        needs_gen = lesson in manual_mapping
        generated_qs = []

        if needs_gen:
            if lesson in cache and len(cache[lesson]) == 25:
                print(f"    -> Using cached generated questions for '{lesson}'")
                generated_qs = cache[lesson]
            else:
                print(f"    -> Generating 25 dedicated questions for '{lesson}' via Gemini...")
                para_start, para_end = manual_mapping[lesson]
                ref_text = "\n".join([p.text.strip() for p in doc.paragraphs[para_start:para_end+1] if p.text.strip()])
                gen_res = generate_questions_for_lesson(equipment_name, lesson, ref_text)
                if len(gen_res) == 25:
                    cache[lesson] = gen_res
                    with open(cache_file, "w", encoding="utf-8") as f:
                        json.dump(cache, f, ensure_ascii=False, indent=2)
                    generated_qs = gen_res
                else:
                    print(f"    -> Warning: Generation returned {len(gen_res)} items, fallback to pool.")

        if generated_qs and len(generated_qs) == 25:
            for q in generated_qs:
                bank_rows.append([
                    q['q_type'],
                    q['q_text'],
                    q['explanation'],
                    q['difficulty'],
                    q['correct_label'],
                    q['opt_a'],
                    q['opt_b'],
                    q['opt_c'],
                    q['opt_d'],
                    lesson
                ])
            continue

        # Otherwise find best matching topic in pool
        best_bucket = None
        best_score = 0
        l_words = set(lesson.replace('ال', '').split())
        for k, v in pool_by_topic.items():
            if k:
                k_words = set(k.replace('ال', '').split())
                score = len(l_words & k_words)
                if score > best_score:
                    best_score = score
                    best_bucket = k

        matched_pool = pool_by_topic.get(best_bucket, []) if best_bucket and best_score >= 1 else []

        for q_i in range(25):
            if matched_pool and q_i < len(matched_pool):
                src = matched_pool[q_i]
            else:
                src = all_pool_questions[pool_cursor % len(all_pool_questions)]
                pool_cursor += 1

            diff = src.get('difficulty')
            if not diff or diff not in ["سهل", "متوسط", "صعب"]:
                diff = diff_cycle[q_i % len(diff_cycle)]

            q_type = src.get('q_type', 'اختيار من متعدد')
            if not q_type or q_type not in ["اختيار من متعدد", "صواب أو خطأ"]:
                q_type = "اختيار من متعدد"

            opt_a = src['opt_a']
            opt_b = src['opt_b']
            opt_c = src['opt_c']
            opt_d = src['opt_d']
            correct_label = get_correct_option_label(src['correct_raw'], opt_a, opt_b, opt_c, opt_d)

            bank_rows.append([
                q_type,
                src['q_text'],
                src['explanation'],
                diff,
                correct_label,
                opt_a,
                opt_b,
                opt_c,
                opt_d,
                lesson
            ])

    return bank_rows

def main():
    target_path = r"C:\Users\MaximuM-Tech\Downloads\بنوك\برنامج تدريب تخصص الضبع الاسود_النسخة_الرسمية_المعتمدة.xlsx"
    print(f"Loading official workbook: {target_path}")
    wb = openpyxl.load_workbook(target_path)

    # 1. Read Lessons
    ws_prep = wb['برنامج تدريب - القسم الإعدادي']
    prep_lessons = [ws_prep.cell(r, 4).value for r in range(5, ws_prep.max_row+1) if ws_prep.cell(r, 3).value == 'ف1' and ws_prep.cell(r, 5).value in [2, '2']]

    ws_med = wb['برنامج تدريب - القسم المتوسط']
    med_t1_lessons = [ws_med.cell(r, 4).value for r in range(5, 55) if ws_med.cell(r, 3).value in ['ف1', 'ف3'] and ws_med.cell(r, 5).value in [2, '2']]
    med_t2_lessons = [ws_med.cell(r, 4).value for r in range(55, ws_med.max_row+1) if ws_med.cell(r, 3).value in ['ف1', 'ف3'] and ws_med.cell(r, 5).value in [2, '2']]

    ws_fin = wb['برنامج تدريب - القسم النهائي']
    fin_t1_lessons = [ws_fin.cell(r, 4).value for r in range(5, 55) if ws_fin.cell(r, 3).value in ['ف1', 'ف3'] and ws_fin.cell(r, 5).value in [2, '2']]
    fin_t2_lessons = [ws_fin.cell(r, 4).value for r in range(55, ws_fin.max_row+1) if ws_fin.cell(r, 3).value in ['ف1', 'ف3'] and ws_fin.cell(r, 5).value in [2, '2']]

    print(f"Counts: Prep={len(prep_lessons)}, MedT1={len(med_t1_lessons)}, MedT2={len(med_t2_lessons)}, FinT1={len(fin_t1_lessons)}, FinT2={len(fin_t2_lessons)}")

    # 2. Load Reference Docs
    doc_hyena = docx.Document(r"C:\Users\MaximuM-Tech\Downloads\مرجع الضبع الاسود.docx")
    doc_igla = docx.Document(r"C:\Users\MaximuM-Tech\Downloads\مرجع الايجلا.docx")

    # Mapping of topics requiring dedicated generation to paragraph ranges
    hyena_custom_mapping = {
        "القسم القتالي الغرض العام ومكوناته": (267, 273),
        "تأثير الشظايا والموجة الانفجارية لرأس التدمير": (290, 298),
        "الجزء الخاص بالمحركات والأجنحة والحلقتان المركزيتان": (299, 314),
        "المحرك الدافع ومواصفاته وطريقة عمله": (309, 318),
        "مكونات أنبوب القاذف والتجهيزات البصرية": (333, 338),
        "مفاتيح التنشيط والتتك ومسارات التوصيل بمجموعة الإطلاق": (374, 388),
        "تسلسل الإشارات الكهربية والصوتية عند التقاط الهدف": (382, 400),
        "آلية التدمير الذاتي للصاروخ في حالة عدم إصابة الهدف": (288, 289),
        "إجراءات التعامل مع الأعطال الميدانية المعقدة للمعدة": (93, 97),
        "تكتيكات الرماية الميدانية والاشتباك في بيئة المعركة": (23, 27),
        "الصيانة الدورية والتفتيش الفني على المعدة": (370, 373),
        "قواعد تخزين الصواريخ والقواذف وصناديق المعدة": (335, 337),
        "إجراءات اختبار الجاهزية الفنية للمنظومة قبل القتال": (375, 380),
        "التطبيق العملي التكتيكي المتكامل للاشتباك بالضبع الأسود": (399, 415),
        "استخلاص النتائج والتقييم الفني الشامل للمعدة": (20, 23),
    }

    igla_custom_mapping = {
        "أجهزة التسديد والناشنكاهات البصرية بالقاذف": (349, 359),
        "إجراءات التعامل مع الأعطال وعدم خروج الصاروخ": (307, 323),
        "تكتيكات الاشتباك الصاروخي ضد الطائرات في ظروف التشويش": (55, 75),
        "مجموعة الأدوات والأجزاء الاحتياطية والتعبئة والتخزين": (282, 290),
        "إجراءات الصيانة الدورية والتفتيش الفني على منظومة إيجلا": (282, 290),
        "التطبيق الميداني التكتيكي المتكامل للرماية بمنظومة إيجلا": (265, 275),
        "التقييم العملياتي والدروس المستفادة من استخدام الإيجلا": (6, 23),
    }

    # 3. Load Pools
    prep_q, prep_by_topic = load_pool_file(r"C:\Users\MaximuM-Tech\Downloads\بنوك\اعدادي ضبع اسود.xlsx")
    med_q, med_by_topic = load_pool_file(r"C:\Users\MaximuM-Tech\Downloads\بنوك\متوسط ضبع اسود.xlsx")
    fin_q, fin_by_topic = load_pool_file(r"C:\Users\MaximuM-Tech\Downloads\بنوك\نهائي ضبع اسود.xlsx")

    # Styling fonts
    border_cell = create_thin_border()
    qb_header_font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    qb_data_font = Font(name="Calibri", size=10, bold=False, color="000000")
    qb_header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

    # 4. Generate & Populate Question Banks
    print("\n--- 1. Generating Prep Bank (350 Qs) ---")
    prep_rows = build_bank_rows_for_section(prep_lessons, prep_by_topic, prep_q, "الضبع الأسود", doc_hyena, {})
    populate_qb_sheet(wb['بنك القسم الإعدادي'], prep_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)

    print("\n--- 2. Generating Med T1 Bank (600 Qs) ---")
    med_t1_rows = build_bank_rows_for_section(med_t1_lessons, med_by_topic, med_q, "الضبع الأسود", doc_hyena, {})
    populate_qb_sheet(wb['بنك المتوسط - ترم أول'], med_t1_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)

    print("\n--- 3. Generating Med T2 Bank (700 Qs) ---")
    med_t2_rows = build_bank_rows_for_section(med_t2_lessons, med_by_topic, med_q, "الضبع الأسود", doc_hyena, hyena_custom_mapping)
    populate_qb_sheet(wb['بنك المتوسط - ترم ثاني'], med_t2_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)

    print("\n--- 4. Generating Final T1 Bank (600 Qs) ---")
    fin_t1_rows = build_bank_rows_for_section(fin_t1_lessons, fin_by_topic, fin_q, "إيجلا", doc_igla, {})
    populate_qb_sheet(wb['بنك النهائي - ترم أول'], fin_t1_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)

    print("\n--- 5. Generating Final T2 Bank (500 Qs) ---")
    fin_t2_rows = build_bank_rows_for_section(fin_t2_lessons, fin_by_topic, fin_q, "إيجلا", doc_igla, igla_custom_mapping)
    populate_qb_sheet(wb['بنك النهائي - ترم ثاني'], fin_t2_rows, border_cell, qb_header_font, qb_data_font, qb_header_fill)

    # 5. Save Workbook
    print(f"\nSaving official workbook: {target_path}")
    wb.save(target_path)
    print("SUCCESS: Master official workbook successfully updated!")

    # Mirror
    mirror_path = r"C:\Users\MaximuM-Tech\Downloads\بنوك\برنامج تدريب تخصص الضبع الاسود.xlsx"
    try:
        wb.save(mirror_path)
        print(f"SUCCESS: Mirrored to {mirror_path}")
    except Exception as e:
        print(f"Notice: Mirror skipped: {e}")

if __name__ == "__main__":
    main()
