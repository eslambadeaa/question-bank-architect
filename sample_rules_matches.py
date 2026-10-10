import docx, re, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'C:\Users\MaximuM-Tech\Downloads\قواعد خدمة الميدان صار محمول 29-11.docx')
print("Total paragraphs in rules:", len(doc.paragraphs))
matches = []
for p in doc.paragraphs:
    t = p.text.strip()
    if any(k in t for k in ['إطلاق', 'صاروخ', 'جيروسكوب', 'توربين', 'محرك', 'توجيه', 'عدسي', 'موازنة', 'قاذف']):
        if len(t) > 30:
            matches.append(t)

print(f"Found {len(matches)} matches. Sample 20:")
for m in matches[:20]:
    print(" - ", m[:100])
