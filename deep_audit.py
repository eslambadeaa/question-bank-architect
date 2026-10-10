import sys, collections
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

files = {
    'النهائي': r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_النهائي_الضبع_الاسود_LMS.xlsx',
    'المتوسط': r'C:\Users\MaximuM-Tech\Downloads\بنك_اسئلة_القسم_المتوسط_الضبع_الاسود_LMS.xlsx'
}

for name, path in files.items():
    print('='*60)
    print(f'فحص ملف: {name}')
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    
    rows = []
    for r in range(2, ws.max_row + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 11)]
        rows.append((r, row_vals))
        
    print(f'إجمالي الأسئلة: {len(rows)}')
    
    # 1. فحص التكرارات في نصوص الأسئلة
    texts = collections.defaultdict(list)
    for r_num, vals in rows:
        t = str(vals[2]).strip() if vals[2] is not None else ''
        texts[t].append(r_num)
        
    duplicates = {k: v for k, v in texts.items() if len(v) > 1}
    print(f'عدد الأسئلة المكررة في نفس الملف: {len(duplicates)}')
    if duplicates:
        print('نماذج من التكرارات:')
        for k, v in list(duplicates.items())[:5]:
            print(f'  - \"{k[:45]}...\" مكرر في الصفوف: {v}')
            
    # 2. فحص توزيع الدروس وعدد الأسئلة ونوعها والصعوبة
    lessons = collections.defaultdict(list)
    for r_num, vals in rows:
        lesson = vals[0]
        lessons[lesson].append((r_num, vals))
        
    print(f'عدد الدروس: {len(lessons)}')
    
    # فحص تفصيلي لكل درس
    mcq_tf_counts = collections.defaultdict(int)
    diff_dist_match = 0
    diff_issues = []
    
    for l_idx, (lesson, l_rows) in enumerate(lessons.items()):
        mcq = [x for x in l_rows if x[1][1] == '1 - متعدد']
        tf = [x for x in l_rows if x[1][1] == '2 - صح/خطأ']
        
        mcq_tf_counts[(len(mcq), len(tf))] += 1
        
        mcq_diff = collections.Counter([x[1][9] for x in mcq])
        tf_diff = collections.Counter([x[1][9] for x in tf])
        tot_diff = collections.Counter([x[1][9] for x in l_rows])
        
        # المطلوب: 30% سهل (7-8), 30% متوسط (7-8), 15% صعب (3-4), 15% صعب جدا (3-4), 10% تفوق (2-3)
        # إجمالي 25: سهل: 7 أو 8، متوسط: 7 أو 8، صعب: 4، صعب جدا: 4، تفوق: 2 (7+8+4+4+2=25)
        
        if l_idx == 0:
            print(f'\nمثال تفصيلي للدرس الأول ({lesson}):')
            print(f'  - عدد أسئلة الاختيار: {len(mcq)}')
            print(f'  - عدد أسئلة الصواب والخطأ: {len(tf)}')
            print(f'  - توزيع صعوبة الاختيار: {dict(mcq_diff)}')
            print(f'  - توزيع صعوبة الصواب والخطأ: {dict(tf_diff)}')
            print(f'  - إجمالي صعوبة الدرس: {dict(tot_diff)}')

    print('\nتوزيع أعداد الأسئلة (اختيار / صح وخطأ) عبر الدروس:')
    for (m_c, t_c), count in mcq_tf_counts.items():
        print(f'  {count} درس يحتوي على: {m_c} اختيار و {t_c} صح وخطأ')

    # 3. فحص الخيارات والإجابات وصلاحية الأعمدة
    missing_ans = 0
    invalid_ans = 0
    missing_mcq_opt = 0
    filled_tf_opt = 0
    empty_cells_cols_A_E = 0
    
    for r_num, vals in rows:
        lesson, q_type, q_text, img, correct, a, b, c, d, diff = vals
        
        # Check A, B, C, E, J
        if not lesson or not q_type or not q_text or correct is None or not diff:
            empty_cells_cols_A_E += 1
            
        if q_type == '1 - متعدد':
            if correct not in ['A', 'B', 'C', 'D']:
                invalid_ans += 1
            if any(opt is None or str(opt).strip() == '' for opt in [a, b, c, d]):
                missing_mcq_opt += 1
        elif q_type == '2 - صح/خطأ':
            if str(correct).strip() not in ['true', 'false', 'True', 'False']:
                invalid_ans += 1
            if any(opt is not None and str(opt).strip() != '' for opt in [a, b, c, d]):
                filled_tf_opt += 1
                
    print('\nفحص الحقول والخلايا:')
    print(f'  - خلايا فارغة في الأعمدة الأساسية (A, B, C, E, J): {empty_cells_cols_A_E}')
    print(f'  - إجابات غير صحيحة أو خارج القيم المسموحة: {invalid_ans}')
    print(f'  - خيارات ناقصة في أسئلة الاختيار: {missing_mcq_opt}')
    print(f'  - خيارات موجودة بالخطأ في أسئلة الصح/خطأ: {filled_tf_opt}')
    print(f'  - قواعد التحقق Data Validation: {len(ws.data_validations)}')
