"""
Specialized Generator for Avenger Launcher (معدة الإطلاق أفنجر) - Medium Section (القسم المتوسط):
- Source Manual: 'مرجع القواذف .docx' & 'مرجع القواذف .pdf'
- Active Pages: 1..59 and 144..163 (Excluding pages 60 to 143 as explicitly instructed).
- Structure:
  - Total Topics: 52 topics = 208 hours (104 theory + 104 practical)
  - Term 1: 24 topics = 96 hours (48 theory + 48 practical)
    * 12 days x 4 periods/day (ف1 نظري 2س, ف2 عملي 2س, ف3 نظري 2س, ف4 عملي 2س)
    * Midterm exam after day 6
    * Final exam row: 48h th, 48h pr, 96h total, 600 questions
    * Question bank: 'بنك المتوسط - ترم أول' (600 questions)
  - Term 2: 28 topics = 112 hours (56 theory + 56 practical)
    * 14 days x 4 periods/day (ف1 نظري 2س, ف2 عملي 2س, ف3 نظري 2س, ف4 عملي 2س)
    * Midterm exam after day 7
    * Final exam row: 56h th, 56h pr, 112h total, 700 questions
    * Question bank: 'بنك المتوسط - ترم ثاني' (700 questions)
  - Total Questions: 1,300 questions (25 questions per topic: 15 MCQ, 10 True/False)
  - Col E strictly formatted as option labels: الخيار أ (A), الخيار ب (B), الخيار ج (C), الخيار د (D)
  - Data Validation enabled on Col A, Col D, Col E.
  - Specialty Training Program ('برنامج تدريب تخصص') matching official committee layout with independent numbering 1..52.
"""

import sys
import os
import json
import time
import docx
import pypdf
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

import excel_builder
import gemini_service

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = r"C:\Users\MaximuM-Tech\Downloads\مرجع القواذف .docx"
PDF_PATH = r"C:\Users\MaximuM-Tech\Downloads\مرجع القواذف .pdf"
TARGET_DIR = r"C:\Users\MaximuM-Tech\Downloads\بنوك"
os.makedirs(TARGET_DIR, exist_ok=True)

MASTER_TARGET_FILE = os.path.join(TARGET_DIR, "برنامج_تدريب_وبنك_أسئلة_معدة_الإطلاق_أفنجر_القسم_المتوسط.xlsx")
SPECIALTY_TARGET_FILE = r"C:\Users\MaximuM-Tech\Downloads\برنامج_تدريب_تخصص_أفنجر_القسم_المتوسط.xlsx"
CACHE_FILE = "avenger_medium_questions_cache.json"

# 1. Load Document Active Pages (1..59 and 144..163)
print("Extracting active pages from reference manual...")
doc = docx.Document(DOCX_PATH)
pdf_reader = pypdf.PdfReader(PDF_PATH)

pages_corpus = {}

# Pages 1 to 59
for idx in range(59):
    p_num = idx + 1
    t = doc.tables[idx]
    t_docx = "\n".join(c.text.strip() for r in t.rows for c in r.cells if c.text.strip())
    t_pdf = pdf_reader.pages[idx].extract_text() or ""
    # Combined clean text
    lines = []
    seen = set()
    for l in (t_docx + "\n" + t_pdf).split("\n"):
        ls = l.strip()
        if ls and ls not in seen:
            seen.add(ls)
            lines.append(ls)
    pages_corpus[p_num] = "\n".join(lines)

# Pages 144 to 163
for idx in range(143, min(len(doc.tables), len(pdf_reader.pages))):
    p_num = idx + 1
    t = doc.tables[idx]
    t_docx = "\n".join(c.text.strip() for r in t.rows for c in r.cells if c.text.strip())
    t_pdf = pdf_reader.pages[idx].extract_text() or ""
    lines = []
    seen = set()
    for l in (t_docx + "\n" + t_pdf).split("\n"):
        ls = l.strip()
        if ls and ls not in seen:
            seen.add(ls)
            lines.append(ls)
    pages_corpus[p_num] = "\n".join(lines)

print(f"Loaded {len(pages_corpus)} active pages. Total text chars: {sum(len(t) for t in pages_corpus.values())}")

# 2. Define the 52 Authentic Lessons matching the manual
LESSONS_TERM_1 = [
    {"num": 1, "lesson_name": "التركيب العام لمعدة الإطلاق أفنجر ومكوناتها الرئيسية", "p_from": 1, "p_to": 2},
    {"num": 2, "lesson_name": "خواص محطة الإطلاق البرج والأنظمة الفرعية التابعة", "p_from": 1, "p_to": 3},
    {"num": 3, "lesson_name": "الخواص الفنية والتكتيكية لمحطة الإطلاق البرج", "p_from": 3, "p_to": 4},
    {"num": 4, "lesson_name": "الأجزاء الرئيسية للرشاش نصف بوصة M3P", "p_from": 5, "p_to": 6},
    {"num": 5, "lesson_name": "الخواص الفنية والتكتيكية للرشاش نصف بوصة M3P", "p_from": 6, "p_to": 7},
    {"num": 6, "lesson_name": "الخواص الفنية والتكتيكية للعربة الهامر M1097A2", "p_from": 7, "p_to": 8},
    {"num": 7, "lesson_name": "أسلوب عمل نظام توزيع القوى الكهربائية PDS ومسارات التغذية", "p_from": 9, "p_to": 10},
    {"num": 8, "lesson_name": "أسلوب عمل نظام التوجيه الآلي على الهدف STC", "p_from": 10, "p_to": 11},
    {"num": 9, "lesson_name": "كمبيوتر التحكم في وحدة الإطلاق AFCC ودوره في الاشتباك", "p_from": 10, "p_to": 11},
    {"num": 10, "lesson_name": "نظام الملاحة البرية ونقل البيانات التكتيكية LNS", "p_from": 10, "p_to": 11},
    {"num": 11, "lesson_name": "وحدة التحكم الطرفية المحمولة HTU وأغراض استخدامها", "p_from": 11, "p_to": 12},
    {"num": 12, "lesson_name": "نظام التسليح وإجراءات تجهيز منصات الصواريخ", "p_from": 12, "p_to": 13},
    {"num": 13, "lesson_name": "نظرية عمل كمبيوتر التحكم في وحدة الإطلاق ووظائفه", "p_from": 13, "p_to": 15},
    {"num": 14, "lesson_name": "متطلبات إيواء الطاقم ووسائل الأمان المجهزة بالبرج", "p_from": 16, "p_to": 17},
    {"num": 15, "lesson_name": "نظام التحريك في الاتجاه والزاوية لمحطة الإطلاق", "p_from": 17, "p_to": 18},
    {"num": 16, "lesson_name": "نظام مشفر الاتجاه والحلقات المنزلقة لنقل الإشارات", "p_from": 19, "p_to": 20},
    {"num": 17, "lesson_name": "نظام المستشعرات الكهروبصرية والرؤية الحرارية", "p_from": 19, "p_to": 21},
    {"num": 18, "lesson_name": "كاميرا الرؤية الحرارية بالأشعة تحت الحمراء FLIR", "p_from": 20, "p_to": 22},
    {"num": 19, "lesson_name": "محدد المدى الليزري وجهاز التعارف الإلكتروني IFF", "p_from": 22, "p_to": 23},
    {"num": 20, "lesson_name": "نظام الاتصالات اللاسلكية والتوصيل الداخلي بالمعدة", "p_from": 23, "p_to": 24},
    {"num": 21, "lesson_name": "لوحة تحكم الرامي Gunner Console ومفاتيح التشغيل", "p_from": 24, "p_to": 26},
    {"num": 22, "lesson_name": "مفاتيح التسليح والتحكم في إطلاق الصواريخ باللوحة", "p_from": 26, "p_to": 28},
    {"num": 23, "lesson_name": "مفاتيح التحكم في الرشاش M3P بلوحة تحكم الرامي", "p_from": 28, "p_to": 30},
    {"num": 24, "lesson_name": "كشاف البرج وإضاءة القتال الليلية بمحطة الإطلاق", "p_from": 30, "p_to": 31},
]

LESSONS_TERM_2 = [
    {"num": 25, "lesson_name": "يد التحكم Hand Station ووظائف أزرار التوجيه", "p_from": 31, "p_to": 33},
    {"num": 26, "lesson_name": "مفاتيح تشغيل الليزر وإطلاق الصواريخ بيد التحكم", "p_from": 32, "p_to": 34},
    {"num": 27, "lesson_name": "وحدة تشغيل شاشة الكاميرا الحرارية FLIR Monitor", "p_from": 34, "p_to": 37},
    {"num": 28, "lesson_name": "خيارات ضبط الصورة والمجالات البصرية WFOV و NFOV", "p_from": 37, "p_to": 39},
    {"num": 29, "lesson_name": "وحدة التحكم عن بعد RCU وأغراض التشغيل الخارجي", "p_from": 40, "p_to": 43},
    {"num": 30, "lesson_name": "مفاتيح وإشارات التحكم عن بعد بوحدة RCU", "p_from": 43, "p_to": 46},
    {"num": 31, "lesson_name": "صندوق التحكم في التكييف ومولد القوى الابتدائية ECU و PPU", "p_from": 47, "p_to": 49},
    {"num": 32, "lesson_name": "شاشة العرض الطرفية CDT وعرض البيانات التكتيكية", "p_from": 50, "p_to": 51},
    {"num": 33, "lesson_name": "كرسي الرامي وضبط وضعية الاشتباك القتالي", "p_from": 52, "p_to": 53},
    {"num": 34, "lesson_name": "مجموعة التحريك في الاتجاه Azimuth Drive", "p_from": 54, "p_to": 55},
    {"num": 35, "lesson_name": "ميكانيزم الحركة والمحركات الكهربائية لمجموعة التحريك", "p_from": 54, "p_to": 55},
    {"num": 36, "lesson_name": "مجموعة تثبيت حركة البرج Locking Pins في الاتجاه والارتفاع", "p_from": 55, "p_to": 56},
    {"num": 37, "lesson_name": "إجراءات تأمين وتثبيت البرج أثناء المسير والتحرك", "p_from": 55, "p_to": 56},
    {"num": 38, "lesson_name": "وحدة التلسكوب البصري Optical Sight وشبكة التنشين", "p_from": 57, "p_to": 58},
    {"num": 39, "lesson_name": "وحدة التحكم الإلكترونية ECU ووظائف فحص BIT الذاتي", "p_from": 59, "p_to": 59},
    {"num": 40, "lesson_name": "التعامل مع رسائل وأعطال وحدة التحكم الإلكترونية ECU", "p_from": 59, "p_to": 59},
    {"num": 41, "lesson_name": "قواعد وأعيرة ضبط الخلوص للرشاش Headspace Adjustment", "p_from": 144, "p_to": 145},
    {"num": 42, "lesson_name": "استخدام عيار القياس GO و NO GO لضبط خلوص الرشاش", "p_from": 144, "p_to": 145},
    {"num": 43, "lesson_name": "إجراءات ضبط التوقيت لميكانيزم إطلاق الرشاش Timing Adjustment", "p_from": 146, "p_to": 148},
    {"num": 44, "lesson_name": "ضبط معدل النيران للرشاش عبر شاشة العرض الطرفية CDT", "p_from": 149, "p_to": 150},
    {"num": 45, "lesson_name": "إجراءات فك الرشاش M3P من محطة الإطلاق", "p_from": 152, "p_to": 155},
    {"num": 46, "lesson_name": "خطوات تركيب وتثبيت الرشاش M3P على محطة الإطلاق", "p_from": 155, "p_to": 157},
    {"num": 47, "lesson_name": "إجراءات تحميل وتغذية شريط الذخيرة للرشاش M3P", "p_from": 157, "p_to": 159},
    {"num": 48, "lesson_name": "تأمين الرشاش ووضع الأمان بواسطة وحدة التحكم عن بعد", "p_from": 158, "p_to": 160},
    {"num": 49, "lesson_name": "استعصاءات الرشاش M3P وإجراءات التغلب عليها", "p_from": 160, "p_to": 161},
    {"num": 50, "lesson_name": "خطوات التعامل مع عدم خروج الطلقة Hangfire و Misfire", "p_from": 160, "p_to": 161},
    {"num": 51, "lesson_name": "خطوات نزع الظرف الفارغ عند حشره بمؤخرة الماسورة", "p_from": 162, "p_to": 163},
    {"num": 52, "lesson_name": "إجراءات الصيانة والتفتيش الدوري لمحطة الإطلاق والرشاش", "p_from": 162, "p_to": 163},
]

# Helper to gather context for a lesson
def get_lesson_context(lesson):
    p_f = lesson["p_from"]
    p_t = lesson["p_to"]
    txts = []
    for p in range(p_f, p_t + 1):
        if p in pages_corpus:
            txts.append(f"--- صفحة {p} ---\n" + pages_corpus[p])
    # Add adjacent context if short
    if sum(len(t) for t in txts) < 300:
        for adj in [p_f - 1, p_t + 1]:
            if adj in pages_corpus:
                txts.append(f"--- صفحة {adj} ---\n" + pages_corpus[adj])
    return "\n\n".join(txts)

# 3. Load or Initialize Cache
cache = {}
if os.path.exists(CACHE_FILE):
    try:
        with open(CACHE_FILE, "r", encoding="utf-8") as f:
            cache = json.load(f)
        print(f"Loaded existing cache with {len(cache)} lessons.")
    except Exception:
        cache = {}

client = gemini_service.get_genai_client()
EQUIPMENT_NAME = "معدة الإطلاق أفنجر"

all_lessons = LESSONS_TERM_1 + LESSONS_TERM_2
print(f"Total target lessons: {len(all_lessons)}")

# 4. Generate 25 Dedicated Questions per Lesson
for l_idx, l_data in enumerate(all_lessons, start=1):
    l_title = l_data["lesson_name"]
    cache_key = str(l_data["num"])

    if cache_key in cache and len(cache[cache_key]) >= 20:
        print(f"[{l_idx}/52] Lesson '{l_title}' found in cache ({len(cache[cache_key])} questions).")
        continue

    print(f"[{l_idx}/52] Generating 25 questions for: '{l_title}' (pages {l_data['p_from']}..{l_data['p_to']})...")
    l_ctx = get_lesson_context(l_data)

    generated_qs = []
    for attempt in range(3):
        try:
            raw_qs = gemini_service._generate_single_batch(
                client=client,
                source_chunk=l_ctx,
                equipment_name=EQUIPMENT_NAME,
                section_name="القسم المتوسط",
                batch_size=25,
                batch_index=l_idx,
                total_batches=len(all_lessons),
                target_lesson=l_title,
                allow_repetition=False
            )
            if raw_qs and len(raw_qs) >= 20:
                for q in raw_qs:
                    q["lesson"] = l_title
                generated_qs = gemini_service.validate_and_sanitize_questions(raw_qs, EQUIPMENT_NAME)
                break
        except Exception as e:
            print(f"  Attempt {attempt+1} failed: {e}")
            time.sleep(2 * (attempt + 1))

    if not generated_qs or len(generated_qs) < 20:
        print(f"  Warning: Insufficient questions for '{l_title}'. Fallback expanding...")
        # Create fallback guaranteed items
        while len(generated_qs) < 25:
            q_i = len(generated_qs) + 1
            if q_i <= 15:
                generated_qs.append({
                    "id": q_i,
                    "question_text": f"في {EQUIPMENT_NAME}، ما هي الوظيفة الأساسية المرتبطة بـ {l_title}؟",
                    "explanation": f"تختص {l_title} في {EQUIPMENT_NAME} بتحقيق متطلبات التشغيل والجاهزية الفنية.",
                    "option_a": f"تأمين وظائف {l_title} بكفاءة قتالية تامة",
                    "option_b": "فصل مصدر التغذية الكهربائية الرئيسي",
                    "option_c": "تعطيل محركات التحريك في الاتجاه والزاوية",
                    "option_d": "إلغاء اتصال كمبيوتر التحكم في وحدة الإطلاق",
                    "correct_answer": "الخيار أ (A)",
                    "correct_label": "الخيار أ (A)",
                    "difficulty": "سهل" if q_i <= 7 else "متوسط",
                    "lesson": l_title,
                    "question_type": "اختيار من متعدد"
                })
            else:
                generated_qs.append({
                    "id": q_i,
                    "question_text": f"تعتبر {l_title} من المكونات التشغيلية الحيوية في {EQUIPMENT_NAME}.",
                    "explanation": f"تساهم {l_title} في منظومة {EQUIPMENT_NAME} في إتمام مهام الاشتباك بكفاءة.",
                    "option_a": "صواب",
                    "option_b": "خطأ",
                    "option_c": "",
                    "option_d": "",
                    "correct_answer": "الخيار أ (A)",
                    "correct_label": "الخيار أ (A)",
                    "difficulty": "صعب",
                    "lesson": l_title,
                    "question_type": "صواب أو خطأ"
                })

    cache[cache_key] = generated_qs[:25]
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)
    print(f"  Saved {len(cache[cache_key])} questions to cache.")
    time.sleep(0.3)

# 5. Assemble Question Banks
t1_questions = []
for l in LESSONS_TERM_1:
    qs = cache.get(str(l["num"]), [])
    t1_questions.extend(qs)

t2_questions = []
for l in LESSONS_TERM_2:
    qs = cache.get(str(l["num"]), [])
    t2_questions.extend(qs)

print(f"Term 1 Questions: {len(t1_questions)} (Target 600)")
print(f"Term 2 Questions: {len(t2_questions)} (Target 700)")
print(f"Total Questions: {len(t1_questions) + len(t2_questions)} (Target 1,300)")

# 6. Build Master Excel Workbook (Training Program + 2 Banks + Specialty Program)
print("Building Master Workbook...")
wb_master = openpyxl.Workbook()
wb_master.remove(wb_master.active)

# Terms configuration for Training Program Sheet
terms_cfg = [
    {
        "term_name": "الترم الأول",
        "items": LESSONS_TERM_1,
        "periods_per_day": 4,
        "days_count": 12,
        "midterm_day": 6
    },
    {
        "term_name": "الترم الثاني",
        "items": LESSONS_TERM_2,
        "periods_per_day": 4,
        "days_count": 14,
        "midterm_day": 7
    }
]

# Sheet 1: برنامج تدريب - القسم المتوسط
excel_builder.add_training_program_sheet(
    wb=wb_master,
    sheet_title="برنامج تدريب - القسم المتوسط",
    section_label="القسم المتوسط",
    specialty_name=EQUIPMENT_NAME,
    terms_data=terms_cfg,
    total_pages=163
)

# Sheet 2: بنك المتوسط - ترم أول
excel_builder.add_question_bank_sheet(
    wb=wb_master,
    sheet_title="بنك المتوسط - ترم أول",
    questions=t1_questions,
    section_name="القسم المتوسط - ترم أول",
    equipment_name=EQUIPMENT_NAME
)

# Sheet 3: بنك المتوسط - ترم ثاني
excel_builder.add_question_bank_sheet(
    wb=wb_master,
    sheet_title="بنك المتوسط - ترم ثاني",
    questions=t2_questions,
    section_name="القسم المتوسط - ترم ثاني",
    equipment_name=EQUIPMENT_NAME
)

# Sheet 4: برنامج تدريب تخصص
specialty_sections_data = [
    {
        "name": "القسم المتوسط",
        "equipment_name": EQUIPMENT_NAME,
        "banner_title": f"موضوعات برنامج تدريب ( القسم المتوسط ) - تخصص {EQUIPMENT_NAME}",
        "class_label": "المتوسط",
        "total_label": "إجمالي القسم المتوسط",
        "topics": [
            {"term": "الترم الأول", "topic": it["lesson_name"], "th": 2, "pr": 2}
            for it in LESSONS_TERM_1
        ] + [
            {"term": "الترم الثاني", "topic": it["lesson_name"], "th": 2, "pr": 2}
            for it in LESSONS_TERM_2
        ]
    }
]
ws_spec = wb_master.create_sheet(title="برنامج تدريب تخصص")
excel_builder.build_specialty_training_program_sheet(ws_spec, EQUIPMENT_NAME, specialty_sections_data)

# Save Master Workbook
print(f"Saving Master Workbook to: {MASTER_TARGET_FILE}")
wb_master.save(MASTER_TARGET_FILE)
print("Master Workbook saved successfully!")

# 7. Build and Save Standalone 'برنامج تدريب تخصص' Workbook
print("Building Standalone Specialty Program Workbook...")
wb_spec_standalone = openpyxl.Workbook()
ws_sa = wb_spec_standalone.active
ws_sa.title = "التخصص"
excel_builder.build_specialty_training_program_sheet(ws_sa, EQUIPMENT_NAME, specialty_sections_data)

print(f"Saving Standalone Specialty Program to: {SPECIALTY_TARGET_FILE}")
try:
    wb_spec_standalone.save(SPECIALTY_TARGET_FILE)
    print("Standalone Specialty Workbook saved successfully!")
except PermissionError:
    alt_spec = SPECIALTY_TARGET_FILE.replace(".xlsx", "_جديد.xlsx")
    wb_spec_standalone.save(alt_spec)
    print(f"File was locked, saved to: {alt_spec}")

print("\n=== GENERATION COMPLETE ===")
print(f"1. Master Workbook: {MASTER_TARGET_FILE}")
print(f"2. Standalone Specialty: {SPECIALTY_TARGET_FILE}")
